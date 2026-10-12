# -*- coding: utf-8 -*-
"""47都道府県 住民一人あたり分野別支出（令和6年度＝R6決算）。
出力: data/pref_expenditure_percapita.json（新規）

分野は市区町村M3（data/m3_expenditure.json）と同一の4分野：
  民生費・教育費・土木費（目的別歳出）／普通建設事業費（性質別歳出）

分子: data/pref_kessan.json の R6 値（目的別＝expenditure_purpose_r6、
      普通建設事業費＝性質別の 普通建設事業費_千円）
分母: 住民基本台帳人口（令和7年1月1日現在）の都道府県計

人口の定義・出所の決め方（2026-09-03の現物確認に基づく）:
  1) 都道府県決算状況調に人口列は存在しない。第1表1-1（001056118.xlsx・列B〜J）、
     第1表1-3（001056121.xlsx・列B〜P）に人口列は無く、「第1 資料のみかた」PDF全文にも
     「人口」の記載が無い（市区町村側の概況表に相当する表が都道府県編に無い）。
  2) よって市区町村M3と同一の人口定義に合わせる。M3の分母は市区町村別決算状況調
     「概況」の住民基本台帳登載人口で、R5決算に対し令6.1.1現在（content-spec-v1.md）。
     同じ市区町村別決算状況調のR6概況（001061669.xlsx）の当該列見出しは
     「住民基本台帳 登載人口（令7.1.1現在）」であることを現物確認した。
     すなわち決算年度内の1月1日時点という同一定義で、R6決算に対応するのは令7.1.1。
  3) 都道府県計の住民基本台帳人口は、同一統計（住民基本台帳に基づく人口、人口動態及び
     世帯数調査）の令7.1.1時点ファイルの都道府県行から取る。市区町村の人口を自前で
     合算しない（政令市とその区が同一ファイルに併存し二重計上になるため）。
     このファイルは data/m1_1a_population_timeseries.json の r7_2025 と同一URL・同一hash。

比較・順位は持たせない（生値と一人あたり額のみ）。
"""
import hashlib
import json
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / "data" / "cache"
OUT = ROOT / "data" / "pref_expenditure_percapita.json"
PREF_KESSAN = ROOT / "data" / "pref_kessan.json"
POP_TIMESERIES = ROOT / "data" / "m1_1a_population_timeseries.json"

TODAY = date.today().isoformat()

FILE_POP = CACHE / "jukikihon_r07_0101_000040306653.xlsx"
URL_POP = "https://www.e-stat.go.jp/stat-search/file-download?statInfId=000040306653&fileKind=0"
PAGE_POP = "https://www.e-stat.go.jp/stat-search/files?toukei=00200241&tstat=000001039591"

# 人口ファイルの列（現物確認：r6が列見出し行、r7=合計、r8以降がデータ）
COL_CODE, COL_PREF, COL_MUNI, COL_POP_TOTAL = 1, 2, 3, 6  # A, B, C, F(人口 計)
POP_DATA_START = 8

# 分野 → (pref_kessan上の取得元, キー)
FIELDS = [
    ("民生費", "expenditure_purpose_r6"),
    ("教育費", "expenditure_purpose_r6"),
    ("土木費", "expenditure_purpose_r6"),
    ("普通建設事業費", "record"),  # 性質別歳出。レコード直下の 普通建設事業費_千円
]


# R6 第5表5-2 の目的別「内訳」列（ヘッダ行を現物確認）。外れ値の裏取りに使う。
SUBITEM_COLS = {
    "民生費": {"1社会福祉費": "O", "2老人福祉費": "P", "3児童福祉費": "Q",
               "4生活保護費": "R", "5災害救助費": "S"},
    "教育費": {"1教育総務費": "BF", "2小学校費": "BG", "3中学校費": "BH", "4高等学校費": "BI",
               "5特別支援学校費": "BJ", "6幼稚園費": "BK", "7社会教育費": "BL",
               "8保健体育費": "BM", "9大学費": "BP"},
    "土木費": {"1土木管理費": "AR", "2道路橋りょう費": "AS", "3河川海岸費": "AT", "4港湾費": "AU",
               "5都市計画費": "AV", "6住宅費": "BA", "7空港費": "BB"},
}
# R6 第6表6-2 の普通建設事業費（EF列・区分「六」）の内訳列。
# EG「うち人件費」は内訳の加算項ではなく内数なので合計に含めない。
SUBITEM_COLS_SEISHITSU = {
    "普通建設事業費": {"1補助事業費": "EH", "2単独事業費": "EN", "3国直轄事業負担金": "ET",
                       "4同級他団体施行事業負担金": "EU", "5受託事業費": "EV"},
}
FILE_R6_5_2 = CACHE / "pref_kessan_5-2.xlsx"
FILE_R6_6_2 = CACHE / "pref_kessan_6-2.xlsx"
R6_5_2_ROWS = (6, 52)
R6_6_2_ROWS = (7, 53)
OUTLIER_RATIO = 2.0


def sha256_of(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def verify_outliers(data: list[dict]) -> dict:
    """一人あたり額が中位値のOUTLIER_RATIO倍以上の項目を抽出し、原典の内訳列で裏を取る。

    列の取り違えなら内訳の合計が款額と合わないため、Σ内訳＝款額の成立を確認できれば
    「実データがそうである」と機械的に言える（目視判断に頼らない）。
    """
    from openpyxl.utils import column_index_from_string

    def load(path: Path, rows: tuple[int, int]):
        ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
        by_name = {}
        for r in range(rows[0], rows[1] + 1):
            v = ws.cell(r, 1).value
            if v is not None:
                by_name[str(v).strip()] = r
        return ws, by_name

    # 目的別（第5表5-2）と性質別（第6表6-2）の両方から内訳を引く
    sheets = {
        "目的別": (*load(FILE_R6_5_2, R6_5_2_ROWS), SUBITEM_COLS),
        "性質別": (*load(FILE_R6_6_2, R6_6_2_ROWS), SUBITEM_COLS_SEISHITSU),
    }

    result: dict = {"判定基準": f"一人あたり額が47件の中位値の{OUTLIER_RATIO}倍以上",
                    "flagged": []}
    for field, _ in FIELDS:
        key = f"{field}_円/人"
        vals = sorted(r[key] for r in data if isinstance(r[key], (int, float)))
        if not vals:
            continue
        median = vals[len(vals) // 2] if len(vals) % 2 else (vals[len(vals) // 2 - 1] + vals[len(vals) // 2]) / 2
        for rec in data:
            v = rec[key]
            if not isinstance(v, (int, float)) or median <= 0 or v / median < OUTLIER_RATIO:
                continue
            item: dict = {
                "pref_name": rec["pref_name"], "field": field,
                "円/人": v, "中位値_円/人": median, "倍率": round(v / median, 2),
                f"{field}_千円": rec[f"{field}_千円"],
            }
            table = next((t for t, (_, _, cols) in sheets.items() if field in cols), None)
            if table is None:
                item["原典の内訳_千円"] = None
                item["判定"] = "内訳列の定義が無いため裏取り未実施"
                result["flagged"].append(item)
                continue
            ws, row_by_name, cols = sheets[table]
            r = row_by_name[rec["pref_name"]]
            breakdown = {}
            for label, col in cols[field].items():
                cv = ws.cell(r, column_index_from_string(col)).value
                breakdown[label] = cv if isinstance(cv, (int, float)) else (
                    "-" if cv is not None and str(cv).strip() else None)
            s = sum(x for x in breakdown.values() if isinstance(x, (int, float)))
            item["内訳の出所"] = f"{table}歳出内訳"
            item["原典の内訳_千円"] = breakdown
            item["Σ内訳_千円"] = s
            item["Σ内訳＝当該費目額"] = (s == rec[f"{field}_千円"])
            item["判定"] = (
                "実データ（原典の内訳列の合計が当該費目額と一致するため、列の取り違えではない）"
                if s == rec[f"{field}_千円"] else "要確認（内訳合計が当該費目額と不一致）"
            )
            result["flagged"].append(item)
    return result


def read_pref_population() -> tuple[dict[str, dict], dict]:
    """都道府県行（市区町村名列が"-"の6桁コード行）から人口を取る。"""
    wb = openpyxl.load_workbook(FILE_POP, data_only=True)
    ws = wb[wb.sheetnames[0]]

    pops: dict[str, dict] = {}
    total_row_value = None
    municipality_rows = 0
    for r in range(POP_DATA_START - 1, ws.max_row + 1):
        code = ws.cell(r, COL_CODE).value
        pref = ws.cell(r, COL_PREF).value
        muni = ws.cell(r, COL_MUNI).value
        pop = ws.cell(r, COL_POP_TOTAL).value
        code_s = "" if code is None else str(code).strip()
        muni_s = "" if muni is None else str(muni).strip()
        pref_s = "" if pref is None else str(pref).strip()

        if pref_s == "合計" and code_s == "-":
            total_row_value = pop
            continue
        if not (code_s.isdigit() and len(code_s) == 6):
            continue
        if muni_s == "-":  # 都道府県行（市区町村名が空＝"-"）
            pops[code_s[:2]] = {"pref_name": pref_s, "団体コード": code_s,
                                "人口_人": pop, "excel_row": r}
        else:
            municipality_rows += 1

    meta = {
        "都道府県行数": len(pops),
        "市区町村行数": municipality_rows,
        "合計行の人口_人": total_row_value,
        "Σ都道府県_人口_人": sum(v["人口_人"] for v in pops.values()),
    }
    meta["合計行との差_人"] = meta["Σ都道府県_人口_人"] - (total_row_value or 0)
    return pops, meta


def main() -> None:
    kessan = json.loads(PREF_KESSAN.read_text(encoding="utf-8"))
    pops, pop_meta = read_pref_population()

    # 人口ファイルの系譜：m1_1a_population_timeseries.json と同一ファイルなら fetched_at を引き継ぐ
    pop_hash = sha256_of(FILE_POP)
    ts = json.loads(POP_TIMESERIES.read_text(encoding="utf-8"))
    prior = next((s for s in ts["sources"] if s.get("source_url") == URL_POP), None)
    pop_fetched_at = (prior["fetched_at"]
                      if prior and prior.get("source_file_hash") == pop_hash else TODAY)
    pop_lineage_note = (
        f"data/m1_1a_population_timeseries.json の r7_2025 と同一URL・同一hash"
        f"（{'一致確認済み・fetched_at引き継ぎ' if prior and prior.get('source_file_hash') == pop_hash else '不一致'}）。"
    )

    kessan_sources = {s["role"]: s for s in kessan["sources"]}
    data: list[dict] = []
    issues: list[dict] = []
    pop_missing: list[dict] = []

    for rec in kessan["data"]:
        code, name = rec["pref_code"], rec["pref_name"]
        p = pops.get(code)
        if p is None:
            pop_missing.append({"pref_code": code, "pref_name": name, "reason": "人口ファイルに都道府県行が無い"})
            continue
        if p["pref_name"] != name:
            issues.append({"pref_code": code, "item": "都道府県名",
                           "pref_kessan": name, "人口ファイル": p["pref_name"]})
        pop = p["人口_人"]

        row: dict = {
            "pref_code": code,
            "pref_name": name,
            "人口_人": pop,
        }
        for field, origin in FIELDS:
            v = rec[f"{field}_千円"] if origin == "record" else rec[origin][field]
            row[f"{field}_千円"] = v
            if not isinstance(v, (int, float)) or isinstance(v, bool):
                issues.append({"pref_code": code, "pref_name": name, "item": field,
                               "raw": v, "reason": "数値でないため一人あたり額を算出しない"})
                row[f"{field}_円/人"] = None
            elif not isinstance(pop, (int, float)) or pop <= 0:
                pop_missing.append({"pref_code": code, "pref_name": name, "reason": f"人口が不正: {pop}"})
                row[f"{field}_円/人"] = None
            else:
                row[f"{field}_円/人"] = round(v * 1000 / pop)
        row["status"] = "確定"
        data.append(row)

    # 一人あたり額が他団体と桁違いになる項目を機械抽出し、原典の内訳列で裏を取る。
    # 「列を取り違えたのか、実データがそうなのか」を人手の目視に委ねないための検査。
    outliers = verify_outliers(data)

    sources = [
        {
            "role": "目的別歳出（民生費・教育費・土木費）",
            "source_url": kessan_sources["R6 目的別歳出内訳（区分「一」〜「二十六」全26項目）"]["source_url"],
            "source_page": kessan_sources["R6 目的別歳出内訳（区分「一」〜「二十六」全26項目）"]["source_page"],
            "source_page_url": kessan_sources["R6 目的別歳出内訳（区分「一」〜「二十六」全26項目）"]["source_page_url"],
            "fetched_at": kessan_sources["R6 目的別歳出内訳（区分「一」〜「二十六」全26項目）"]["fetched_at"],
            "source_file_hash_sha256": kessan_sources["R6 目的別歳出内訳（区分「一」〜「二十六」全26項目）"]["source_file_hash_sha256"],
            "provides": ["民生費_千円", "教育費_千円", "土木費_千円"],
            "note": "値は data/pref_kessan.json の expenditure_purpose_r6 を経由して読んでいる（同一ファイルを二重にパースしない）。",
        },
        {
            "role": "性質別歳出（普通建設事業費）",
            "source_url": kessan_sources["R6 性質別歳出内訳（普通建設事業費）"]["source_url"],
            "source_page": kessan_sources["R6 性質別歳出内訳（普通建設事業費）"]["source_page"],
            "source_page_url": kessan_sources["R6 性質別歳出内訳（普通建設事業費）"]["source_page_url"],
            "fetched_at": kessan_sources["R6 性質別歳出内訳（普通建設事業費）"]["fetched_at"],
            "source_file_hash_sha256": kessan_sources["R6 性質別歳出内訳（普通建設事業費）"]["source_file_hash_sha256"],
            "provides": ["普通建設事業費_千円"],
            "note": "普通建設事業費は目的別ではなく性質別の区分（市区町村M3と同じ扱い）。"
                    "値は data/pref_kessan.json の 普通建設事業費_千円 を経由して読んでいる。",
        },
        {
            "role": "分母（住民基本台帳人口・都道府県計）",
            "source_url": URL_POP,
            "source_page": "e-Stat「住民基本台帳に基づく人口、人口動態及び世帯数調査」"
                           "【総計】市区町村別人口、人口動態及び世帯数（令和7年1月1日時点）都道府県行",
            "source_page_url": PAGE_POP,
            "fetched_at": pop_fetched_at,
            "source_file_hash_sha256": pop_hash,
            "cache_file": FILE_POP.name,
            "provides": ["人口_人"],
            "note": pop_lineage_note +
                    "都道府県決算状況調（都道府県編）に人口列が存在しないことを現物確認した上で、"
                    "市区町村M3と同一の人口定義（住民基本台帳登載人口・決算年度内の1月1日現在）に合わせて選定した。"
                    "R6決算に対応する時点は令7.1.1で、市区町村別決算状況調R6概況（001061669.xlsx）の"
                    "人口列見出し「住民基本台帳 登載人口（令7.1.1現在）」と同一時点。"
                    "都道府県行（市区町村名列が\"-\"の6桁団体コード行）のみを使い、市区町村行の合算はしていない"
                    "（政令市とその行政区が同一ファイルに併存するため合算は二重計上になる）。",
        },
        {
            "role": "中間生成物（分子の読み出し元）",
            "source_url": "data/pref_kessan.json",
            "source_page": None,
            "source_page_url": None,
            "fetched_at": kessan.get("updated"),
            "source_file_hash_sha256": sha256_of(PREF_KESSAN),
            "provides": ["民生費_千円", "教育費_千円", "土木費_千円", "普通建設事業費_千円"],
            "note": "リポジトリ内の中間生成物。上流の原典URL・hashは本配列の他要素（第5表5-2・第6表6-2）に記録済み。"
                    "再生成順序は build_pref_kessan.py → build_pref_expenditure_percapita.py。",
        },
    ]

    out = {
        "schema_version": "1.0",
        "title": "都道府県 住民一人あたり分野別支出",
        "provider": "総務省 / e-Stat",
        "layer": "都道府県（市区町村指標の集計ではなく都道府県自身の歳出）",
        "year": "2024年度（令和6年度）",
        "year_label": "2024年度（令和6年度、FY April 2024 - March 2025）決算",
        "updated": TODAY,
        "unit_amount": "千円（都道府県決算状況調 各表ヘッダ「(単位 千円)」）",
        "unit_percapita": "円/人",
        "population_source": "住民基本台帳登載人口（令和7年1月1日現在・都道府県計）",
        "population_year_gap": "約0年（FY2024中 vs 令7.1.1）",
        "fields": ["民生費", "教育費", "土木費", "普通建設事業費"],
        "notes": [
            "分野の選び方は市区町村M3（data/m3_expenditure.json）と同一の4分野。"
            "民生費・教育費・土木費は目的別歳出、普通建設事業費は性質別歳出から取る（M3と同じ扱い）。",
            "一人あたり額 = 当該費目_千円 × 1000 ÷ 人口_人（四捨五入して円単位）。生値（費目額・人口）も併せて格納する。",
            "都道府県決算状況調（都道府県編）には人口列が存在しない（第1表1-1・第1表1-3・資料のみかたPDFを現物確認）。"
            "このため分母は市区町村M3と同一定義の住民基本台帳人口を、住基統計の都道府県行から取得している。",
            "比較・順位は持たせない（全国平均・都道府県平均・順位・偏差値・複合スコアのいずれも格納しない）。"
            "都道府県ページは市区町村指標の集計層ではなく、都道府県自身の歳出をそのまま示す層である（CLAUDE.md 2026-08-04）。",
            "本ファイルの数値は事実の提示のみ。評価語・色・ランキングは持たせない。",
        ],
        "sources": sources,
        "data": data,
        "data_check_log": {
            "count": {"expected": 47, "actual": len(data)},
            "population_extraction": pop_meta,
            "population_missing": pop_missing,
            "value_issues": issues,
            "outlier_verification": outliers,
        },
    }

    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")

    lines = [
        f"[出力] {OUT.relative_to(ROOT)}",
        f"[件数] {len(data)}/47",
        f"[人口抽出] {json.dumps(pop_meta, ensure_ascii=False)}",
        f"[人口欠損] {len(pop_missing)} 件 {pop_missing}",
        f"[値の問題] {len(issues)} 件 {issues}",
        "",
        f"[外れ値の裏取り] {len(outliers['flagged'])} 件（判定基準: {outliers['判定基準']}）",
        *[f"  {o['pref_name']} {o['field']}: {o['円/人']:,}円/人（中位値の{o['倍率']}倍） → {o['判定']}"
          + (f"\n      原典内訳: {json.dumps(o['原典の内訳_千円'], ensure_ascii=False)}"
             if o.get("原典の内訳_千円") else "")
          for o in outliers["flagged"]],
        "",
        "[岐阜県レコード]",
        json.dumps(next(r for r in data if r["pref_code"] == "21"), ensure_ascii=False, indent=2),
        "",
        "[全47件 一人あたり額（円/人）]",
    ]
    for r in data:
        lines.append(
            f"  {r['pref_code']} {r['pref_name']:<5s} 人口={r['人口_人']:>11,} "
            f"民生費={r['民生費_円/人']:>8,} 教育費={r['教育費_円/人']:>8,} "
            f"土木費={r['土木費_円/人']:>8,} 普通建設事業費={r['普通建設事業費_円/人']:>8,}"
        )
    report = ROOT / "scripts" / "_tmp_report_build_percapita.txt"
    report.write_text("\n".join(lines), encoding="utf-8")
    print(f"[report] {report.name}")


if __name__ == "__main__":
    main()
