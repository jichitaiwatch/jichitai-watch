# -*- coding: utf-8 -*-
"""まちのすがた「産業別就業者構成比」市区町村版JSONに全国順位・類似団体比較を追加する。

対象（docs/content-spec-v1.md §7.2 / docs/similar-group-aggregation-design.md）:
  2a 産業別就業者構成比（第1次・第2次・第3次の3指標。岐阜42市区町村・比較値=全国平均・類似団体
  平均・全国順位・百分位つき）。content-spec-v1.md §7.2の定義上の主比較は「比較=全国」だが、
  2026-09-27（1-1a）・2026-09-28（本件含む）運営者裁定により、既存実装（1-2b/1-2c）と揃え
  「全国順位＋類似団体平均」の形で提示する。

入力（取得済みローカルファイルのみ。ネットワークアクセスをしない）:
  poc/recon-20260828/1-2a/estat_000032214499_kokusei2020_industry_placeofwork.xlsx
    （統計表ID 000032214499・第８表・従業地集計。既存 data/m1_2a_industry.json と同一の全国集約
    ファイル。5桁地域コード＝地域名セル "21201_岐阜市" 形式。列19「1_従業地・通学地による人口」
    が従業地ベースの値。産業列には再掲の合算行 R1_（再掲）第1次産業／R2_（再掲）第2次産業／
    R3_（再掲）第3次産業／T_分類不能の産業／0_総数 が既にあり、これをそのまま使う
    （大分類の個別合算はしない＝既存m1_2a_industry.jsonの値と2026-09-28コード直接確認で一致済み）。
  poc/recon-20260828/0/ruikei_kubun_r06.xlsx（類似団体区分マスタ・市町村一覧シート・1741団体）

更新（既存ファイルを書き換え・新規ファイルは作らない）:
  data/m1_2a_industry.json の各レコード（42件）に comparison フィールド
  （第1次構成比・第2次構成比・第3次構成比の3キー）を追加する。

集計（docs/similar-group-aggregation-design.md §3・§4・§9・§10）:
  母集団M = 区分マスタに載る団体（1741）∩ 全国集約ファイルに値がある団体（産業計>0）。
  構成比 = 再掲合算値 ÷ 0_総数 × 100（既存JSONと同一の計算式・2026-09-28確認）。
  束ね写像：政令指定都市・中核市・施行時特例市・特別区 → "大都市"。他は類団区分の生ラベルをそのまま使う。
  順位：降順 competition ranking（同値は同順位、次の順位を飛ばす）。値が大きいほど上位＝優劣評価ではない。
  percentile = (rank_out_of − rank + 1) / rank_out_of × 100（小数第1位）。
  average_method = "unweighted_mean"（単純算術平均。人口加重等はしない）。

規律:
  - 列位置はハードコードせず、ヘッダのラベル文字列から同定し自己検証する。
  - 表題をコードでassertしてから使う（推測しない）。
  - 評価語・評価色・複合スコアは付けない（中立性・D-4/D-6/D-18）。
"""
from __future__ import annotations

import json
import re
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

INDUSTRY_SRC = ROOT / "poc" / "recon-20260828" / "1-2a" / "estat_000032214499_kokusei2020_industry_placeofwork.xlsx"
RUIKEI_SRC = ROOT / "poc" / "recon-20260828" / "0" / "ruikei_kubun_r06.xlsx"
OUT = ROOT / "data" / "m1_2a_industry.json"

UPDATED = date.today().isoformat()

DAITOSHI_KUBUN = {1, 2, 8, 9}

ANCHOR = {
    "gifu_shi_group": "大都市",
    "gifu_shi_group_n": 128,
    "ogaki_shi_group": "Ⅳ－２",
    "ogaki_shi_group_n": 75,
    "takayama_shi_group": "Ⅱ－１",
    "takayama_shi_group_n": 127,
}

# 2026-09-28コード直接確認によるアンカー（岐阜市・既存JSON値と一致することを別途
# 「既存レコードとの整合確認」で検証するため、ここでは自己算出の構造検証に用いる）。
GIFU_SHI_INDUSTRY_ANCHOR = {"第1次構成比": 1.44, "第2次構成比": 17.55, "第3次構成比": 77.79}

FIELD_TO_SANKEI_LABEL = {
    "第1次構成比": "R1_（再掲）第1次産業",
    "第2次構成比": "R2_（再掲）第2次産業",
    "第3次構成比": "R3_（再掲）第3次産業",
}


class Abort(Exception):
    """欠損・単位不一致・アンカー不一致などで処理を止めるための例外。推測で続行しない。"""


def s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def cell(row: tuple, c: int):
    if row is None or c < 0 or c >= len(row):
        return None
    return row[c]


# --------------------------------------------------------------------------------------
# 区分マスタ（類似団体区分）※既存スクリプトと同一実装
# --------------------------------------------------------------------------------------
def load_ruikei_master() -> dict[str, dict]:
    from openpyxl import load_workbook

    wb = load_workbook(RUIKEI_SRC, read_only=True, data_only=True)
    sheet_name = next((n for n in wb.sheetnames if "市町村一覧" in n), None)
    if sheet_name is None:
        raise Abort(f"ruikei: シート『市町村一覧』が見つからない（実在={wb.sheetnames}）")
    ws = wb[sheet_name]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    print(f"  [ruikei] engine=openpyxl sheet={sheet_name!r}")
    hidx = None
    for i, row in enumerate(rows):
        if s(row[0] if row else None) == "団体コード":
            hidx = i
            break
    if hidx is None:
        raise Abort("ruikei: ヘッダ行（1列目='団体コード'）が見つからない")

    out: dict[str, dict] = {}
    for row in rows[hidx + 1 :]:
        code = s(row[0]) if len(row) > 0 else ""
        if not re.fullmatch(r"\d{6}", code):
            continue
        name = s(row[1])
        danntai_kubun_raw = row[2]
        ruikei_raw = s(row[4])
        if not ruikei_raw:
            continue
        try:
            danntai_kubun = int(danntai_kubun_raw)
        except (TypeError, ValueError):
            raise Abort(f"ruikei: {code} の団体区分が数値でない: {danntai_kubun_raw!r}")
        group_code = "大都市" if danntai_kubun in DAITOSHI_KUBUN else ruikei_raw
        out[code] = {"name": name, "danntai_kubun": danntai_kubun, "ruikei_raw": ruikei_raw, "group_code": group_code}
    if len(out) != 1741:
        raise Abort(f"ruikei: 市町村一覧の件数が1741ではない（実測={len(out)}）")
    print(f"  [ruikei] 市町村一覧 {len(out)}件を読み込み")
    return out


# --------------------------------------------------------------------------------------
# 産業別就業者数（国勢調査 第８表・従業地集計）
# --------------------------------------------------------------------------------------
def resolve_industry_columns(ws):
    """(地域名列, 男女列, 従業上の地位列, 産業3部門列, 産業列, 従業地ベース値列, データ開始行) を返す。"""
    header_rows = list(ws.iter_rows(min_row=1, max_row=10, values_only=True))

    label_row_idx = None
    for i, row in enumerate(header_rows):
        if s(cell(row, 0)) == "地域識別コード":
            label_row_idx = i
            break
    if label_row_idx is None:
        raise Abort("industry: ヘッダ行（『地域識別コード』を含む行）が見つからない")
    label_row = header_rows[label_row_idx]

    def col_of(label: str, row):
        for c, v in enumerate(row):
            if s(v) == label:
                return c
        return None

    area_col = col_of("地域名", label_row)
    sex_col = col_of("男女", label_row)
    status_col = col_of("従業上の地位", label_row)
    sanbumon_col = col_of("産業3部門", label_row)
    sangyo_col = col_of("産業", label_row)
    if None in (area_col, sex_col, status_col, sanbumon_col, sangyo_col):
        raise Abort(
            f"industry: ヘッダ列が見つからない area={area_col} sex={sex_col} "
            f"status={status_col} sanbumon={sanbumon_col} sangyo={sangyo_col}"
        )

    # 「1_従業地・通学地による人口」列（就業者数区分の1つ）。label_row_idx より上の行にある。
    place_col = None
    for i in range(0, label_row_idx):
        for c, v in enumerate(header_rows[i]):
            if s(v) == "1_従業地・通学地による人口":
                place_col = c
                break
        if place_col is not None:
            break
    if place_col is None:
        raise Abort("industry: 『1_従業地・通学地による人口』列が見つからない")

    title_texts = " ".join(s(v) for row in header_rows[:2] for v in row if v)
    if "第８表" not in title_texts or "従業地" not in title_texts:
        raise Abort(f"industry: 表題に想定語を確認できない（実測={title_texts!r}）。推測で続行しない。")
    print(f"  [industry] 表題確認: {title_texts[:60]}...")

    data_start = label_row_idx + 1
    return area_col, sex_col, status_col, sanbumon_col, sangyo_col, place_col, data_start


def load_industry_by_code5() -> dict[str, dict]:
    from openpyxl import load_workbook

    wb = load_workbook(INDUSTRY_SRC, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    area_col, sex_col, status_col, sanbumon_col, sangyo_col, place_col, data_start = resolve_industry_columns(ws)
    print(
        f"  [industry] 列: 地域名列=c{area_col + 1} 男女列=c{sex_col + 1} 従業上の地位列=c{status_col + 1} "
        f"産業3部門列=c{sanbumon_col + 1} 産業列=c{sangyo_col + 1} 従業地ベース値列=c{place_col + 1} "
        f"データ開始行=r{data_start + 1}"
    )

    wanted_sangyo = {"0_総数": "計", "R1_（再掲）第1次産業": "1次", "R2_（再掲）第2次産業": "2次", "R3_（再掲）第3次産業": "3次"}
    raw: dict[str, dict[str, float]] = {}
    n_scanned = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < data_start:
            continue
        if s(cell(row, sex_col)) != "0_総数" or s(cell(row, status_col)) != "0_総数":
            continue
        sangyo = s(cell(row, sangyo_col))
        if sangyo not in wanted_sangyo:
            continue
        area_v = s(cell(row, area_col))
        m = re.match(r"^(\d{5})_", area_v)
        if not m:
            continue
        code5 = m.group(1)
        if code5 == "00000" or code5[-3:] == "000":
            continue  # 全国・都道府県の集計行を除く
        value = cell(row, place_col)
        if not isinstance(value, (int, float)) or isinstance(value, bool):
            continue  # '-' 等の非該当値
        raw.setdefault(code5, {})
        key = wanted_sangyo[sangyo]
        if key in raw[code5]:
            raise Abort(f"industry: コード{code5}の{sangyo}が重複している")
        raw[code5][key] = float(value)
        n_scanned += 1

    out: dict[str, dict] = {}
    for code5, v in raw.items():
        if not all(k in v for k in ("計", "1次", "2次", "3次")):
            continue
        if v["計"] <= 0:
            continue
        out[code5] = {
            "第1次構成比": round(v["1次"] / v["計"] * 100, 2),
            "第2次構成比": round(v["2次"] / v["計"] * 100, 2),
            "第3次構成比": round(v["3次"] / v["計"] * 100, 2),
        }
    print(f"  [industry] {n_scanned}セル走査・市区町村{len(out)}件で3指標完備")
    return out


# --------------------------------------------------------------------------------------
# 集計
# --------------------------------------------------------------------------------------
def ranks_desc(values: list[float]) -> list[int]:
    return [1 + sum(1 for w in values if w > v) for v in values]


def build_comparison_for_field(field: str, ruikei: dict, industry_by_code6: dict) -> tuple[dict[str, dict], dict[str, int]]:
    population: dict[str, dict] = {}
    for code6, meta in ruikei.items():
        v = industry_by_code6.get(code6)
        if v is None:
            continue
        population[code6] = {"value": v[field], "group_code": meta["group_code"]}

    codes = list(population.keys())
    values = [population[c]["value"] for c in codes]
    n_total = len(codes)
    national_average = round(sum(values) / n_total, 1)

    group_values: dict[str, list[float]] = {}
    for c in codes:
        group_values.setdefault(population[c]["group_code"], []).append(population[c]["value"])
    group_average = {g: round(sum(vs) / len(vs), 1) for g, vs in group_values.items()}
    group_n = {g: len(vs) for g, vs in group_values.items()}

    ranks = dict(zip(codes, ranks_desc(values)))

    out: dict[str, dict] = {}
    for c in codes:
        p = population[c]
        rank = ranks[c]
        pct = round((n_total - rank + 1) / n_total * 100, 1)
        out[c] = {
            "value": p["value"],
            "national_average": national_average,
            "group_code": p["group_code"],
            "group_average": group_average[p["group_code"]],
            "rank": rank,
            "rank_out_of": n_total,
            "percentile": pct,
            "average_method": "unweighted_mean",
            "source_note": (
                f"全国{n_total}市区町村中の順位（母集団=類似団体区分マスタ登載団体1741のうち、"
                "本統計に数値がある団体）。類似団体平均は総務省の類似団体区分（束ね後）の単純算術平均。"
                "順位は値の大きい順であり、大小の望ましさを含意しない。"
            ),
        }
    return out, group_n


def check(label, got, want):
    ok = got == want
    print(f"    {'OK ' if ok else 'NG '} {label}: 実測={got!r} 期待={want!r}")
    return ok


def main():
    for p in (INDUSTRY_SRC, RUIKEI_SRC, OUT):
        if not p.exists():
            raise Abort(f"入力ファイルが無い: {p}")

    print("== 既存JSON読み込み ==")
    payload = json.loads(OUT.read_text(encoding="utf-8"))
    existing = payload["data"]
    if len(existing) != 42:
        raise Abort(f"既存JSONのレコード数が42ではない（実測={len(existing)}）")

    print("== 区分マスタ読み込み ==")
    ruikei = load_ruikei_master()

    print("== 全国産業別就業者構成比ファイル読み込み ==")
    industry_by_code5 = load_industry_by_code5()
    industry_by_code6 = {code6: industry_by_code5[code6[:5]] for code6 in ruikei if code6[:5] in industry_by_code5}

    print("== 母集団Mの構築・集計（第1次・第2次・第3次構成比） ==")
    comp1, group_n1 = build_comparison_for_field("第1次構成比", ruikei, industry_by_code6)
    comp2, group_n2 = build_comparison_for_field("第2次構成比", ruikei, industry_by_code6)
    comp3, group_n3 = build_comparison_for_field("第3次構成比", ruikei, industry_by_code6)

    print("== アンカー自己検証 ==")
    gifu_shi_meta = ruikei.get("212016")
    ogaki_shi_meta = ruikei.get("212024")
    takayama_shi_meta = ruikei.get("212032")
    if not (gifu_shi_meta and ogaki_shi_meta and takayama_shi_meta):
        raise Abort("岐阜市/大垣市/高山市のいずれかが区分マスタに無い")
    results = [
        check("岐阜市 束ね後区分", gifu_shi_meta["group_code"], ANCHOR["gifu_shi_group"]),
        check("岐阜市 束ね後区分の団体数", group_n1[gifu_shi_meta["group_code"]], ANCHOR["gifu_shi_group_n"]),
        check("大垣市 束ね後区分", ogaki_shi_meta["group_code"], ANCHOR["ogaki_shi_group"]),
        check("大垣市 束ね後区分の団体数", group_n1[ogaki_shi_meta["group_code"]], ANCHOR["ogaki_shi_group_n"]),
        check("高山市 束ね後区分", takayama_shi_meta["group_code"], ANCHOR["takayama_shi_group"]),
        check("高山市 束ね後区分の団体数", group_n1[takayama_shi_meta["group_code"]], ANCHOR["takayama_shi_group_n"]),
        check("岐阜市 第1次構成比 自己算出値", comp1["212016"]["value"], GIFU_SHI_INDUSTRY_ANCHOR["第1次構成比"]),
        check("岐阜市 第2次構成比 自己算出値", comp2["212016"]["value"], GIFU_SHI_INDUSTRY_ANCHOR["第2次構成比"]),
        check("岐阜市 第3次構成比 自己算出値", comp3["212016"]["value"], GIFU_SHI_INDUSTRY_ANCHOR["第3次構成比"]),
    ]
    if not all(results):
        raise Abort("アンカー値が一致しない。JSONを書かずに中断した。")

    print("== 既存42件との整合確認 ==")
    gifu_codes = sorted(c for c in industry_by_code6 if c.startswith("21"))
    if len(gifu_codes) != 42:
        raise Abort(f"岐阜県内市区町村が{len(gifu_codes)}件（42件ではない）")
    by_area_code = {r["area_code"]: r for r in existing}
    if set(by_area_code.keys()) != set(gifu_codes):
        raise Abort("既存JSONのarea_codeと母集団Mの岐阜42件のキー集合が一致しない。JSONを書かずに中断した。")

    print("== 既存レコードへ comparison を付与（生値の再抽出一致も併せて検証） ==")
    for rec in existing:
        code6 = rec["area_code"]
        c1, c2, c3 = comp1[code6], comp2[code6], comp3[code6]
        for field, c in (("第1次構成比", c1), ("第2次構成比", c2), ("第3次構成比", c3)):
            if abs(rec[field] - c["value"]) > 0.011:
                raise Abort(
                    f"{rec['市区町村名']}（{code6}）: 既存{field}{rec[field]}と全国ファイルからの再抽出値"
                    f"{c['value']}が不一致。JSONを書かずに中断した。"
                )
        rec["comparison"] = {"第1次構成比": c1, "第2次構成比": c2, "第3次構成比": c3}

    payload["updated"] = UPDATED
    payload.setdefault("notes", [])
    prefix = "comparison（全国平均・類似団体平均・全国順位・百分位）は"
    payload["notes"] = [n for n in payload["notes"] if not n.startswith(prefix)]
    payload["notes"].append(
        "comparison（全国平均・類似団体平均・全国順位・百分位）は2026-09-28追加（F-5）。対象は第1次・"
        "第2次・第3次構成比の3指標。docs/similar-group-aggregation-design.md §3・§4の集計段の定義に従う。"
        "母集団M = 区分マスタ（総務省 類似団体別市町村財政指数表・1741団体）∩ 全国集約ファイル"
        "（統計表ID000032214499・従業地集計の再掲合算行）に数値がある団体。"
        "national_average は母集団Mの単純算術平均、group_average は束ね後区分ごとの単純算術平均"
        "（average_method=\"unweighted_mean\"）。束ね写像：政令指定都市・中核市・施行時特例市・特別区は"
        "『大都市』に統合。それ以外は類団区分の原ラベルをそのまま使う。"
        "content-spec-v1.md §7.2の定義上の主比較は「比較=全国」だが、2026-09-27・2026-09-28運営者裁定に"
        "より全国順位＋類似団体平均の形で提示する。"
        "rank は母集団M内での値の降順順位（1=最大値。同値は同順位・次の順位を飛ばすcompetition ranking。"
        "値が大きいほど上位＝優劣評価ではない）。"
        "percentile = (rank_out_of − rank + 1) / rank_out_of × 100（小数第1位）。"
        "評価語・評価色・複合スコアは付与していない（20260808_decisions.md D-4／D-6／D-18）。"
    )

    print("== 書き出し ==")
    with OUT.open("w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)
    print(f"  wrote {OUT.relative_to(ROOT)}  ({len(existing)}件にcomparison付与)")


if __name__ == "__main__":
    try:
        main()
    except Abort as exc:
        print(f"\n!! 中断: {exc}", file=sys.stderr)
        sys.exit(1)
