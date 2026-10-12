# -*- coding: utf-8 -*-
"""くらし「ふるさと納税」カード（F-6）の市区町村版JSONを生成する。

対象（docs/launch-backlog.md F-6・2026-10-07 運営者裁定）:
  岐阜42市区町村の
    ・受入額・受入件数（令和7年度・決算見込）
    ・住民税（市町村民税分）の寄附金税額控除額（令和8年度課税・2026-06-01時点の事前調査）
  を、人口一人あたりに換算した値とあわせて出す。
  比較値＝全国平均・類似団体平均・全国順位・百分位（F-5と同じ形式）。

入力（取得済みローカルファイルのみ。ネットワークアクセスをしない）:
  poc/recon-20261007-furusato/soumu_001084989_r7_ukeire.xlsx   総務省「各自治体の令和7年度受入額等」
  poc/recon-20261007-furusato/soumu_001085012_r8_koujo.xlsx    総務省「各自治体の令和8年度課税における住民税控除額等」
  poc/recon-20260828/1-1a/estat_000040306653_juki_shichoson_pop_hh_r7.xlsx  住基人口（令和7年1月1日）
  poc/recon-20260828/0/ruikei_kubun_r06.xlsx                   類似団体区分マスタ（1741団体）

出力（新規1本）:
  data/m4_5_furusato.json

算出式:
  一人あたり受入額_円   = 受入額_円 ÷ 人口_計（令和7年1月1日住基）  ※整数円に丸め
  一人あたり控除額_円   = 控除額（市町村民税分）_円 ÷ 人口_計         ※整数円に丸め
  人口は、受入期間（令和7年度）・控除の基礎年（令和7年）のどちらにも近い令和7年1月1日を共通に使う。

規律:
  - 列位置はハードコードせず、ヘッダのラベル文字列から同定し自己検証する。
  - 単位（円・件・人）をコードでassertしてから計算する（推測しない）。
  - 受入額は「決算見込」、控除額は「※推計値含む」。生成物のnotesに明記する。
  - 道府県民税分の控除額は市町村の減収ではないため載せない。
  - 差引・収支・評価語・評価色・複合スコアは作らない（中立性・D-4/D-6/D-18）。
  - 母集団・束ね写像・順位の定義は docs/similar-group-aggregation-design.md §3・§4・§9・§10（F-5と同一）。
"""
from __future__ import annotations

import json
import re
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from lineage import build_sources  # noqa: E402

UKEIRE_SRC = ROOT / "poc" / "recon-20261007-furusato" / "soumu_001084989_r7_ukeire.xlsx"
KOUJO_SRC = ROOT / "poc" / "recon-20261007-furusato" / "soumu_001085012_r8_koujo.xlsx"
POP_SRC = ROOT / "poc" / "recon-20260828" / "1-1a" / "estat_000040306653_juki_shichoson_pop_hh_r7.xlsx"
RUIKEI_SRC = ROOT / "poc" / "recon-20260828" / "0" / "ruikei_kubun_r06.xlsx"
OUT = ROOT / "data" / "m4_5_furusato.json"

UPDATED = date.today().isoformat()
PAGE_URL = "https://www.soumu.go.jp/main_sosiki/jichi_zeisei/czaisei/czaisei_seido/furusato/archive/"
UKEIRE_URL = "https://www.soumu.go.jp/main_content/001084989.xlsx"
KOUJO_URL = "https://www.soumu.go.jp/main_content/001085012.xlsx"
POP_URL = "https://www.e-stat.go.jp/stat-search/file-download?statInfId=000040306653&fileKind=0"
POP_PAGE = "https://www.e-stat.go.jp/stat-search/files?toukei=00200241&tstat=000001039591"

DAITOSHI_KUBUN = {1, 2, 8, 9}  # 政令指定都市・特別区・中核市・施行時特例市 → "大都市"
ANCHOR = {  # 設計書§9（岐阜実測）。F-5と同じアンカー。
    "212016": ("大都市", 128),
    "212024": ("Ⅳ－２", 75),
    "212032": ("Ⅱ－１", 127),
}


class Abort(Exception):
    """欠損・単位不一致・アンカー不一致などで処理を止める。推測で続行しない。"""


def s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def norm(v) -> str:
    """改行・全角空白を除いた比較用文字列。"""
    return re.sub(r"[\s　]+", "", s(v))


def load_rows(path: Path, sheet: str | None = None) -> list[list]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def cell(rows, r, c):
    if r < 0 or r >= len(rows) or c < 0 or c >= len(rows[r]):
        return None
    return rows[r][c]


def num(v):
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return None
    return int(v) if float(v).is_integer() else float(v)


# ---------------------------------------------------------------- 区分マスタ
def load_ruikei() -> dict[str, dict]:
    from openpyxl import load_workbook

    wb = load_workbook(RUIKEI_SRC, read_only=True, data_only=True)
    name = next((n for n in wb.sheetnames if "市町村一覧" in n), None)
    if name is None:
        raise Abort(f"ruikei: シート『市町村一覧』が無い（実在={wb.sheetnames}）")
    rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
    h = next((i for i, r in enumerate(rows) if s(cell(rows, i, 0)) == "団体コード"), None)
    if h is None:
        raise Abort("ruikei: ヘッダ行が無い")
    out = {}
    for row in rows[h + 1 :]:
        code = s(row[0]) if row else ""
        if not re.fullmatch(r"\d{6}", code):
            continue
        ruikei_raw = s(row[4])
        if not ruikei_raw:
            continue
        kubun = int(row[2])
        out[code] = {"group_code": "大都市" if kubun in DAITOSHI_KUBUN else ruikei_raw}
    if len(out) != 1741:
        raise Abort(f"ruikei: 件数が1741でない（実測={len(out)}）")
    return out


# ---------------------------------------------------------------- 受入額
def load_ukeire() -> tuple[dict[str, dict], str]:
    rows = load_rows(UKEIRE_SRC)
    title = " ".join(norm(rows[i][0]) for i in range(0, 6) if rows[i] and rows[i][0])
    if "令和７年度" not in title or "決算見込" not in title:
        raise Abort(f"受入額: 表題に『令和７年度』『決算見込』が無い: {title[:80]!r}")
    if "件数：件、金額：円" not in title.replace("　", ""):
        raise Abort("受入額: 単位表記『件数：件、金額：円』が確認できない")
    # ①列（令和7年度に受け入れた寄附件数、金額）を見出し文字列から同定
    r_head = c_cnt = None
    for i in range(0, 20):
        for c, v in enumerate(rows[i] if i < len(rows) else []):
            t = norm(v)
            if t.startswith("①") and "令和７年度に受け入れた寄附件数" in t:
                r_head, c_cnt = i, c
    if r_head is None:
        raise Abort("受入額: 見出し『①令和７年度に受け入れた寄附件数、金額』が見つからない")
    if norm(cell(rows, r_head + 1, c_cnt)) != "件数" or norm(cell(rows, r_head + 1, c_cnt + 1)) != "金額":
        raise Abort("受入額: ①列の下の副見出しが『件数』『金額』でない")
    c_code = next((c for c, v in enumerate(rows[7]) if "団体コード" in norm(v)), None)
    if c_code is None:
        raise Abort("受入額: 団体コード列が見つからない")
    out = {}
    for r in range(r_head + 2, len(rows)):
        code = s(cell(rows, r, c_code))
        if not re.fullmatch(r"\d{5,6}", code):
            continue
        code = code.zfill(6)
        if code in out:
            raise Abort(f"受入額: コード {code} が重複")
        out[code] = {"count": num(cell(rows, r, c_cnt)), "amount": num(cell(rows, r, c_cnt + 1))}
    return out, f"r{r_head + 1}行・c{c_cnt + 1},{c_cnt + 2}列"


# ---------------------------------------------------------------- 控除額
def load_koujo() -> tuple[dict[str, dict], str]:
    rows = load_rows(KOUJO_SRC, "集計表")
    title = norm(rows[1][0]) if len(rows) > 1 and rows[1] else ""
    if "令和８年度課税" not in title:
        raise Abort(f"控除額: 表題に『令和８年度課税』が無い: {title[:60]!r}")
    # 「ふるさと納税に係る寄附金税額控除」→ 下段の「市町村民税」→ 人数／寄附金額／控除額
    r_g = c_g = None
    for i in range(0, 20):
        for c, v in enumerate(rows[i] if i < len(rows) else []):
            if norm(v) == "ふるさと納税に係る寄附金税額控除":
                r_g, c_g = i, c
    if r_g is None:
        raise Abort("控除額: 見出し『ふるさと納税に係る寄附金税額控除』が見つからない")
    r_sub = next((i for i in range(r_g + 1, r_g + 4) if norm(cell(rows, i, c_g)) == "市町村民税"), None)
    if r_sub is None:
        raise Abort("控除額: 『市町村民税』の見出しが見つからない")
    # 見出しの人数／寄附金額／控除額は「市町村民税」見出しの1〜3行下にある（間に空行が入る）
    r_unit = next((i for i in range(r_sub + 1, r_sub + 4) if norm(cell(rows, i, c_g)).startswith("人数")), None)
    if r_unit is None:
        raise Abort("控除額: 『人数』見出し行が見つからない")
    lab = [norm(cell(rows, r_unit, c_g + k)) for k in range(3)]
    if not (lab[0].startswith("人数") and lab[1].startswith("寄附金額") and lab[2].startswith("控除額")):
        raise Abort(f"控除額: 列見出しが人数／寄附金額／控除額でない: {lab}")
    if "（人）" not in lab[0] or "（円）" not in lab[1] or "（円）" not in lab[2]:
        raise Abort(f"控除額: 単位が人／円でない: {lab}")
    c_code = next((c for c, v in enumerate(rows[r_g]) if "団体" in norm(v) and "コード" in norm(v)), None)
    if c_code is None:
        raise Abort("控除額: 団体コード列が見つからない")
    out = {}
    for r in range(r_unit + 1, len(rows)):
        code = s(cell(rows, r, c_code))
        if not re.fullmatch(r"\d{5,6}", code):
            continue
        code = code.zfill(6)  # この表は先頭0が欠けた5桁で入っている行がある
        if code in out:
            raise Abort(f"控除額: コード {code} が重複")
        out[code] = {
            "persons": num(cell(rows, r, c_g)),
            "donation": num(cell(rows, r, c_g + 1)),
            "deduction": num(cell(rows, r, c_g + 2)),
        }
    return out, f"r{r_unit}行（見出し）・c{c_g + 1}〜{c_g + 3}列・控除額見出し={lab[2]!r}"


# ---------------------------------------------------------------- 人口
def load_population() -> dict[str, int]:
    rows = load_rows(POP_SRC)
    title = norm(rows[0][0])
    if "住民基本台帳" not in title or "令和7年1月1日" not in title:
        raise Abort(f"人口: 表題が令和7年1月1日住基でない: {title[:50]!r}")
    h = next((i for i in range(0, 15) if s(cell(rows, i, 0)) == "団体コード"), None)
    if h is None:
        raise Abort("人口: 団体コード行が無い")
    c_pop = None
    for c in range(len(rows[h])):
        if norm(cell(rows, h - 2, c)) == "人口" and norm(cell(rows, h - 1, c)) == "計":
            c_pop = c
    if c_pop is None:
        raise Abort("人口: 『人口／計』列が見つからない")
    if norm(cell(rows, h, c_pop)) != "人":
        raise Abort("人口: 単位が『人』でない")
    out = {}
    for r in range(h + 1, len(rows)):
        code = s(cell(rows, r, 0))
        v = num(cell(rows, r, c_pop))
        if re.fullmatch(r"\d{6}", code) and v is not None:
            out[code] = int(v)
    return out


# ---------------------------------------------------------------- 集計
def ranks_desc(values):
    return [1 + sum(1 for w in values if w > v) for v in values]


def build_comparison(pop_m: dict[str, dict], key: str):
    """pop_m: code -> {group_code, <key>: value}。母集団全体の平均・順位を返す。"""
    codes = list(pop_m)
    vals = [pop_m[c][key] for c in codes]
    n = len(codes)
    nat = round(sum(vals) / n)
    gv: dict[str, list] = {}
    for c in codes:
        gv.setdefault(pop_m[c]["group_code"], []).append(pop_m[c][key])
    gavg = {g: round(sum(v) / len(v)) for g, v in gv.items()}
    gn = {g: len(v) for g, v in gv.items()}
    rk = dict(zip(codes, ranks_desc(vals)))
    return n, nat, gavg, gn, rk


def main():
    for p in (UKEIRE_SRC, KOUJO_SRC, POP_SRC, RUIKEI_SRC):
        if not p.exists():
            raise Abort(f"入力ファイルが無い: {p}")

    print("== 読み込み ==")
    ruikei = load_ruikei()
    ukeire, u_info = load_ukeire()
    koujo, k_info = load_koujo()
    pop = load_population()
    print(f"  区分マスタ={len(ruikei)} 受入額={len(ukeire)}（{u_info}） 控除額={len(koujo)}（{k_info}） 人口={len(pop)}")

    print("== 母集団Mの構築 ==")
    M = {}
    drop = {"受入額なし": [], "控除額なし": [], "人口なし": [], "値が数値でない": [], "人口0": []}
    for code, meta in ruikei.items():
        u, k, p = ukeire.get(code), koujo.get(code), pop.get(code)
        if u is None:
            drop["受入額なし"].append(code); continue
        if k is None:
            drop["控除額なし"].append(code); continue
        if p is None:
            drop["人口なし"].append(code); continue
        if u["amount"] is None or u["count"] is None or k["deduction"] is None:
            drop["値が数値でない"].append(code); continue
        if p <= 0:
            drop["人口0"].append(code); continue
        M[code] = {
            "group_code": meta["group_code"],
            "pop": p,
            "count": u["count"], "amount": u["amount"],
            "persons": k["persons"], "donation": k["donation"], "deduction": k["deduction"],
            "per_in": round(u["amount"] / p),
            "per_out": round(k["deduction"] / p),
        }
    for kk, v in drop.items():
        if v:
            print(f"  母集団外（{kk}）: {len(v)}件 {v[:8]}")
    print(f"  母集団M = {len(M)}件")
    if len(M) != 1741:
        raise Abort(f"母集団Mが1741でない（{len(M)}）。除外の理由は上記。JSONを書かずに中断した。")

    print("== 比較値の算出 ==")
    n_in, nat_in, gavg_in, gn, rk_in = build_comparison(M, "per_in")
    n_out, nat_out, gavg_out, _, rk_out = build_comparison(M, "per_out")

    print("== アンカー自己検証（設計書§9） ==")
    ok_all = True
    for code, (g, n) in ANCHOR.items():
        got = (M[code]["group_code"], gn[M[code]["group_code"]])
        ok = got == (g, n)
        ok_all &= ok
        print(f"    {'OK ' if ok else 'NG '} {code}: 実測={got} 期待={(g, n)}")
    if not ok_all:
        raise Abort("アンカー不一致。JSONを書かずに中断した。")

    note_rank = (
        f"全国{n_in}市区町村中の順位（母集団=類似団体区分マスタ登載団体1741のうち、受入額・控除額・人口の全てに数値がある団体）。"
        "類似団体平均は総務省の類似団体区分（束ね後）の単純算術平均。順位は値の大きい順であり、大小の望ましさを含意しない。"
    )

    def comp(code, key, rk, nat, gavg):
        m = M[code]
        return {
            "value": m[key],
            "national_average": nat,
            "group_code": m["group_code"],
            "group_average": gavg[m["group_code"]],
            "rank": rk[code],
            "rank_out_of": n_in,
            "percentile": round((n_in - rk[code] + 1) / n_in * 100, 1),
            "average_method": "unweighted_mean",
            "source_note": note_rank,
        }

    gifu = sorted(c for c in M if c.startswith("21"))
    if len(gifu) != 42:
        raise Abort(f"岐阜県内が{len(gifu)}件（42件でない）")
    names = {r["area_code"]: r["市区町村名"] for r in json.load(open(ROOT / "data" / "m1_kodomo_iryohi.json", encoding="utf-8"))["data"]}
    if set(names) != set(gifu):
        raise Abort("岐阜42件のコードがdata/m1_kodomo_iryohi.jsonと一致しない")

    data = []
    for code in gifu:
        m = M[code]
        data.append({
            "area_code": code,
            "市区町村名": names[code],
            "受入_件数": m["count"],
            "受入_金額_円": m["amount"],
            "控除_市町村民税_人数": m["persons"],
            "控除_市町村民税_寄附金額_円": m["donation"],
            "控除_市町村民税_控除額_円": round(m["deduction"]),
            "人口_計": m["pop"],
            "一人あたり受入額_円": m["per_in"],
            "一人あたり控除額_円": m["per_out"],
            "comparison": {
                "一人あたり受入額_円": comp(code, "per_in", rk_in, nat_in, gavg_in),
                "一人あたり控除額_円": comp(code, "per_out", rk_out, nat_out, gavg_out),
            },
        })

    sources = build_sources(OUT, [
        {"role": "受入額（令和7年度・決算見込）", "cache_path": UKEIRE_SRC, "source_url": UKEIRE_URL, "source_page": PAGE_URL},
        {"role": "住民税控除額（令和8年度課税）", "cache_path": KOUJO_SRC, "source_url": KOUJO_URL, "source_page": PAGE_URL},
        {"role": "人口（住基・令和7年1月1日）", "cache_path": POP_SRC, "source_url": POP_URL, "source_page": POP_PAGE},
    ])
    for e in sources[:2]:
        e["_provenance_note"] = ("2026-10-07 運営者がブラウザでダウンロードし会話に添付→poc/recon-20261007-furusato/に保存。"
                                 "fetched_atは本スクリプトの実行時刻（ダウンロード時刻の記録は無い）。")

    notes = [
        "受入額・受入件数は総務省『ふるさと納税に関する現況調査』令和8年度実施（2026-07-31公表）の『各自治体の令和7年度受入額等』"
        "（令和7年4月1日〜令和8年3月31日の**決算見込**）。確定値ではない。件数：件、金額：円。",
        "住民税控除額は同調査の『各自治体の令和8年度課税における住民税控除額等』（令和8年6月1日時点の事前調査）の"
        f"『ふるさと納税に係る寄附金税額控除』のうち**市町村民税**の控除額（円）。表の見出しに『※推計値含む』の注記がある（{k_info}）。"
        "道府県民税分は市町村の減収ではないため載せていない。受入額と期間が異なる（受入＝令和7年度、控除＝令和8年度課税で、"
        "令和7年1〜12月分の寄附に対応）。",
        "一人あたり値＝金額÷人口_計（住民基本台帳・令和7年1月1日）。整数円に丸め。受入・控除の双方に同じ人口を使う。",
        "受入額と控除額の差引・収支は算出していない。受入が多いこと・控除が多いことのいずれにも優劣の評価を付けない（D-4／D-6／D-18）。",
        "比較値（comparison）は docs/similar-group-aggregation-design.md §3・§4 の集計段の定義に従う（F-5と同一）。"
        f"母集団M={n_in}団体。national_average／group_average は母集団Mの単純算術平均（整数円）。"
        "束ね写像：政令指定都市・中核市・施行時特例市・特別区は『大都市』に統合。"
        "rank は値の降順（1=最大値）、同値は同順位・次の順位を飛ばす。percentile=(rank_out_of−rank+1)/rank_out_of×100。",
        "団体コードは控除額表で先頭の0が欠けた5桁の行があるため6桁に補正して結合した。",
        "入力は取得済みのローカルファイルのみで、生成時に再ダウンロードはしていない。",
    ]
    payload = {
        "sources": sources,
        # 画面の「対象年度」表示は lib/staleness.ts formatPeriodLabel が先頭の和暦年度だけを拾うため受入額の年度のみとし、
        # 控除（令和8年度課税）・人口（令和7年1月1日）の期間はカード本文と notes に明記する。
        "period": "令和7年度（2025）",
        "updated": UPDATED,
        "notes": notes,
        "data": data,
    }
    with OUT.open("w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)
    print(f"== 書き出し == {OUT.relative_to(ROOT)} ({len(data)}件)")

    print("== サニティ：岐阜42件の一人あたり受入額（円）上位3・下位3 ==")
    for r in sorted(data, key=lambda r: -r["一人あたり受入額_円"])[:3] + sorted(data, key=lambda r: -r["一人あたり受入額_円"])[-3:]:
        c = r["comparison"]["一人あたり受入額_円"]
        d = r["comparison"]["一人あたり控除額_円"]
        print(f"  {r['市区町村名']}: 受入/人={r['一人あたり受入額_円']:,}円(全国{c['rank']}位) 控除/人={r['一人あたり控除額_円']:,}円(全国{d['rank']}位) 区分={c['group_code']}")
    print(f"  全国平均: 受入/人={nat_in:,}円 控除/人={nat_out:,}円")


if __name__ == "__main__":
    try:
        main()
    except Abort as exc:
        print(f"\n!! 中断: {exc}", file=sys.stderr)
        sys.exit(1)
