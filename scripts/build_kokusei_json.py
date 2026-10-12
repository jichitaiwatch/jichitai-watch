# -*- coding: utf-8 -*-
"""まちのすがた 国勢調査系3指標の生値JSONを、市区町村（岐阜42）と都道府県（47＋全国）で生成する。

対象（content-spec-v1.md §7.11 確定表）:
  #2 単身世帯割合            … 令和2年国勢調査 人口等基本集計 第9-1表（statInfId=000032142495 / シート b09_01）
  #4 産業別就労者数・構成比  … 令和2年国勢調査 従業地・通学地集計 第8表（statInfId=000032214499 / シート e08-01）
  #5 上位3業種               … #4の派生（従業地就業者数の降順。評価語は付けない）
  #6 昼夜間人口比率          … 令和2年国勢調査 従業地・通学地集計 第1-1表（statInfId=000032214141 / シート e01_01）

出力（新規6本。既存JSONは読まない・書き換えない）:
  data/m1_1b_single_household.json   data/pref_1b_single_household.json
  data/m1_2a_industry.json           data/pref_2a_industry.json
  data/m1_2b_daynight.json           data/pref_2b_daynight.json

規律:
  - 入力は取得済みローカルファイルのみ。ネットワークアクセスをしない。
  - 列位置はハードコードせず、ヘッダ上部のラベル文字列から同定する。
  - 比較値（全国平均・類団平均・順位）は算出しない（集計段で付与）。
  - アンカー値と内部整合（部門合計＝産業計）を自己検証し、不一致なら例外で中断する。
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "poc" / "recon-20260828"
OUT = ROOT / "data"

UPDATED = date.today().isoformat()
PERIOD = "令和2年(2020)国勢調査"
CODE_RE = re.compile(r"^(\d{5})_(.+)$")
DAIBUNRUI_RE = re.compile(r"^([A-Z])_")

# 入力ファイルの定義。fetched_at と ledger_hash は poc/recon-20260828/ledger-batch1.csv の実測値。
INPUTS = {
    "single": {
        "path": SRC / "1-1b" / "estat_000032142495_kokusei2020_family_type_shichoson.xlsx",
        "stat_inf_id": "000032142495",
        "stat_name": "令和2年国勢調査 人口等基本集計",
        "source_url": "https://www.e-stat.go.jp/stat-search/file-download?statInfId=000032142495&fileKind=0",
        "source_page": "https://www.e-stat.go.jp/stat-search/files?page=1&layout=datalist&toukei=00200521&tstat=000001136464",
        "fetched_at": "2026-08-28T20:06:41+09:00",
        "ledger_hash": "343c7a2797071cb63215bcaee640e2b43c6b212efb3b31dd379cd84206981f25",
    },
    "industry": {
        "path": SRC / "1-2a" / "estat_000032214499_kokusei2020_industry_placeofwork.xlsx",
        "stat_inf_id": "000032214499",
        "stat_name": "令和2年国勢調査 従業地・通学地による人口・就業状態等集計",
        "source_url": "https://www.e-stat.go.jp/stat-search/file-download?statInfId=000032214499&fileKind=0",
        "source_page": (
            "https://www.e-stat.go.jp/stat-search/files?page=1&layout=dataset&toukei=00200521"
            "&tstat=000001136464&tclass1=000001136469"
        ),
        "fetched_at": "2026-08-28T20:11:39+09:00",
        "ledger_hash": "15bda4748c49cef07b3d5d364844ea0a1064a2667e938b27e0a73cc281f3cf6f",
    },
    "daynight": {
        "path": SRC / "1-2b" / "estat_000032214141_kokusei2020_daynight.xlsx",
        "stat_inf_id": "000032214141",
        "stat_name": "令和2年国勢調査 従業地・通学地による人口・就業状態等集計",
        "source_url": "https://www.e-stat.go.jp/stat-search/file-download?statInfId=000032214141&fileKind=0",
        "source_page": (
            "https://www.e-stat.go.jp/stat-search/files?page=1&layout=dataset&toukei=00200521"
            "&tstat=000001136464&tclass1=000001136469&stat_infid=000032214141"
        ),
        "fetched_at": "2026-08-28T20:03:48+09:00",
        "ledger_hash": "8c1a3b731617c421ddae723439531985308332d8941849e8a9b6dfa6b72b37a5",
    },
}


class Abort(Exception):
    """アンカー不一致・構造不一致で処理を止めるための例外。推測で続行しない。"""


def s(v) -> str:
    return "" if v is None else str(v).strip()


def num(v, where: str):
    """数値化。'-' や空は欠損として中断させる（本3指標では欠損は想定されない）。"""
    t = s(v)
    if t == "" or t == "-" or t == "***" or t == "X":
        raise Abort(f"欠損値を検出（0に丸めず中断）: {where} 値={t!r}")
    t = t.replace(",", "")
    try:
        if re.fullmatch(r"-?\d+", t):
            return int(t)
        return float(t)
    except ValueError as exc:
        raise Abort(f"数値化できない値: {where} 値={t!r}") from exc


def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def open_sheet(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    return wb, wb[wb.sheetnames[0]]


def read_header(ws, limit: int = 40):
    """1列目が '地域識別コード' の行をヘッダ行とし、(ヘッダ行番号, キー列マップ, 上部ラベル行群) を返す。"""
    label_rows: list[list[str]] = []
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=limit, values_only=True), start=1):
        vals = [s(v) for v in row]
        if vals and vals[0] == "地域識別コード":
            keycols = {lab: c for c, lab in enumerate(vals, start=1) if lab and lab != " "}
            return i, keycols, label_rows
        label_rows.append(vals)
    raise Abort(f"ヘッダ行（1列目='地域識別コード'）が {limit} 行以内に見つからない")


def cols_with_label(label_rows: list[list[str]], target: str) -> set[int]:
    """上部ラベル行群のいずれかの行で target と完全一致する列番号の集合。"""
    found = set()
    for vals in label_rows:
        for c, v in enumerate(vals, start=1):
            if v == target:
                found.add(c)
    return found


def only_col(label_rows, target: str, what: str) -> int:
    got = cols_with_label(label_rows, target)
    if len(got) != 1:
        raise Abort(f"{what}: ラベル {target!r} に一致する列が {sorted(got)}（1列に確定できない）")
    return next(iter(got))


def table_title_of(label_rows: list[list[str]], what: str) -> str:
    """ヘッダより上の行から '第…表' で始まる表題行を拾う（read_onlyシートでも安全）。"""
    for vals in label_rows:
        head = vals[0] if vals else ""
        if head.startswith("第") and "表" in head:
            return head
    raise Abort(f"{what}: 表題行（'第…表'）が見つからない")


def one_col(cands: set[int], what: str) -> int:
    if len(cands) != 1:
        raise Abort(f"{what}: 列を1つに確定できない（候補={sorted(cands)}）")
    return next(iter(cands))


def need(keycols: dict[str, int], label: str, what: str) -> int:
    if label not in keycols:
        raise Abort(f"{what}: キー列 {label!r} がヘッダ行に無い（実在={sorted(keycols)}）")
    return keycols[label]


def classify(area_cell: str):
    """地域名セル '21201_岐阜市' → ('21201', '岐阜市', 種別)。種別は 'national'/'pref'/'muni'/None。"""
    m = CODE_RE.match(area_cell)
    if not m:
        return None, None, None
    code, name = m.group(1), m.group(2)
    if code == "00000":
        return code, "全国", "national"
    if code.endswith("000"):
        return code, name, "pref"
    return code, name, "muni"


DASH = "-"


def num_or_dash(v, where: str):
    """原表の『-』（該当数値なし＝構造的ゼロ）はリテラル文字列 '-' のまま返す。

    既存 data/pref_fiscal_indicators.json の規約に合わせる（'-'＝該当なし。null＝確認できず、
    とは区別する）。0 に変換しない。それ以外の非数値トークンは num() が中断させる。
    """
    if s(v) == DASH:
        return DASH
    return num(v, where)


def zero(v):
    """算術用。'-'（構造的ゼロ）を 0 として扱う。格納値は変換しない。"""
    return 0 if v == DASH else v


def in_scope(code: str, kind: str) -> bool:
    """対象地域のみを解析する。全国／47都道府県／岐阜県内の市区町村（先頭2桁=21）。

    範囲外の市区町村は値を読まない。国勢調査R2には対象外自治体の欠損（例 07546_双葉町の
    一般世帯数='-'）が実在するため、範囲外まで数値化すると対象外の欠損で止まってしまう。
    範囲内で欠損が出た場合は num() が中断させる。
    """
    if kind in ("pref", "national"):
        return True
    return kind == "muni" and code.startswith("21")


def pct(numerator, denominator, where: str):
    if not denominator:
        raise Abort(f"構成比の分母が0または欠損: {where}")
    return round(numerator / denominator * 100, 2)


# --------------------------------------------------------------------------------------
# #2 単身世帯割合
# --------------------------------------------------------------------------------------
def build_single(meta):
    wb, ws = open_sheet(meta["path"])
    hidx, keycols, label_rows = read_header(ws)
    c_area = need(keycols, "地域名", "single")
    c_cat = need(keycols, "世帯の家族類型", "single")
    # 一般世帯数の「総数」列＝グループ見出し'一般世帯数'かつ小見出し'0_総数'の列
    c_val = one_col(
        cols_with_label(label_rows, "一般世帯数") & cols_with_label(label_rows, "0_総数"),
        "single 一般世帯数の総数列",
    )
    table_title = table_title_of(label_rows, "single")

    picked: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=hidx + 1, values_only=True):
        area = s(row[c_area - 1]) if len(row) >= c_area else ""
        code, name, kind = classify(area)
        if kind is None or not in_scope(code, kind):
            continue
        cat = s(row[c_cat - 1]) if len(row) >= c_cat else ""
        if cat not in ("0_総数", "3_単独世帯"):
            continue
        rec = picked.setdefault(code, {"code": code, "name": name, "kind": kind})
        key = "一般世帯数" if cat == "0_総数" else "単独世帯数"
        rec[key] = num(row[c_val - 1], f"single {area} {cat}")
    wb.close()

    for code, rec in picked.items():
        for k in ("一般世帯数", "単独世帯数"):
            if k not in rec:
                raise Abort(f"single: {code} に {k} の行が無い")
        rec["単身世帯割合"] = pct(rec["単独世帯数"], rec["一般世帯数"], f"single {code}")

    print(f"  [single] 列: 地域名=c{c_area} 家族類型=c{c_cat} 一般世帯数総数=c{c_val} / 地域数={len(picked)}")
    return picked, {"table_title": table_title, "cols": {"area": c_area, "cat": c_cat, "val": c_val}}


def single_values(rec):
    return {
        "一般世帯数": rec["一般世帯数"],
        "単独世帯数": rec["単独世帯数"],
        "単身世帯割合": rec["単身世帯割合"],
    }


# --------------------------------------------------------------------------------------
# #4 産業別就労者数・構成比 ＋ #5 上位3業種
# --------------------------------------------------------------------------------------
SECTOR_ROWS = {
    "R1": "第1次就業者数",
    "R2": "第2次就業者数",
    "R3": "第3次就業者数",
}


def build_industry(meta):
    wb, ws = open_sheet(meta["path"])
    hidx, keycols, label_rows = read_header(ws)
    c_area = need(keycols, "地域名", "industry")
    c_sex = need(keycols, "男女", "industry")
    c_pos = need(keycols, "従業上の地位", "industry")
    c_ind = need(keycols, "産業", "industry")
    # 常住地ベースと従業地ベースを判別し、従業地ベースの総数列を採る（§7.11＝従業地集計）
    c_home = only_col(label_rows, "0_常住地による人口", "industry 常住地ベース列")
    c_work = only_col(label_rows, "1_従業地・通学地による人口", "industry 従業地ベース列")
    if c_home == c_work:
        raise Abort("industry: 常住地列と従業地列が同一に解決された（判別不能）")
    table_title = table_title_of(label_rows, "industry")

    picked: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=hidx + 1, values_only=True):
        if len(row) < c_work:
            continue
        area = s(row[c_area - 1])
        code, name, kind = classify(area)
        if kind is None or not in_scope(code, kind):
            continue
        if s(row[c_sex - 1]) != "0_総数" or s(row[c_pos - 1]) != "0_総数":
            continue
        ind = s(row[c_ind - 1])
        rec = picked.setdefault(code, {"code": code, "name": name, "kind": kind, "daibunrui": {}})
        where = f"industry {area} {ind}"
        if ind == "0_総数":
            rec["産業計"] = num(row[c_work - 1], where)
        elif ind.split("_", 1)[0] in SECTOR_ROWS:
            rec[SECTOR_ROWS[ind.split("_", 1)[0]]] = num(row[c_work - 1], where)
        else:
            m = DAIBUNRUI_RE.match(ind)
            if m:  # 大分類（A〜T）。'01_総合農業' 等の中分類は数字始まりなので除外される
                rec["daibunrui"][m.group(1)] = {
                    "産業コード": m.group(1),
                    "産業名": ind.split("_", 1)[1],
                    "就業者数": num_or_dash(row[c_work - 1], where),
                }
    wb.close()

    top3_changes_if_T_included = []
    for code, rec in picked.items():
        for k in ["産業計", *SECTOR_ROWS.values()]:
            if k not in rec:
                raise Abort(f"industry: {code} に {k} の行が無い")
        dai = rec.pop("daibunrui")
        if "T" not in dai:
            raise Abort(f"industry: {code} に T_分類不能の産業 の行が無い")
        rec["分類不能就業者数"] = dai["T"]["就業者数"]

        # 内部整合1: 3部門（再掲行）＋分類不能 ＝ 産業計（'-' は算術上0として計算）
        parts = sum(zero(rec[k]) for k in SECTOR_ROWS.values()) + zero(rec["分類不能就業者数"])
        if parts != zero(rec["産業計"]):
            raise Abort(f"industry: {code} 部門合計{parts} ≠ 産業計{rec['産業計']}")
        # 内部整合2: 大分類A〜Sの合計 ＝ 3部門合計（再掲行と大分類の突合）
        dai_sum = sum(zero(v["就業者数"]) for k, v in dai.items() if k != "T")
        sector_sum = sum(zero(rec[k]) for k in SECTOR_ROWS.values())
        if dai_sum != sector_sum:
            raise Abort(f"industry: {code} 大分類合計{dai_sum} ≠ 3部門合計{sector_sum}")

        for label, key in (("第1次", "第1次構成比"), ("第2次", "第2次構成比"), ("第3次", "第3次構成比")):
            rec[key] = pct(zero(rec[f"{label}就業者数"]), zero(rec["産業計"]), f"industry {code} {label}")

        # #5 上位3業種：大分類A〜S（'0_総数'・部門再掲行・中分類・分類不能Tを除く）を従業地就業者数の降順で3件
        ranked = sorted(
            (v for k, v in dai.items() if k != "T"),
            key=lambda d: (-zero(d["就業者数"]), d["産業コード"]),
        )
        rec["上位3業種"] = [dict(d) for d in ranked[:3]]
        with_t = sorted(dai.values(), key=lambda d: (-zero(d["就業者数"]), d["産業コード"]))[:3]
        if [d["産業コード"] for d in with_t] != [d["産業コード"] for d in rec["上位3業種"]]:
            top3_changes_if_T_included.append(code)

    print(
        f"  [industry] 列: 地域名=c{c_area} 男女=c{c_sex} 従業上の地位=c{c_pos} 産業=c{c_ind} "
        f"常住地=c{c_home} 従業地=c{c_work} / 地域数={len(picked)}"
    )
    print(f"  [industry] Tを含めると上位3業種が変わる地域: {top3_changes_if_T_included or 'なし'}")
    return picked, {
        "table_title": table_title,
        "t_affects_top3": top3_changes_if_T_included,
        "cols": {"home": c_home, "work": c_work},
    }


def industry_values(rec):
    out = {"産業計": rec["産業計"]}
    for label in ("第1次", "第2次", "第3次"):
        out[f"{label}就業者数"] = rec[f"{label}就業者数"]
    out["分類不能就業者数"] = rec["分類不能就業者数"]
    for label in ("第1次", "第2次", "第3次"):
        out[f"{label}構成比"] = rec[f"{label}構成比"]
    out["上位3業種"] = rec["上位3業種"]
    return out


# --------------------------------------------------------------------------------------
# #6 昼夜間人口比率
# --------------------------------------------------------------------------------------
def build_daynight(meta):
    wb, ws = open_sheet(meta["path"])
    hidx, keycols, label_rows = read_header(ws)
    c_area = need(keycols, "地域名", "daynight")
    c_sex = need(keycols, "男女", "daynight")
    c_age = need(keycols, "年齢", "daynight")
    c_ratio = only_col(label_rows, "昼夜間人口比率", "daynight 昼夜間人口比率列")
    table_title = table_title_of(label_rows, "daynight")

    picked: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=hidx + 1, values_only=True):
        if len(row) < c_ratio:
            continue
        area = s(row[c_area - 1])
        code, name, kind = classify(area)
        if kind is None or not in_scope(code, kind):
            continue
        if s(row[c_sex - 1]) != "0_総数" or s(row[c_age - 1]) != "00_総数":
            continue
        if code in picked:
            raise Abort(f"daynight: {code} の総数行が重複している")
        picked[code] = {
            "code": code,
            "name": name,
            "kind": kind,
            "昼夜間人口比率": num(row[c_ratio - 1], f"daynight {area}"),
        }
    wb.close()
    print(f"  [daynight] 列: 地域名=c{c_area} 男女=c{c_sex} 年齢=c{c_age} 昼夜間人口比率=c{c_ratio} / 地域数={len(picked)}")
    return picked, {"table_title": table_title, "cols": {"ratio": c_ratio}}


def daynight_values(rec):
    return {"昼夜間人口比率": rec["昼夜間人口比率"]}


# --------------------------------------------------------------------------------------
# 書き出し
# --------------------------------------------------------------------------------------
def source_entry(meta, table_title: str, hash_key: str, file_hash: str):
    return {
        "source_url": meta["source_url"],
        "source_page": meta["source_page"],
        "stat_name": meta["stat_name"],
        "table_title": table_title,
        "statInfId": meta["stat_inf_id"],
        "fetched_at": meta["fetched_at"],
        hash_key: file_hash,
    }


COMMON_NOTES = [
    "本ファイルは生値のみを格納する。比較値（全国平均・類団平均・全国順位）は含めず、集計段で付与する。",
    "地域コードは国勢調査の5桁（例 21201=岐阜市、21000=岐阜県、00000=全国）。住民基本台帳の6桁体系（例 212016）とは別体系であり、"
    "data/m1_1a_population_timeseries.json の 市区町村コード と直接結合できない（先頭5桁での結合が必要）。",
    "入力は取得済みのローカルファイルのみを使用し、本ファイル生成時に再ダウンロードはしていない。",
]

MACHI_NOTE = "対象は岐阜県内の市区町村42件（地域コード先頭2桁=21・末尾3桁≠000）。都道府県行・全国行は含まない。"
PREF_NOTE = (
    "対象は47都道府県（地域コード末尾3桁=000）に全国行（00000_全国）を加えた48件。"
    "全国行は後段の全国比較用に同梱したものであり、都道府県の1件として数えない。"
)

NOTES = {
    "single": [
        "単身世帯割合(％) = 世帯の家族類型『3_単独世帯』の一般世帯数 ÷ 『0_総数』の一般世帯数 × 100。小数第2位で丸め。"
        "分子・分母の原数値を併載しているため再計算可能。",
        "『一般世帯数』は世帯員の年齢による世帯の種類が『0_総数』の列（＝一般世帯数の総数）を採った。単位は世帯。",
    ],
    "industry": [
        "就業者数は全て『従業地・通学地による人口』ベース（§7.11＝従業地集計）。常住地ベースの列は使用していない。",
        "男女=『0_総数』・従業上の地位=『0_総数』の行群から取得した（15歳以上就業者）。単位は人。",
        "第1次／第2次／第3次の就業者数は、統計表が持つ再掲行『R1_（再掲）第1次産業』『R2_（再掲）第2次産業』"
        "『R3_（再掲）第3次産業』の値をそのまま格納した（自前で合算していない）。生成時に大分類A〜Sの合計と一致することを全地域で検証済み。",
        "産業計には『T_分類不能の産業』が含まれる。このため第1次・第2次・第3次の構成比の合計は100％にならない。"
        "差分を追えるよう 分類不能就業者数 を併載した（第1次＋第2次＋第3次＋分類不能＝産業計 を全地域で検証済み）。",
        "構成比(％) = 各部門の就業者数 ÷ 産業計 × 100。小数第2位で丸め。原数値を併載しているため再計算可能。",
        "産業大分類の \"-\" はe-Stat原典の構造的ゼロ（該当計上なし）。0円/0人と同義のリテラル文字列として格納し、集計時は0として扱う。"
        "null（＝確認できず）とは区別する（既存 data/pref_fiscal_indicators.json の規約に同じ）。生成時の整合検証では0として計算し、成立を確認済み。",
        "上位3業種は産業大分類A〜Sを従業地就業者数の降順に並べた上位3件で、コードと名称と就業者数のみを持つ。"
        "『0_総数』・部門再掲行（R1〜R3）・中分類行（『01_総合農業』等）・『T_分類不能の産業』は対象外。"
        "評価語・順位語は付けない（降順の機械抽出のみ）。",
    ],
    "daynight": [
        "昼夜間人口比率は統計表に算出済みの値（男女=『0_総数』・年齢=『00_総数』の行）をそのまま格納した。自前計算していない。",
        "100が昼夜均衡。単位表記は原表にも無い。全国は100。",
    ],
}


def write_json(path: Path, sources, notes, data):
    payload = {
        "sources": sources,
        "period": PERIOD,
        "updated": UPDATED,
        "notes": notes,
        "data": data,
    }
    with path.open("w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)
    print(f"  wrote {path.relative_to(ROOT)}  ({len(data)} records)")


def emit(kind_key, meta, picked, extra, values_fn, machi_name, pref_name):
    file_hash = sha256_of(meta["path"])
    if file_hash != meta["ledger_hash"]:
        raise Abort(
            f"{kind_key}: 入力ファイルのsha256がledger-batch1.csvの記録と不一致\n"
            f"  実測={file_hash}\n  台帳={meta['ledger_hash']}"
        )
    title = extra["table_title"]

    munis = sorted((r for r in picked.values() if r["kind"] == "muni" and r["code"].startswith("21")),
                   key=lambda r: r["code"])
    if len(munis) != 42:
        raise Abort(f"{kind_key}: 岐阜県内市区町村が {len(munis)} 件（42件ではない）")
    machi_rows = [
        {"市区町村コード": r["code"], "都道府県": "岐阜県", "市区町村名": r["name"], **values_fn(r)}
        for r in munis
    ]
    write_json(
        OUT / machi_name,
        [source_entry(meta, title, "source_file_hash", file_hash)],
        [*NOTES[kind_key], MACHI_NOTE, *COMMON_NOTES],
        machi_rows,
    )

    prefs = sorted((r for r in picked.values() if r["kind"] in ("pref", "national")),
                   key=lambda r: r["code"])
    n_pref = sum(1 for r in prefs if r["kind"] == "pref")
    if n_pref != 47 or len(prefs) != 48:
        raise Abort(f"{kind_key}: 都道府県 {n_pref} 件・全体 {len(prefs)} 件（47＋全国=48ではない）")
    pref_rows = [
        {"都道府県コード": r["code"], "都道府県名": r["name"], **values_fn(r)}
        for r in prefs
    ]
    write_json(
        OUT / pref_name,
        [source_entry(meta, title, "source_file_hash_sha256", file_hash)],
        [*NOTES[kind_key], PREF_NOTE, *COMMON_NOTES],
        pref_rows,
    )
    return picked


# --------------------------------------------------------------------------------------
# アンカー自己検証
# --------------------------------------------------------------------------------------
def check(label, got, want):
    ok = got == want
    print(f"    {'OK ' if ok else 'NG '} {label}: 実測={got!r} 期待={want!r}")
    return ok


def main():
    for meta in INPUTS.values():
        if not meta["path"].exists():
            raise Abort(f"入力ファイルが無い: {meta['path']}")

    print("== 抽出 ==")
    single, single_x = build_single(INPUTS["single"])
    industry, industry_x = build_industry(INPUTS["industry"])
    daynight, daynight_x = build_daynight(INPUTS["daynight"])

    print("== アンカー自己検証 ==")
    gifu_ind = industry["21000"]
    e_mfg = next((d for d in gifu_ind["上位3業種"] if d["産業コード"] == "E"), None)
    results = [
        check("#2 岐阜県(21000) 一般世帯総数", single["21000"]["一般世帯数"], 779029),
        check("#2 岐阜県(21000) 単独世帯数", single["21000"]["単独世帯数"], 228715),
        check("#2 全国(00000) 一般世帯総数", single["00000"]["一般世帯数"], 55704949),
        check("#4 岐阜県(21000) 製造業E 従業地就業者数", e_mfg["就業者数"] if e_mfg else None, 222550),
        check("#6 岐阜県(21000) 昼夜間人口比率", daynight["21000"]["昼夜間人口比率"], 96.62078),
        check("#6 全国(00000) 昼夜間人口比率", daynight["00000"]["昼夜間人口比率"], 100),
    ]
    if not all(results):
        raise Abort("アンカー値が一致しない。JSONを書かずに中断した。")

    print("== 書き出し ==")
    emit("single", INPUTS["single"], single, single_x, single_values,
         "m1_1b_single_household.json", "pref_1b_single_household.json")
    emit("industry", INPUTS["industry"], industry, industry_x, industry_values,
         "m1_2a_industry.json", "pref_2a_industry.json")
    emit("daynight", INPUTS["daynight"], daynight, daynight_x, daynight_values,
         "m1_2b_daynight.json", "pref_2b_daynight.json")

    print("== 岐阜県・全国の該当値 ==")
    for code in ("21000", "00000"):
        nm = single[code]["name"]
        si, ind, dn = single[code], industry[code], daynight[code]
        print(f"  [{code} {nm}]")
        print(f"    #2 一般世帯数={si['一般世帯数']} 単独世帯数={si['単独世帯数']} 単身世帯割合={si['単身世帯割合']}%")
        print(f"    #4 産業計={ind['産業計']} 第1次={ind['第1次就業者数']}({ind['第1次構成比']}%) "
              f"第2次={ind['第2次就業者数']}({ind['第2次構成比']}%) 第3次={ind['第3次就業者数']}({ind['第3次構成比']}%) "
              f"分類不能={ind['分類不能就業者数']}")
        print(f"    #5 上位3業種=" + " / ".join(
            f"{d['産業コード']}_{d['産業名']}:{d['就業者数']}" for d in ind["上位3業種"]))
        print(f"    #6 昼夜間人口比率={dn['昼夜間人口比率']}")
    print(f"  岐阜市(21201) #6 昼夜間人口比率={daynight['21201']['昼夜間人口比率']}")


if __name__ == "__main__":
    try:
        main()
    except Abort as exc:
        print(f"\n!! 中断: {exc}", file=sys.stderr)
        sys.exit(1)
