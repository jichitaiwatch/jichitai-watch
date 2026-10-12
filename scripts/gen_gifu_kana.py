"""
岐阜県の読み（カナ）抽出。

一次ソース: 総務省「全国地方公共団体コード」（都道府県コード及び市区町村コード）
  公表ページ: https://www.soumu.go.jp/denshijiti/code.html
出力: data/gifu_kana.json（手書き禁止。本スクリプトが単一ソース）

カナは総務省コード表の生値（半角カタカナ）を保存する。全角化・ローマ字化しない。
推測で読みを埋めない。取得できない団体は中断して報告する。

  python scripts/gen_gifu_kana.py
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from datetime import date, datetime, timezone, timedelta
from pathlib import Path

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE_PAGE = "https://www.soumu.go.jp/denshijiti/code.html"
CACHE_PATH = ROOT / "data" / "cache" / "soumu_lg_code.xlsx"
OUTPUT_PATH = ROOT / "data" / "gifu_kana.json"
PREF_FILTER = "21"
USER_AGENT = "Mozilla/5.0 (jichitai-watch gifu-kana fetch)"
JST = timezone(timedelta(hours=9))

EXPECTED = {
    "210005": "岐阜県",
    "212016": "岐阜市",
    "212024": "大垣市",
    "212032": "高山市",
    "212041": "多治見市",
    "212059": "関市",
    "212067": "中津川市",
    "212075": "美濃市",
    "212083": "瑞浪市",
    "212091": "羽島市",
    "212105": "恵那市",
    "212113": "美濃加茂市",
    "212121": "土岐市",
    "212130": "各務原市",
    "212148": "可児市",
    "212156": "山県市",
    "212164": "瑞穂市",
    "212172": "飛騨市",
    "212181": "本巣市",
    "212199": "郡上市",
    "212202": "下呂市",
    "212211": "海津市",
    "213021": "岐南町",
    "213039": "笠松町",
    "213411": "養老町",
    "213616": "垂井町",
    "213624": "関ケ原町",
    "213811": "神戸町",
    "213829": "輪之内町",
    "213837": "安八町",
    "214019": "揖斐川町",
    "214035": "大野町",
    "214043": "池田町",
    "214213": "北方町",
    "215015": "坂祝町",
    "215023": "富加町",
    "215031": "川辺町",
    "215040": "七宗町",
    "215058": "八百津町",
    "215066": "白川町",
    "215074": "東白川村",
    "215210": "御嵩町",
    "216046": "白川村",
}


def cell_str(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def normalize_code(raw: object) -> str | None:
    s = cell_str(raw).replace(" ", "").replace("　", "")
    if s.endswith(".0"):
        s = s[:-2]
    if not re.fullmatch(r"\d{5,6}", s):
        return None
    return s.zfill(6) if len(s) == 5 else s


def resolve_xlsx_url() -> str:
    print(f"[fetch] {SOURCE_PAGE}")
    resp = requests.get(SOURCE_PAGE, timeout=60, headers={"User-Agent": USER_AGENT})
    resp.raise_for_status()
    hrefs = re.findall(r'href=["\']([^"\']+\.xlsx)["\']', resp.text, flags=re.I)
    abs_urls: list[str] = []
    for href in hrefs:
        if href.startswith("http"):
            abs_urls.append(href)
        else:
            abs_urls.append("https://www.soumu.go.jp" + (href if href.startswith("/") else "/" + href))
    if not abs_urls:
        raise SystemExit("公表ページから xlsx のダウンロードURLを特定できませんでした")
    print(f"[xlsx] {abs_urls[0]}")
    return abs_urls[0]


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def parse_source_as_of(wb) -> str:
    for name in wb.sheetnames:
        m = re.search(r"R(\d+)\.(\d+)\.(\d+)", name)
        if m:
            year = 2018 + int(m.group(1))
            return date(year, int(m.group(2)), int(m.group(3))).isoformat()
    raise SystemExit("コード表シート名から基準時点（source_as_of）を読めませんでした")


def detect_columns(ws) -> dict[str, int]:
    for r in range(1, min(16, (ws.max_row or 1) + 1)):
        labels = {c: cell_str(ws.cell(r, c).value) for c in range(1, (ws.max_column or 1) + 1)}
        joined = " ".join(labels.values())
        if "団体コード" not in joined and "コード" not in joined:
            continue
        mapping: dict[str, int] = {}
        for c, lab in labels.items():
            lab_one = lab.replace("\n", "")
            if "団体コード" in lab_one or lab_one in {"コード", "code"}:
                mapping.setdefault("code", c)
            elif "都道府県名" in lab_one and "カナ" in lab_one:
                mapping["pref_kana"] = c
            elif "市区町村名" in lab_one and "カナ" in lab_one:
                mapping["city_kana"] = c
            elif "都道府県名" in lab_one:
                mapping.setdefault("pref_name", c)
            elif "市区町村名" in lab_one:
                mapping.setdefault("city_name", c)
        if {"code", "pref_name", "city_name", "pref_kana", "city_kana"} <= mapping.keys():
            mapping["header_row"] = r
            return mapping
    raise SystemExit("ヘッダー行（団体コード・漢字名・カナ列）を特定できませんでした")


def choose_current_orgs_sheet(wb):
    for name in wb.sheetnames:
        if "政令" in name or "改正" in name or "仕様" in name:
            continue
        return wb[name]
    return wb[wb.sheetnames[0]]


def extract_gifu_kana(ws) -> list[dict[str, str]]:
    cols = detect_columns(ws)
    start = cols["header_row"] + 1
    rows: list[dict[str, str]] = []
    seen: set[str] = set()
    missing_kana: list[str] = []
    for r in range(start, (ws.max_row or 0) + 1):
        code = normalize_code(ws.cell(r, cols["code"]).value)
        if not code or not code.startswith(PREF_FILTER):
            continue
        if len(code) != 6:
            raise SystemExit(f"area_code が6桁でない: {code!r}")
        pref_name = cell_str(ws.cell(r, cols["pref_name"]).value)
        city_name = cell_str(ws.cell(r, cols["city_name"]).value)
        pref_kana = cell_str(ws.cell(r, cols["pref_kana"]).value)
        city_kana = cell_str(ws.cell(r, cols["city_kana"]).value)
        if pref_name == "" and city_name == "":
            continue
        if code in seen:
            continue
        seen.add(code)
        is_pref = city_name == ""
        name = pref_name if is_pref else city_name
        kana = pref_kana if is_pref else city_kana
        if not name:
            raise SystemExit(f"名称が空: area_code={code}")
        if not kana:
            missing_kana.append(name)
            continue
        rows.append({"area_code": code, "name": name, "kana": kana})
    if missing_kana:
        raise SystemExit("カナを取得できなかった団体名: " + " / ".join(missing_kana))
    rows.sort(key=lambda r: r["area_code"])
    return rows


def validate(rows: list[dict[str, str]]) -> None:
    got = {r["area_code"]: r["name"] for r in rows}
    expected_codes = set(EXPECTED)
    got_codes = set(got)
    missing = sorted(expected_codes - got_codes)
    extra = sorted(got_codes - expected_codes)
    errors: list[str] = []
    if len(rows) != 43:
        errors.append(f"件数不一致: {len(rows)} 件（期待 43）")
    if missing:
        names = [EXPECTED[c] for c in missing]
        errors.append("欠けている団体名: " + " / ".join(f"{n}({c})" for n, c in zip(names, missing)))
    if extra:
        errors.append("余剰の area_code: " + " / ".join(f"{got[c]}({c})" for c in extra))
    name_mismatch = []
    for code, expected_name in EXPECTED.items():
        if code in got and got[code] != expected_name:
            name_mismatch.append(f"{code} ファイル={got[code]!r} 検算一覧={expected_name!r}")
    if name_mismatch:
        errors.append("名称不一致（ファイルの生値を優先し、推測では埋めない）: " + " / ".join(name_mismatch))
    if errors:
        print("VALIDATION FAILED:", file=sys.stderr)
        for e in errors:
            print(f"  - {e}", file=sys.stderr)
        raise SystemExit(1)


def main() -> None:
    xlsx_url = resolve_xlsx_url()
    if not CACHE_PATH.exists() or CACHE_PATH.stat().st_size == 0:
        raise SystemExit(f"キャッシュが無い: {CACHE_PATH}。先に実ファイルをバイトのまま保存すること。")

    file_hash = sha256_file(CACHE_PATH)
    # fetched_at は実ファイルを書き込んだ時刻（=ダウンロード完了時刻）。実行時刻は入れない（data-schema.md §5.3）
    fetched_at = datetime.fromtimestamp(CACHE_PATH.stat().st_mtime, tz=JST).isoformat(timespec="seconds")
    if OUTPUT_PATH.exists():
        prev = json.loads(OUTPUT_PATH.read_text(encoding="utf-8"))
        prev_sources = prev.get("sources") or []
        if prev_sources:
            prev_hash = prev_sources[0].get("source_file_hash")
            prev_fetched = prev_sources[0].get("fetched_at")
            if prev_hash == file_hash and prev_fetched:
                fetched_at = prev_fetched

    wb = load_workbook(CACHE_PATH, read_only=True, data_only=True)
    sheet = choose_current_orgs_sheet(wb)
    source_as_of = parse_source_as_of(wb)
    rows = extract_gifu_kana(sheet)
    print(f"[sheet] {sheet.title}  gifu={len(rows)}  source_as_of={source_as_of}")
    validate(rows)

    payload = {
        "sources": [
            {
                "source_url": xlsx_url,
                "source_page": SOURCE_PAGE,
                "fetched_at": fetched_at,
                "source_file_hash": file_hash,
            }
        ],
        "source_as_of": source_as_of,
        "notes": [
            "検索の読み一致用。カナは総務省コード表の生値・半角カタカナ",
            "総務省「全国地方公共団体コード」（都道府県コード及び市区町村コード）から岐阜県(21)分を抽出。推測による読みの補完はしない。",
        ],
        "data": rows,
    }
    OUTPUT_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"[write] {OUTPUT_PATH}  bytes={OUTPUT_PATH.stat().st_size}  hash={file_hash}")


if __name__ == "__main__":
    main()
