"""
全国の読み（カナ）抽出。

一次ソース: 総務省「全国地方公共団体コード」（都道府県コード及び市区町村コード）
  公表ページ: https://www.soumu.go.jp/denshijiti/code.html
出力: data/municipality_kana.json（手書き禁止。本スクリプトが単一ソース）

カナは総務省コード表の生値（半角カタカナ）を保存する。全角化・ローマ字化しない。
推測で読みを埋めない。取得できない行は中断して団体名・コードを報告する。
政令市の区は除外せず全件保存する。area_code 重複は同一内容なら1件に畳む。

  python scripts/gen_municipality_kana.py
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from datetime import date, datetime, timezone, timedelta
from pathlib import Path

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE_PAGE = "https://www.soumu.go.jp/denshijiti/code.html"
CACHE_PATH = ROOT / "data" / "cache" / "soumu_lg_code.xlsx"
OUTPUT_PATH = ROOT / "data" / "municipality_kana.json"
GIFU_KANA_PATH = ROOT / "data" / "gifu_kana.json"
USER_AGENT = "Mozilla/5.0 (jichitai-watch municipality-kana fetch)"
JST = timezone(timedelta(hours=9))

EXPECTED_PREF_COUNT = 47
GIFU_MIN_CODES = {
    "210005",
    "212016",
    "216046",
}


def cell_str(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def normalize_code(raw: object) -> str | None:
    s = cell_str(raw).replace(" ", "").replace("　", "")
    if s.endswith(".0"):
        s = s[:-2]
    if not re.fullmatch(r"\d{5,6}", s):
        return None
    return s.zfill(6) if len(s) == 5 else s


def resolve_xlsx_url() -> str:
    print(f"[fetch] {SOURCE_PAGE}")
    resp = requests.get(SOURCE_PAGE, timeout=60, headers={"User-Agent": USER_AGENT})
    resp.raise_for_status()
    hrefs = re.findall(r'href=["\']([^"\']+\.xlsx)["\']', resp.text, flags=re.I)
    abs_urls: list[str] = []
    for href in hrefs:
        if href.startswith("http"):
            abs_urls.append(href)
        else:
            abs_urls.append("https://www.soumu.go.jp" + (href if href.startswith("/") else "/" + href))
    if not abs_urls:
        raise SystemExit("公表ページから xlsx のダウンロードURLを特定できませんでした")
    print(f"[xlsx] {abs_urls[0]}")
    return abs_urls[0]


def download_xlsx(url: str, dest: Path) -> None:
    print(f"[download] {url}")
    resp = requests.get(url, timeout=120, headers={"User-Agent": USER_AGENT})
    resp.raise_for_status()
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(resp.content)
    print(f"[saved] {dest} bytes={len(resp.content)}")


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def parse_source_as_of(wb) -> str:
    for name in wb.sheetnames:
        m = re.search(r"R(\d+)\.(\d+)\.(\d+)", name)
        if m:
            year = 2018 + int(m.group(1))
            return date(year, int(m.group(2)), int(m.group(3))).isoformat()
    raise SystemExit("コード表シート名から基準時点（source_as_of）を読めませんでした")


def detect_columns(ws) -> dict[str, int]:
    for r in range(1, min(16, (ws.max_row or 1) + 1)):
        labels = {c: cell_str(ws.cell(r, c).value) for c in range(1, (ws.max_column or 1) + 1)}
        joined = " ".join(labels.values())
        if "団体コード" not in joined and "コード" not in joined:
            continue
        mapping: dict[str, int] = {}
        for c, lab in labels.items():
            lab_one = lab.replace("\n", "")
            if "団体コード" in lab_one or lab_one in {"コード", "code"}:
                mapping.setdefault("code", c)
            elif "都道府県名" in lab_one and ("カナ" in lab_one or "ｶﾅ" in lab_one):
                mapping["pref_kana"] = c
            elif "市区町村名" in lab_one and ("カナ" in lab_one or "ｶﾅ" in lab_one):
                mapping["city_kana"] = c
            elif "都道府県名" in lab_one:
                mapping.setdefault("pref_name", c)
            elif "市区町村名" in lab_one:
                mapping.setdefault("city_name", c)
        if {"code", "pref_name", "city_name", "pref_kana", "city_kana"} <= mapping.keys():
            mapping["header_row"] = r
            return mapping
    raise SystemExit(f"ヘッダー行を特定できませんでした: sheet={ws.title!r}")


def code_table_sheets(wb) -> list:
    sheets = []
    for name in wb.sheetnames:
        if "改正" in name or "仕様" in name:
            continue
        sheets.append(wb[name])
    if not sheets:
        raise SystemExit("コード表シートが見つかりませんでした")
    return sheets


def extract_sheet_rows(ws) -> tuple[list[dict[str, str]], int, list[str]]:
    cols = detect_columns(ws)
    start = cols["header_row"] + 1
    rows: list[dict[str, str]] = []
    missing: list[str] = []
    source_rows = 0
    for r in range(start, (ws.max_row or 0) + 1):
        code = normalize_code(ws.cell(r, cols["code"]).value)
        if not code:
            continue
        source_rows += 1
        if len(code) != 6:
            raise SystemExit(f"area_code が6桁でない: {code!r} sheet={ws.title!r} row={r}")
        pref_name = cell_str(ws.cell(r, cols["pref_name"]).value)
        city_name = cell_str(ws.cell(r, cols["city_name"]).value)
        pref_kana = cell_str(ws.cell(r, cols["pref_kana"]).value)
        city_kana = cell_str(ws.cell(r, cols["city_kana"]).value)
        if pref_name == "" and city_name == "":
            missing.append(f"{code} (名称空) sheet={ws.title} row={r}")
            continue
        is_pref = city_name == ""
        name = pref_name if is_pref else city_name
        kana = pref_kana if is_pref else city_kana
        if not name:
            missing.append(f"{code} (名称空) sheet={ws.title} row={r}")
            continue
        if not kana:
            missing.append(f"{name}({code}) sheet={ws.title} row={r}")
            continue
        rows.append(
            {
                "area_code": code,
                "pref_code": code[:2],
                "level": "prefecture" if is_pref else "municipality",
                "name": name,
                "kana": kana,
            }
        )
    return rows, source_rows, missing


def merge_rows(sheet_rows: list[tuple[str, list[dict[str, str]]]]) -> tuple[list[dict[str, str]], dict]:
    by_code: dict[str, dict[str, str]] = {}
    overlap = 0
    conflicts: list[str] = []
    added_by_sheet: dict[str, int] = {}
    for sheet_name, rows in sheet_rows:
        added = 0
        for rec in rows:
            code = rec["area_code"]
            prev = by_code.get(code)
            if prev is None:
                by_code[code] = rec
                added += 1
                continue
            overlap += 1
            key = ("level", "name", "kana", "pref_code")
            if any(prev[k] != rec[k] for k in key):
                conflicts.append(
                    f"{code} {prev['name']}/{prev['kana']} vs {rec['name']}/{rec['kana']} (sheet={sheet_name})"
                )
        added_by_sheet[sheet_name] = added
    if conflicts:
        raise SystemExit("同一 area_code で内容が食い違う行: " + " / ".join(conflicts))
    data = sorted(by_code.values(), key=lambda r: r["area_code"])
    stats = {"overlap": overlap, "added_by_sheet": added_by_sheet}
    return data, stats


def validate(rows: list[dict[str, str]], source_total: int, unique_total: int) -> None:
    prefs = [r for r in rows if r["level"] == "prefecture"]
    munis = [r for r in rows if r["level"] == "municipality"]
    errors: list[str] = []
    if len(prefs) != EXPECTED_PREF_COUNT:
        errors.append(f"都道府県数が {len(prefs)} 件（期待 {EXPECTED_PREF_COUNT}）")
    pref_codes = {r["pref_code"] for r in prefs}
    if len(pref_codes) != EXPECTED_PREF_COUNT:
        errors.append(f"都道府県 pref_code のユニーク数が {len(pref_codes)}")
    if len(rows) != unique_total:
        errors.append(f"出力件数 {len(rows)} がユニオン {unique_total} と不一致")
    codes = [r["area_code"] for r in rows]
    if len(codes) != len(set(codes)):
        errors.append("area_code 重複が出力に残っている")
    if any(len(r["area_code"]) != 6 for r in rows):
        errors.append("6桁でない area_code がある")
    missing_gifu = sorted(c for c in GIFU_MIN_CODES if c not in set(codes))
    if missing_gifu:
        errors.append("岐阜の既知コードが欠けている: " + ",".join(missing_gifu))
    if source_total < unique_total:
        errors.append(f"ソース行数 {source_total} がユニオン {unique_total} より少ない")
    if errors:
        print("VALIDATION FAILED:", file=sys.stderr)
        for e in errors:
            print(f"  - {e}", file=sys.stderr)
        raise SystemExit(1)
    print(
        json.dumps(
            {
                "total": len(rows),
                "prefecture": len(prefs),
                "municipality": len(munis),
                "source_rows": source_total,
                "unique": unique_total,
            },
            ensure_ascii=True,
        )
    )


def resolve_fetched_at(file_hash: str) -> str:
    fetched_at = datetime.fromtimestamp(CACHE_PATH.stat().st_mtime, tz=JST).isoformat(timespec="seconds")
    for path in (OUTPUT_PATH, GIFU_KANA_PATH):
        if not path.exists():
            continue
        prev = json.loads(path.read_text(encoding="utf-8"))
        prev_sources = prev.get("sources") or []
        if not prev_sources:
            continue
        if prev_sources[0].get("source_file_hash") == file_hash and prev_sources[0].get("fetched_at"):
            return prev_sources[0]["fetched_at"]
    return fetched_at


def main() -> None:
    xlsx_url = resolve_xlsx_url()
    if not CACHE_PATH.exists() or CACHE_PATH.stat().st_size == 0:
        download_xlsx(xlsx_url, CACHE_PATH)

    file_hash = sha256_file(CACHE_PATH)
    fetched_at = resolve_fetched_at(file_hash)

    wb = load_workbook(CACHE_PATH, read_only=False, data_only=True)
    source_as_of = parse_source_as_of(wb)
    sheets = code_table_sheets(wb)

    sheet_rows: list[tuple[str, list[dict[str, str]]]] = []
    source_counts: dict[str, int] = {}
    missing: list[str] = []
    for ws in sheets:
        rows, source_n, miss = extract_sheet_rows(ws)
        sheet_rows.append((ws.title, rows))
        source_counts[ws.title] = source_n
        missing.extend(miss)
        print(f"[sheet] {ws.title} source_rows={source_n} parsed={len(rows)}")

    if missing:
        print("カナまたは名称を取得できなかった行:", file=sys.stderr)
        for item in missing:
            print(f"  - {item}", file=sys.stderr)
        raise SystemExit(1)

    data, stats = merge_rows(sheet_rows)
    source_total = sum(source_counts.values())
    unique_total = len(data)
    print(
        json.dumps(
            {"source_counts": source_counts, "overlap": stats["overlap"], "added_by_sheet": stats["added_by_sheet"]},
            ensure_ascii=True,
        )
    )
    validate(data, source_total, unique_total)

    payload = {
        "sources": [
            {
                "source_url": xlsx_url,
                "source_page": SOURCE_PAGE,
                "fetched_at": fetched_at,
                "source_file_hash": file_hash,
            }
        ],
        "source_as_of": source_as_of,
        "notes": [
            "検索の読み一致用。カナは総務省コード表の生値・半角カタカナ。全国全件。",
            "総務省「全国地方公共団体コード」（都道府県コード及び市区町村コード）から抽出。現在の団体シートと政令指定都市シートを取り込み、政令市の区は除外しない。推測による読みの補完はしない。",
        ],
        "data": data,
    }
    OUTPUT_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"[write] {OUTPUT_PATH} bytes={OUTPUT_PATH.stat().st_size} n={len(data)} hash={file_hash}")


if __name__ == "__main__":
    main()
