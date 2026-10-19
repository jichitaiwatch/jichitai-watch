# -*- coding: utf-8 -*-
import json
import os
import sys
from pathlib import Path

# 県引数（2026-10-17・D-24 共通化）: --pref 23 など。対象は自治体マスタの市区町村行（scripts/pref_args.py）。
# 岐阜（既定）の中間出力は従来の out_merged.json、それ以外は out_merged_{県コード}.json（互いに上書きしない）。
sys.path.insert(0, str(Path(__file__).resolve().parent))
from pref_args import DEFAULT_PREF, load_master_cities, load_master_pref_name, load_master_wards, parse_pref  # noqa: E402

PREF = parse_pref(sys.argv[1:])
CITIES = load_master_cities(PREF)
PREF_NAME = load_master_pref_name(PREF)
SUFFIX = "" if PREF == DEFAULT_PREF else f"_{PREF}"

# 原本と中間出力は poc/recon-20260830-1-1a-timeseries/ にある。相対パスをそのフォルダ基準にそろえる
os.chdir(Path(__file__).resolve().parents[1] / "poc" / "recon-20260830-1-1a-timeseries")
import openpyxl
import xlrd

FILES = [
    {
        "key": "r2_2020",
        "reiwa_label": "令和2年",
        "date": "2020-01-01",
        "path": "estat_000031971203_juki_shichoson_pop_hh_r2_2020.xls",
        "engine": "xlrd",
        "stat_inf_id": "000031971203",
    },
    {
        "key": "r3_2021",
        "reiwa_label": "令和3年",
        "date": "2021-01-01",
        "path": "estat_000040306659_juki_shichoson_pop_hh_r3_2021.xlsx",
        "engine": "openpyxl",
        "stat_inf_id": "000040306659",
    },
    {
        "key": "r4_2022",
        "reiwa_label": "令和4年",
        "date": "2022-01-01",
        "path": "estat_000032224636_juki_shichoson_pop_hh_r4_2022.xlsx",
        "engine": "openpyxl",
        "stat_inf_id": "000032224636",
    },
    {
        "key": "r5_2023",
        "reiwa_label": "令和5年",
        "date": "2023-01-01",
        "path": "estat_000040306647_juki_shichoson_pop_hh_r5_2023.xlsx",
        "engine": "openpyxl",
        "stat_inf_id": "000040306647",
    },
    {
        "key": "r6_2024",
        "reiwa_label": "令和6年",
        "date": "2024-01-01",
        "path": "estat_000040306672_juki_shichoson_pop_hh_r6_2024.xlsx",
        "engine": "openpyxl",
        "stat_inf_id": "000040306672",
    },
    {
        "key": "r7_2025",
        "reiwa_label": "令和7年",
        "date": "2025-01-01",
        "path": "../recon-20260828/1-1a/estat_000040306653_juki_shichoson_pop_hh_r7.xlsx",
        "engine": "openpyxl",
        "stat_inf_id": "000040306653",
    },
    {
        "key": "r8_2026",
        "reiwa_label": "令和8年",
        "date": "2026-01-01",
        # 総務省サイト掲載ファイル（e-Stat 経由ではない）。ファイルID 000892952
        "path": "../recon-20261012-juki-r8/soumu_000892952_juki_shichoson_pop_hh_r8.xlsx",
        "engine": "openpyxl",
        "stat_inf_id": None,
    },
]


def load_rows_openpyxl(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = []
    for r in range(1, ws.max_row + 1):
        rows.append([ws.cell(row=r, column=c).value for c in range(1, 8)])
    return rows


def load_rows_xlrd(path):
    wb = xlrd.open_workbook(path)
    sh = wb.sheet_by_index(0)
    rows = []
    for r in range(sh.nrows):
        row = []
        for c in range(7):
            v = sh.cell_value(r, c) if c < sh.ncols else None
            row.append(v)
        rows.append(row)
    return rows


def find_header_row(rows):
    for i, row in enumerate(rows):
        if row and row[0] == "団体コード":
            return i
    raise ValueError("header row (団体コード) not found")


def is_real_municipality(code, name):
    if not name:
        return False
    if name == "-":
        return False
    if name.endswith("郡"):
        return False
    return True


def extract_gifu(fmeta, codes=None):  # 名前は従来のまま。実際は PREF の県を抜く。codes 省略時は自治体マスタの市区町村
    if fmeta["engine"] == "xlrd":
        rows = load_rows_xlrd(fmeta["path"])
    else:
        rows = load_rows_openpyxl(fmeta["path"])

    header_idx = find_header_row(rows)
    data_rows = rows[header_idx + 1 :]

    result = {}
    for row in data_rows:
        code, pref, name, male, female, total, households = row[:7]
        if pref != PREF_NAME:
            continue
        if not is_real_municipality(code, name):
            continue
        code = str(code)
        if code not in (CITIES if codes is None else codes):
            continue  # 自治体マスタの団体だけ（政令市の区の行などを除く）。岐阜では従来どおり42件すべてが通る
        result[code] = {
            "市区町村コード": code,
            "市区町村名": name,
            "男": int(male) if male is not None else None,
            "女": int(female) if female is not None else None,
            "計": int(total) if total is not None else None,
            "世帯数": int(households) if households is not None else None,
        }
    return result


def main():
    per_year = {}
    counts = {}
    for fmeta in FILES:
        key = fmeta["key"]
        gifu = extract_gifu(fmeta)
        per_year[key] = gifu
        counts[key] = len(gifu)

    with open(f"out_extract_summary{SUFFIX}.json", "w", encoding="utf-8") as f:
        json.dump(counts, f, ensure_ascii=False, indent=1)

    # merge by code
    all_codes = set()
    for gifu in per_year.values():
        all_codes.update(gifu.keys())

    merged = []
    for code in sorted(all_codes):
        rec = {"市区町村コード": code}
        name_set = set()
        for fmeta in FILES:
            key = fmeta["key"]
            entry = per_year[key].get(code)
            if entry:
                name_set.add(entry["市区町村名"])
                rec[key] = {
                    "人口_男": entry["男"],
                    "人口_女": entry["女"],
                    "人口_計": entry["計"],
                    "世帯数": entry["世帯数"],
                }
            else:
                rec[key] = None
        rec["市区町村名"] = sorted(name_set)
        merged.append(rec)

    with open(f"out_merged{SUFFIX}.json", "w", encoding="utf-8") as f:
        json.dump(merged, f, ensure_ascii=False, indent=1)

    # 政令市の区（試作のみ）：市の中間出力 out_merged とは別ファイル out_merged_{県}_wards.json に書く
    wards = load_master_wards(PREF)
    if wards:
        w_year = {fm["key"]: extract_gifu(fm, set(wards)) for fm in FILES}
        w_merged = []
        for code in sorted(wards):
            rec = {"市区町村コード": code, "市区町村名": [wards[code][0]], "parent_area_code": wards[code][1]}
            for fm in FILES:
                e = w_year[fm["key"]].get(code)
                rec[fm["key"]] = None if e is None else {"人口_男": e["男"], "人口_女": e["女"], "人口_計": e["計"], "世帯数": e["世帯数"]}
            w_merged.append(rec)
        with open(f"out_merged{SUFFIX}_wards.json", "w", encoding="utf-8") as f:
            json.dump(w_merged, f, ensure_ascii=False, indent=1)
        print("wards:", {k: len(v) for k, v in w_year.items()})

    print("counts:", counts)
    print("total distinct codes:", len(all_codes))
    if set(all_codes) != set(CITIES):
        raise SystemExit(f"自治体マスタ{len(CITIES)}件と抽出{len(all_codes)}件が一致しない: "
                         f"マスタのみ={sorted(set(CITIES) - all_codes)[:8]} 抽出のみ={sorted(all_codes - set(CITIES))[:8]}")


if __name__ == "__main__":
    main()
