# -*- coding: utf-8 -*-
"""まちのすがた「単身世帯割合」「高齢化率」市区町村版JSONに全国順位・類似団体比較を追加する。

対象（docs/content-spec-v1.md §7.2 / docs/similar-group-aggregation-design.md）:
  1-1b 単身世帯割合・高齢化率（岐阜42市区町村・比較値=全国平均・類似団体平均・全国順位・百分位つき）
  content-spec-v1.md §7.2 の定義：1-1bの主比較は「比較=類団」。2026-09-27運営者裁定（1-1a人口・
  世帯数）と同じ「全国順位＋類似団体平均」の形で、既存実装（1-2b/1-2c）と統一する。2026-09-28、
  この形をF-5残り項目（1-1b・2a）にも適用することで運営者裁定確定。

入力（取得済みローカルファイルのみ。ネットワークアクセスをしない）:
  poc/recon-20261012-kokusei2025/estat_000040506456_kokusei2025_family_type_shichoson.xlsx
    （統計表ID 000040506456・令和7年国勢調査 第９－１表。既存 data/m1_1b_single_household.json と同一の全国集約
    ファイル。5桁地域コード＝地域名セル "21201_岐阜市" 形式。単独世帯数(3_単独世帯)／
    一般世帯数(0_総数)の2行を抽出）
  poc/recon-20261012-juki-r8/soumu_000892953_juki_shichoson_age_r8.xlsx
    （総務省ファイルID 000892953・令和8年1月1日時点。既存 data/m1_1c_aging_rate.json と同一の全国集約
    ファイル。6桁団体コード直接（住基）。総数・65歳以上（8つの年齢帯の合計・性別「計」行）を抽出。列はラベルから同定）
  poc/recon-20260828/0/ruikei_kubun_r06.xlsx（類似団体区分マスタ・市町村一覧シート・1741団体）

更新（既存ファイルを書き換え・新規ファイルは作らない）:
  data/m1_1b_single_household.json の各レコード（42件）に comparison フィールドを追加する
  （対象値＝単身世帯割合）。
  data/m1_1c_aging_rate.json の各レコード（42件）に comparison フィールドを追加する
  （対象値＝高齢化率）。

集計（docs/similar-group-aggregation-design.md §3・§4・§9・§10。m1_2b_daynight_comparison.py と
  同一の実装パターン）:
  母集団M = 区分マスタに載る団体（1741）∩ 全国集約ファイルに値がある団体。
  束ね写像：政令指定都市・中核市・施行時特例市・特別区 → "大都市"。他は類団区分の生ラベルをそのまま使う。
  順位：降順 competition ranking（同値は同順位、次の順位を飛ばす）。値が大きいほど上位＝優劣評価ではない。
  percentile = (rank_out_of − rank + 1) / rank_out_of × 100（小数第1位）。
  average_method = "unweighted_mean"（単純算術平均。人口加重等はしない）。

規律:
  - 列位置はハードコードせず、ヘッダのラベル文字列から同定し自己検証する。
  - 表題をコードでassertしてから使う（推測しない）。
  - 評価語・評価色・複合スコアは付けない（中立性・D-4/D-6/D-18）。
"""
from __future__ import annotations

import json
import re
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

FAMILY_SRC = ROOT / "poc" / "recon-20261012-kokusei2025" / "estat_000040506456_kokusei2025_family_type_shichoson.xlsx"
AGE_SRC = ROOT / "poc" / "recon-20261012-juki-r8" / "soumu_000892953_juki_shichoson_age_r8.xlsx"
RUIKEI_SRC = ROOT / "poc" / "recon-20260828" / "0" / "ruikei_kubun_r06.xlsx"
OUT_SINGLE = ROOT / "data" / "m1_1b_single_household.json"
OUT_AGING = ROOT / "data" / "m1_1c_aging_rate.json"

UPDATED = date.today().isoformat()

DAITOSHI_KUBUN = {1, 2, 8, 9}

ANCHOR = {
    "gifu_shi_group": "大都市",
    "gifu_shi_group_n": 128,
    "ogaki_shi_group": "Ⅳ－２",
    "ogaki_shi_group_n": 75,
    "takayama_shi_group": "Ⅱ－１",
    "takayama_shi_group_n": 127,
}

# 2026-09-28コード直接確認によるアンカー（岐阜市・自己算出値。既存JSONの値と一致することを
# 別途「既存レコードとの整合確認」で検証するため、ここでは構造検証のみに用いる）。
# 単身世帯割合は令和7年国勢調査へ更新。令和2年版の固定アンカー（36.47）は使わず、既存JSON（同一原本から生成）との
# 一致確認（apply_to_json）と独立再計算で検証する。
# 高齢化率は令和8年版へ更新。令和7年版の固定アンカー（29.32）は使わず、既存JSON（同一原本から生成）との
# 一致確認（apply_to_json）と、下の独立再計算で検証する。


class Abort(Exception):
    """欠損・単位不一致・アンカー不一致などで処理を止めるための例外。推測で続行しない。"""


def s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def cell(row: tuple, c: int):
    if row is None or c < 0 or c >= len(row):
        return None
    return row[c]


# --------------------------------------------------------------------------------------
# 区分マスタ（類似団体区分）※既存スクリプトと同一実装
# --------------------------------------------------------------------------------------
def load_ruikei_master() -> dict[str, dict]:
    from openpyxl import load_workbook

    wb = load_workbook(RUIKEI_SRC, read_only=True, data_only=True)
    sheet_name = next((n for n in wb.sheetnames if "市町村一覧" in n), None)
    if sheet_name is None:
        raise Abort(f"ruikei: シート『市町村一覧』が見つからない（実在={wb.sheetnames}）")
    ws = wb[sheet_name]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    print(f"  [ruikei] engine=openpyxl sheet={sheet_name!r}")
    hidx = None
    for i, row in enumerate(rows):
        if s(row[0] if row else None) == "団体コード":
            hidx = i
            break
    if hidx is None:
        raise Abort("ruikei: ヘッダ行（1列目='団体コード'）が見つからない")

    out: dict[str, dict] = {}
    for row in rows[hidx + 1 :]:
        code = s(row[0]) if len(row) > 0 else ""
        if not re.fullmatch(r"\d{6}", code):
            continue
        name = s(row[1])
        danntai_kubun_raw = row[2]
        ruikei_raw = s(row[4])
        if not ruikei_raw:
            continue
        try:
            danntai_kubun = int(danntai_kubun_raw)
        except (TypeError, ValueError):
            raise Abort(f"ruikei: {code} の団体区分が数値でない: {danntai_kubun_raw!r}")
        group_code = "大都市" if danntai_kubun in DAITOSHI_KUBUN else ruikei_raw
        out[code] = {"name": name, "danntai_kubun": danntai_kubun, "ruikei_raw": ruikei_raw, "group_code": group_code}
    if len(out) != 1741:
        raise Abort(f"ruikei: 市町村一覧の件数が1741ではない（実測={len(out)}）")
    print(f"  [ruikei] 市町村一覧 {len(out)}件を読み込み")
    return out


# --------------------------------------------------------------------------------------
# 単身世帯割合（国勢調査 第９－１表・世帯の家族類型）
# --------------------------------------------------------------------------------------
def load_single_household_ratio_by_code5() -> dict[str, float]:
    from openpyxl import load_workbook

    wb = load_workbook(FAMILY_SRC, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header_rows = list(ws.iter_rows(min_row=1, max_row=4, values_only=True))
    title_texts = " ".join(s(v) for row in header_rows[:2] for v in row if v)
    if "第９－１表" not in title_texts or "世帯の家族類型" not in title_texts or "令和７年" not in title_texts:
        raise Abort(f"family: 表題に想定語を確認できない（実測={title_texts!r}）。推測で続行しない。")
    print(f"  [family] 表題確認: {title_texts[:60]}...")
    # 列位置（地域名=3列目・家族類型=5列目・一般世帯数の総数=6列目）をヘッダ行で自己検証する
    hdr = list(ws.iter_rows(min_row=1, max_row=40, values_only=True))
    hrow = next((r for r in hdr if s(cell(r, 0)) == "地域識別コード"), None)
    if hrow is None or s(cell(hrow, 2)) != "地域名" or s(cell(hrow, 4)) != "世帯の家族類型":
        raise Abort("family: ヘッダ行の列（地域名=3列目・世帯の家族類型=5列目）を確認できない")
    grp = [r for r in hdr if s(cell(r, 5)) == "一般世帯数"]
    sub = [r for r in hdr if s(cell(r, 5)) == "0_総数"]
    if not grp or not sub:
        raise Abort("family: 6列目が『一般世帯数／0_総数』であることを確認できない")

    general: dict[str, float] = {}
    single: dict[str, float] = {}
    n_scanned = 0
    for row in ws.iter_rows(values_only=True):
        area_v = s(cell(row, 2))
        m = re.match(r"^(\d{5})_", area_v)
        if not m:
            continue
        code5 = m.group(1)
        if code5 == "00000" or code5[-3:] == "000":
            continue  # 全国・都道府県の集計行を除く
        family_type = s(cell(row, 4))
        value = cell(row, 5)
        if not isinstance(value, (int, float)) or isinstance(value, bool):
            continue
        if family_type == "0_総数":
            if code5 in general:
                raise Abort(f"family: コード{code5}の一般世帯数(0_総数)が重複している")
            general[code5] = float(value)
            n_scanned += 1
        elif family_type == "3_単独世帯":
            if code5 in single:
                raise Abort(f"family: コード{code5}の単独世帯数(3_単独世帯)が重複している")
            single[code5] = float(value)
    print(f"  [family] 一般世帯数{len(general)}件・単独世帯数{len(single)}件抽出")

    out: dict[str, float] = {}
    for code5, gen_v in general.items():
        sng_v = single.get(code5)
        if sng_v is None or gen_v <= 0:
            continue
        out[code5] = round(sng_v / gen_v * 100, 2)
    return out


# --------------------------------------------------------------------------------------
# 高齢化率（住基「年齢階級別人口」）
# --------------------------------------------------------------------------------------
def resolve_age_columns(ws):
    """(団体コード列, 性別列, 総数列, 65歳以上列リスト, データ開始行) を返す。列位置は自己同定する。"""
    header_rows = list(ws.iter_rows(min_row=1, max_row=4, values_only=True))

    code_row_idx = None
    for i, row in enumerate(header_rows):
        if s(cell(row, 0)) == "団体コード" and s(cell(row, 1)) == "都道府県名" and s(cell(row, 2)) == "市区町村名":
            code_row_idx = i
            break
    if code_row_idx is None:
        raise Abort("age: ヘッダ行（団体コード/都道府県名/市区町村名）が見つからない")
    label_row = header_rows[code_row_idx - 1]  # 総数/0歳～4歳/.../100歳以上

    def col_of(label: str):
        for c, v in enumerate(label_row):
            if s(v) == label:
                return c
        return None

    sex_col = 3  # 団体コード(0)/都道府県名(1)/市区町村名(2)/性別(3) を想定・下で自己検証
    if s(cell(header_rows[code_row_idx], sex_col)) != "性別":
        raise Abort("age: 4列目が『性別』ではない（想定した表構造と異なる）")

    total_col = col_of("総数")
    over65_cols = [col_of(f"{a}歳～{a + 4}歳") for a in range(65, 95, 5)]
    over65_cols += [col_of("95歳～99歳"), col_of("100歳以上")]
    if total_col is None or None in over65_cols:
        raise Abort(f"age: 総数列または65歳以上の年齢帯列が見つからない total={total_col} over65={over65_cols}")

    title_texts = " ".join(s(v) for row in header_rows[:1] for v in row if v)
    if "住民基本台帳" not in title_texts or "年齢階級別人口" not in title_texts or "令和8年" not in title_texts:
        raise Abort(f"age: 表題に想定語を確認できない（実測={title_texts!r}）。推測で続行しない。")
    print(f"  [age] 表題確認: {title_texts[:60]}...")

    data_start = code_row_idx + 1
    return 0, sex_col, total_col, over65_cols, data_start


def load_aging_rate_by_code6() -> dict[str, float]:
    from openpyxl import load_workbook

    wb = load_workbook(AGE_SRC, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    code_col, sex_col, total_col, over65_cols, data_start = resolve_age_columns(ws)
    print(
        f"  [age] 列: 団体コード列=c{code_col + 1} 性別列=c{sex_col + 1} 総数列=c{total_col + 1} "
        f"65歳以上列={[c + 1 for c in over65_cols]} データ開始行=r{data_start + 1}"
    )

    out: dict[str, float] = {}
    n_scanned = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < data_start:
            continue
        code = s(cell(row, code_col))
        if not re.fullmatch(r"\d{6}", code):
            continue
        if s(cell(row, sex_col)) != "計":
            continue  # 男女計のみ対象（男/女の内訳行は除く）
        n_scanned += 1
        total_v = cell(row, total_col)
        over65_vals = [cell(row, c) for c in over65_cols]
        if not isinstance(total_v, (int, float)) or isinstance(total_v, bool):
            continue
        if any(not isinstance(v, (int, float)) or isinstance(v, bool) for v in over65_vals):
            continue
        if total_v <= 0:
            continue
        if code in out:
            raise Abort(f"age: コード{code}が重複している")
        over65_total = sum(over65_vals)
        out[code] = round(over65_total / total_v * 100, 2)
    print(f"  [age] 市区町村{n_scanned}行走査・有効値{len(out)}件抽出")
    return out


# --------------------------------------------------------------------------------------
# 集計（共通）
# --------------------------------------------------------------------------------------
def ranks_desc(values: list[float]) -> list[int]:
    return [1 + sum(1 for w in values if w > v) for v in values]


def build_comparison(ruikei: dict, value_by_code6: dict[str, float]) -> tuple[dict[str, dict], dict[str, int]]:
    population: dict[str, dict] = {}
    for code6, meta in ruikei.items():
        v = value_by_code6.get(code6)
        if v is None:
            continue
        population[code6] = {"value": v, "group_code": meta["group_code"]}

    codes = list(population.keys())
    values = [population[c]["value"] for c in codes]
    n_total = len(codes)
    national_average = round(sum(values) / n_total, 1)

    group_values: dict[str, list[float]] = {}
    for c in codes:
        group_values.setdefault(population[c]["group_code"], []).append(population[c]["value"])
    group_average = {g: round(sum(vs) / len(vs), 1) for g, vs in group_values.items()}
    group_n = {g: len(vs) for g, vs in group_values.items()}

    ranks = dict(zip(codes, ranks_desc(values)))

    out: dict[str, dict] = {}
    for c in codes:
        p = population[c]
        rank = ranks[c]
        pct = round((n_total - rank + 1) / n_total * 100, 1)
        out[c] = {
            "value": p["value"],
            "national_average": national_average,
            "group_code": p["group_code"],
            "group_average": group_average[p["group_code"]],
            "rank": rank,
            "rank_out_of": n_total,
            "percentile": pct,
            "average_method": "unweighted_mean",
            "source_note": (
                f"全国{n_total}市区町村中の順位（母集団=類似団体区分マスタ登載団体1741のうち、"
                "本統計に数値がある団体）。類似団体平均は総務省の類似団体区分（束ね後）の単純算術平均。"
                "順位は値の大きい順であり、大小の望ましさを含意しない。"
            ),
        }
    return out, group_n


def check_anchor(label, got, want):
    ok = got == want
    print(f"    {'OK ' if ok else 'NG '} {label}: 実測={got!r} 期待={want!r}")
    return ok


def apply_to_json(out_path: Path, comparison_by_code6: dict[str, dict], value_field: str, metric_label: str, statinfid: str):
    payload = json.loads(out_path.read_text(encoding="utf-8"))
    existing = payload["data"]
    if len(existing) != 42:
        raise Abort(f"{out_path.name}: 既存JSONのレコード数が42ではない（実測={len(existing)}）")

    gifu_codes = sorted(c for c in comparison_by_code6 if c.startswith("21"))
    if len(gifu_codes) != 42:
        raise Abort(f"{out_path.name}: 岐阜県内市区町村が母集団Mで{len(gifu_codes)}件（42件ではない）")

    by_area_code = {r["area_code"]: r for r in existing}
    if set(by_area_code.keys()) != set(gifu_codes):
        raise Abort(f"{out_path.name}: 既存JSONのarea_codeと母集団Mの岐阜42件のキー集合が一致しない。")

    for rec in existing:
        code6 = rec["area_code"]
        c = comparison_by_code6[code6]
        if abs(rec[value_field] - c["value"]) > 0.011:  # 四捨五入誤差許容（小数第2位）
            raise Abort(
                f"{out_path.name}: {rec['市区町村名']}（{code6}）: 既存値{rec[value_field]}と"
                f"全国ファイルからの再算出値{c['value']}が不一致（差={abs(rec[value_field] - c['value'])}）。"
            )
        rec["comparison"] = c

    payload["updated"] = UPDATED
    payload.setdefault("notes", [])
    prefix = "comparison（全国平均・類似団体平均・全国順位・百分位）は"
    payload["notes"] = [n for n in payload["notes"] if not n.startswith(prefix)]
    payload["notes"].append(
        f"comparison（全国平均・類似団体平均・全国順位・百分位）は2026-09-28追加（F-5）。"
        "docs/similar-group-aggregation-design.md §3・§4の集計段の定義に従う。"
        f"母集団M = 区分マスタ（総務省 類似団体別市町村財政指数表・1741団体）"
        f"∩ 全国集約ファイル（{statinfid}）に数値がある団体。"
        "national_average は母集団Mの単純算術平均、group_average は束ね後区分ごとの単純算術平均"
        "（average_method=\"unweighted_mean\"）。束ね写像：政令指定都市・中核市・施行時特例市・特別区は"
        "『大都市』に統合。それ以外は類団区分の原ラベルをそのまま使う。"
        f"content-spec-v1.md §7.2の定義上の主比較は「比較=類団」（{metric_label}）。"
        "2026-09-27・2026-09-28運営者裁定により全国順位＋類似団体平均の形で提示する。"
        "rank は母集団M内での値の降順順位（1=最大値。同値は同順位・次の順位を飛ばすcompetition ranking。"
        "値が大きいほど上位＝優劣評価ではない）。"
        "percentile = (rank_out_of − rank + 1) / rank_out_of × 100（小数第1位）。"
        "評価語・評価色・複合スコアは付与していない（20260808_decisions.md D-4／D-6／D-18）。"
    )

    with out_path.open("w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)
    print(f"  wrote {out_path.relative_to(ROOT)}  ({len(existing)}件にcomparison付与)")


def main():
    for p in (FAMILY_SRC, AGE_SRC, RUIKEI_SRC, OUT_SINGLE, OUT_AGING):
        if not p.exists():
            raise Abort(f"入力ファイルが無い: {p}")

    print("== 区分マスタ読み込み ==")
    ruikei = load_ruikei_master()

    print("== 単身世帯割合（全国・国勢調査）読み込み ==")
    single_ratio_by_code5 = load_single_household_ratio_by_code5()
    single_ratio_by_code6 = {code6: single_ratio_by_code5[code6[:5]] for code6 in ruikei if code6[:5] in single_ratio_by_code5}

    print("== 高齢化率（全国・住基）読み込み ==")
    aging_by_code6 = load_aging_rate_by_code6()

    print("== 区分アンカー自己検証 ==")
    gifu_shi_meta = ruikei.get("212016")
    ogaki_shi_meta = ruikei.get("212024")
    takayama_shi_meta = ruikei.get("212032")
    if not (gifu_shi_meta and ogaki_shi_meta and takayama_shi_meta):
        raise Abort("岐阜市/大垣市/高山市のいずれかが区分マスタに無い")

    print("== 単身世帯割合の集計・書き出し ==")
    comp_single, group_n_s = build_comparison(ruikei, single_ratio_by_code6)
    results_s = [
        check_anchor("岐阜市 束ね後区分", gifu_shi_meta["group_code"], ANCHOR["gifu_shi_group"]),
        check_anchor("岐阜市 束ね後区分の団体数", group_n_s[gifu_shi_meta["group_code"]], ANCHOR["gifu_shi_group_n"]),
        check_anchor("大垣市 束ね後区分", ogaki_shi_meta["group_code"], ANCHOR["ogaki_shi_group"]),
        check_anchor("大垣市 束ね後区分の団体数", group_n_s[ogaki_shi_meta["group_code"]], ANCHOR["ogaki_shi_group_n"]),
        check_anchor("高山市 束ね後区分", takayama_shi_meta["group_code"], ANCHOR["takayama_shi_group"]),
        check_anchor("高山市 束ね後区分の団体数", group_n_s[takayama_shi_meta["group_code"]], ANCHOR["takayama_shi_group_n"]),
        check_anchor("岐阜市 単身世帯割合 母集団Mの件数", comp_single["212016"]["rank_out_of"], len(single_ratio_by_code6)),
        check_anchor(
            "岐阜市 単身世帯割合 順位（独立再計算）",
            comp_single["212016"]["rank"],
            1 + sum(1 for v in single_ratio_by_code6.values() if v > single_ratio_by_code6["212016"]),
        ),
    ]
    if not all(results_s):
        raise Abort("単身世帯割合: アンカー値が一致しない。JSONを書かずに中断した。")
    apply_to_json(OUT_SINGLE, comp_single, "単身世帯割合", "単身世帯割合", "統計表ID000040506456（令和7年国勢調査）")

    print("== 高齢化率の集計・書き出し ==")
    comp_aging, group_n_a = build_comparison(ruikei, aging_by_code6)
    results_a = [
        check_anchor("岐阜市 束ね後区分", gifu_shi_meta["group_code"], ANCHOR["gifu_shi_group"]),
        check_anchor("岐阜市 束ね後区分の団体数", group_n_a[gifu_shi_meta["group_code"]], ANCHOR["gifu_shi_group_n"]),
        check_anchor("岐阜市 高齢化率 母集団Mの件数", comp_aging["212016"]["rank_out_of"], len([c for c in ruikei if c in aging_by_code6])),
        check_anchor(
            "岐阜市 高齢化率 順位（独立再計算）",
            comp_aging["212016"]["rank"],
            1 + sum(1 for c in ruikei if c in aging_by_code6 and aging_by_code6[c] > aging_by_code6["212016"]),
        ),
    ]
    if not all(results_a):
        raise Abort("高齢化率: アンカー値が一致しない。JSONを書かずに中断した。")
    apply_to_json(OUT_AGING, comp_aging, "高齢化率", "高齢化率", "総務省ファイルID000892953（令和8年1月1日）")

    print("== 岐阜42件 単身世帯割合順（サニティチェック用） ==")
    d = json.loads(OUT_SINGLE.read_text(encoding="utf-8"))
    ranked = sorted(d["data"], key=lambda r: -r["comparison"]["value"])
    for r in ranked[:3] + ranked[-3:]:
        c = r["comparison"]
        print(
            f"  {r['市区町村名']}: {c['value']}% 全国{c['rank']}位/{c['rank_out_of']} "
            f"区分={c['group_code']} 区分平均={c['group_average']}"
        )


if __name__ == "__main__":
    try:
        main()
    except Abort as exc:
        print(f"\n!! 中断: {exc}", file=sys.stderr)
        sys.exit(1)
