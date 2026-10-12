# -*- coding: utf-8 -*-
"""Extract 2-2a〜2-2g for Gifu 42 into data/m2_2_finance.json. Keys from SCHEMA_*."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RECON = ROOT / "poc" / "recon-20260830"
RAW = RECON / "raw"
INSPECT = RECON / "_inspect"
OUT_JSON = ROOT / "data" / "m2_2_finance.json"
VERIFY_PATH = INSPECT / "verify.json"

GIFU_42 = [
    ("212016", "岐阜市"),
    ("212024", "大垣市"),
    ("212032", "高山市"),
    ("212041", "多治見市"),
    ("212059", "関市"),
    ("212067", "中津川市"),
    ("212075", "美濃市"),
    ("212083", "瑞浪市"),
    ("212091", "羽島市"),
    ("212105", "恵那市"),
    ("212113", "美濃加茂市"),
    ("212121", "土岐市"),
    ("212130", "各務原市"),
    ("212148", "可児市"),
    ("212156", "山県市"),
    ("212164", "瑞穂市"),
    ("212172", "飛騨市"),
    ("212181", "本巣市"),
    ("212199", "郡上市"),
    ("212202", "下呂市"),
    ("212211", "海津市"),
    ("213021", "岐南町"),
    ("213039", "笠松町"),
    ("213411", "養老町"),
    ("213616", "垂井町"),
    ("213624", "関ケ原町"),
    ("213811", "神戸町"),
    ("213829", "輪之内町"),
    ("213837", "安八町"),
    ("214019", "揖斐川町"),
    ("214035", "大野町"),
    ("214043", "池田町"),
    ("214213", "北方町"),
    ("215015", "坂祝町"),
    ("215023", "富加町"),
    ("215031", "川辺町"),
    ("215040", "七宗町"),
    ("215058", "八百津町"),
    ("215066", "白川町"),
    ("215074", "東白川村"),
    ("215210", "御嵩町"),
    ("216046", "白川村"),
]
GIFU_BY_CODE = {c: n for c, n in GIFU_42}
GIFU_BY_NAME = {n: c for c, n in GIFU_42}

REVENUE_KAN_KEYS_R6 = [
    "地方税",
    "地方譲与税",
    "利子割交付金",
    "配当割交付金",
    "株式等譲渡所得割交付金",
    "分離課税所得割交付金",
    "地方消費税交付金",
    "ゴルフ場利用税交付金",
    "自動車取得税交付金",
    "軽油引取税交付金",
    "自動車税環境性能割交付金",
    "法人事業税交付金",
    "地方特例交付金等",
    "地方交付税",
    "交通安全対策特別交付金",
    "分担金及び負担金",
    "使用料",
    "手数料",
    "国庫支出金",
    "国有提供施設等所在市町村助成交付金",
    "都道府県支出金",
    "財産収入",
    "寄附金",
    "繰入金",
    "繰越金",
    "諸収入",
    "地方債",
    "特別区財政調整交付金",
]
REVENUE_KAN_KEYS_R5 = [
    "地方税",
    "地方譲与税",
    "利子割交付金",
    "配当割交付金",
    "株式等譲渡所得割交付金",
    "分離課税所得割交付金",
    "地方消費税交付金",
    "ゴルフ場利用税交付金",
    "特別地方消費税交付金",
    "自動車取得税交付金",
    "軽油引取税交付金",
    "自動車税環境性能割交付金",
    "法人事業税交付金",
    "地方特例交付金等",
    "地方交付税",
    "交通安全対策特別交付金",
    "分担金及び負担金",
    "使用料",
    "手数料",
    "国庫支出金",
    "国有提供施設等所在市町村助成交付金",
    "都道府県支出金",
    "財産収入",
    "寄附金",
    "繰入金",
    "繰越金",
    "諸収入",
    "地方債",
    "特別区財政調整交付金",
]
EXPENDITURE_PURPOSE_KEYS = [
    "議会費",
    "総務費",
    "民生費",
    "衛生費",
    "労働費",
    "農林水産業費",
    "商工費",
    "土木費",
    "消防費",
    "教育費",
    "災害復旧費",
    "公債費",
    "諸支出金",
    "前年度繰上充用金",
]
RECORD_KEYS = [
    "area_code",
    "市区町村名",
    "revenue_r6",
    "expenditure_purpose_r6",
    "revenue_r5",
    "expenditure_purpose_r5",
    "実質公債費比率",
    "将来負担比率",
    "実質赤字比率",
    "連結実質赤字比率",
    "実質赤字比率_早期健全化基準",
    "連結実質赤字比率_早期健全化基準",
    "ratio_source_tokens",
    "excel_rows",
]
TOP_KEYS = [
    "year",
    "prior_year",
    "updated",
    "unit_revenue_expenditure",
    "unit_ratio",
    "code_format",
    "revenue_kan_keys_r6",
    "revenue_kan_keys_r5",
    "expenditure_purpose_keys",
    "record_keys",
    "statutory_criteria_common",
    "prefecture_layer",
    "sources",
    "verification",
    "notes",
    "data",
]

DASH = {"-", "－", "―", "─"}


def norm(v) -> str:
    if v is None:
        return ""
    return str(v).replace("\n", "").replace(" ", "").replace("　", "")


def to_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if s in DASH or s == "":
        return None
    try:
        n = float(s)
        return int(n) if n.is_integer() else n
    except ValueError:
        return None


def token(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return None
    s = str(v).strip()
    return s if s in DASH else None


def header_map(ws, row_idx: int) -> dict[str, int]:
    row = next(ws.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True))
    out: dict[str, int] = {}
    for i, v in enumerate(row):
        key = norm(v)
        if key and key not in out:
            out[key] = i
    return out


def load_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


def index_muni(rows, header_row: int, code_col: int = 14, name_col: int = 15):
    by_code = {}
    totals = {}
    pref_headers = []
    for i, row in enumerate(rows, start=1):
        if i <= header_row:
            continue
        code = row[code_col] if len(row) > code_col else None
        name = row[name_col] if len(row) > name_col else None
        code_s = "" if code is None else str(code).strip()
        name_s = "" if name is None else str(name).strip().replace("　", "")
        if code_s.isdigit() and len(code_s) == 6:
            by_code[code_s] = (i, row)
        elif name_s in ("全国計", "市計", "町村計", "合計"):
            totals[name_s] = (i, row)
        elif name_s and not code_s:
            pref_headers.append((i, name_s))
    return by_code, totals, pref_headers


def extract_named(row, colmap: dict[str, int], keys: list[str]) -> dict:
    out = {}
    for k in keys:
        idx = colmap.get(norm(k))
        if idx is None:
            continue
        if idx >= len(row):
            continue
        out[k] = to_num(row[idx])
    return out


def extract_exp(row, colmap: dict[str, int]) -> dict:
    out = extract_named(
        row,
        colmap,
        [k for k in EXPENDITURE_PURPOSE_KEYS if k not in ("労働費", "諸支出金")],
    )
    # 労働費・諸支出金は款合計列がヘッダに無く、項列の合計。
    labor = 0
    labor_ok = False
    for idx in (35, 36):
        if idx < len(row):
            n = to_num(row[idx])
            if n is not None:
                labor += n
                labor_ok = True
    out["労働費"] = labor if labor_ok else None
    shos = 0
    shos_ok = False
    for idx in (91, 92, 93):
        if idx < len(row):
            n = to_num(row[idx])
            if n is not None:
                shos += n
                shos_ok = True
    out["諸支出金"] = shos if shos_ok else None
    return {k: out.get(k) for k in EXPENDITURE_PURPOSE_KEYS}


def sum_vals(d: dict) -> float:
    return float(sum(v for v in d.values() if isinstance(v, (int, float))))


def main() -> None:
    files = {
        "gaikyo_shi": RAW / "001061669.xlsx",
        "gaikyo_cho": RAW / "001061674.xlsx",
        "sainyu_shi_r6": RAW / "001061670.xlsx",
        "sainyu_cho_r6": RAW / "001061675.xlsx",
        "mokuteki_shi_r6": RAW / "001061671.xlsx",
        "mokuteki_cho_r6": RAW / "001061676.xlsx",
        "sainyu_shi_r5": RAW / "000999901.xlsx",
        "sainyu_cho_r5": RAW / "000999906.xlsx",
        "mokuteki_shi_r5": RAW / "000999902.xlsx",
        "mokuteki_cho_r5": RAW / "000999908.xlsx",
        "kenzen": RAW / "001060123.xlsx",
    }
    loaded = {k: load_rows(p) for k, p in files.items() if k != "kenzen"}

    maps = {}
    indexed = {}
    for key, rows in loaded.items():
        header_r = 10 if "gaikyo" in key else 10
        maps[key] = header_map_from_rows(rows, header_r)
        indexed[key] = index_muni(rows, header_r)

    kenzen_by_name = {}
    kenzen_pref = None
    wb = load_workbook(files["kenzen"], read_only=True, data_only=True)
    ws = wb["健全化（市区町村）"]
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if len(row) < 11:
            continue
        if row[2] == "岐阜県" and row[3] in GIFU_BY_NAME:
            kenzen_by_name[str(row[3])] = (i, row)
    ws2 = wb["健全化（都道府県）"]
    for i, row in enumerate(ws2.iter_rows(values_only=True), start=1):
        if row and len(row) > 6 and row[2] == "岐阜県":
            kenzen_pref = {"row": i, "実質公債費比率": to_num(row[6]), "将来負担比率": to_num(row[7]),
                           "実質赤字比率": to_num(row[4]), "連結実質赤字比率": to_num(row[5])}
    wb.close()

    data = []
    verify_rows = []
    missing = {k: [] for k in ("gaikyo", "sainyu_r6", "mokuteki_r6", "sainyu_r5", "mokuteki_r5", "kenzen")}

    for code, name in GIFU_42:
        layer = "shi" if name.endswith("市") else "cho"
        gaikyo_key = f"gaikyo_{layer}"
        rec_src = {}
        excel_rows = {}

        g_idx, g_row = indexed[gaikyo_key][0].get(code, (None, None))
        if g_row is None:
            missing["gaikyo"].append(code)
            continue
        excel_rows["gaikyo_r6"] = g_idx
        ratios = {
            "実質赤字比率": to_num(g_row[35]),
            "連結実質赤字比率": to_num(g_row[36]),
            "実質公債費比率": to_num(g_row[37]),
            "将来負担比率": to_num(g_row[38]),
        }
        tokens = {
            "実質赤字比率": token(g_row[35]) or ("-" if ratios["実質赤字比率"] is None else None),
            "連結実質赤字比率": token(g_row[36]) or ("-" if ratios["連結実質赤字比率"] is None else None),
            "実質公債費比率": token(g_row[37]),
            "将来負担比率": token(g_row[38]) or ("-" if ratios["将来負担比率"] is None else None),
        }
        gaikyo_in = to_num(g_row[39])
        gaikyo_out = to_num(g_row[40])

        def take(filekey, keys, kind):
            idx, row = indexed[filekey][0].get(code, (None, None))
            if row is None:
                missing[kind].append(code)
                return None, None
            excel_rows[kind] = idx
            if kind.startswith("sainyu"):
                return extract_named(row, maps[filekey], keys), idx
            return extract_exp(row, maps[filekey]), idx

        rev6, _ = take(f"sainyu_{layer}_r6", REVENUE_KAN_KEYS_R6, "sainyu_r6")
        exp6, _ = take(f"mokuteki_{layer}_r6", EXPENDITURE_PURPOSE_KEYS, "mokuteki_r6")
        rev5, _ = take(f"sainyu_{layer}_r5", REVENUE_KAN_KEYS_R5, "sainyu_r5")
        exp5, _ = take(f"mokuteki_{layer}_r5", EXPENDITURE_PURPOSE_KEYS, "mokuteki_r5")

        krow = kenzen_by_name.get(name)
        early_def = early_ren = None
        if krow:
            excel_rows["kenzen_r6"] = krow[0]
            kr = krow[1]
            early_def = to_num(kr[6])
            early_ren = to_num(kr[8])
            rec_src["kenzen_match"] = {
                "実質公債費比率": to_num(kr[9]),
                "将来負担比率": to_num(kr[10]),
                "実質赤字比率": to_num(kr[5]),
                "連結実質赤字比率": to_num(kr[7]),
            }
        else:
            missing["kenzen"].append(code)

        record = {
            "area_code": code,
            "市区町村名": name,
            "revenue_r6": rev6 or {},
            "expenditure_purpose_r6": exp6 or {},
            "revenue_r5": rev5 or {},
            "expenditure_purpose_r5": exp5 or {},
            "実質公債費比率": ratios["実質公債費比率"],
            "将来負担比率": ratios["将来負担比率"],
            "実質赤字比率": ratios["実質赤字比率"],
            "連結実質赤字比率": ratios["連結実質赤字比率"],
            "実質赤字比率_早期健全化基準": early_def,
            "連結実質赤字比率_早期健全化基準": early_ren,
            "ratio_source_tokens": {k: tokens[k] for k in ("実質赤字比率", "連結実質赤字比率", "実質公債費比率", "将来負担比率")},
            "excel_rows": excel_rows,
        }
        data.append({k: record[k] for k in RECORD_KEYS})

        rev6_sum = sum_vals(rev6 or {})
        exp6_sum = sum_vals(exp6 or {})
        verify_rows.append({
            "area_code": code,
            "市区町村名": name,
            "gaikyo_row": g_idx,
            "歳入総額_概況": gaikyo_in,
            "歳出総額_概況": gaikyo_out,
            "revenue_r6_sum": rev6_sum,
            "expenditure_r6_sum": exp6_sum,
            "revenue_minus_gaikyo": None if gaikyo_in is None else rev6_sum - gaikyo_in,
            "expenditure_minus_gaikyo": None if gaikyo_out is None else exp6_sum - gaikyo_out,
            "kenzen_vs_gaikyo": rec_src.get("kenzen_match"),
        })

    # national totals from 都市/町村 概況
    nat = {}
    for key, label in (("gaikyo_shi", "都市"), ("gaikyo_cho", "町村")):
        by_code, totals, prefs = indexed[key]
        muni_sum_in = 0
        muni_sum_out = 0
        n_muni = 0
        for code, (i, row) in by_code.items():
            inn = to_num(row[39])
            outv = to_num(row[40])
            if inn is not None:
                muni_sum_in += inn
                n_muni += 1
            if outv is not None:
                muni_sum_out += outv
        nat_row = totals.get("全国計")
        nat[label] = {
            "muni_rows_6digit": len(by_code),
            "muni_with_revenue": n_muni,
            "sum_歳入総額": muni_sum_in,
            "sum_歳出総額": muni_sum_out,
            "全国計_row": None if nat_row is None else nat_row[0],
            "全国計_歳入総額": None if nat_row is None else to_num(nat_row[1][39]),
            "全国計_歳出総額": None if nat_row is None else to_num(nat_row[1][40]),
            "pref_header_sample": prefs[:3],
        }

    gifu_rev = sum(v["revenue_r6_sum"] for v in verify_rows)
    gifu_gaikyo_in = sum(v["歳入総額_概況"] or 0 for v in verify_rows)
    gifu_exp = sum(v["expenditure_r6_sum"] for v in verify_rows)
    gifu_gaikyo_out = sum(v["歳出総額_概況"] or 0 for v in verify_rows)
    mismatches_in = [v for v in verify_rows if v["revenue_minus_gaikyo"] not in (0, 0.0, None)]
    mismatches_out = [v for v in verify_rows if v["expenditure_minus_gaikyo"] not in (0, 0.0, None)]

    gifu_shi = next(v for v in verify_rows if v["area_code"] == "212016")
    gifu_rec = next(r for r in data if r["area_code"] == "212016")

    kenzen_mismatch = []
    for v in verify_rows:
        km = v.get("kenzen_vs_gaikyo")
        if not km:
            continue
        rec = next(r for r in data if r["area_code"] == v["area_code"])
        for field in ("実質公債費比率", "将来負担比率", "実質赤字比率", "連結実質赤字比率"):
            if km.get(field) != rec.get(field):
                kenzen_mismatch.append({"area_code": v["area_code"], "field": field, "gaikyo": rec.get(field), "kenzen": km.get(field)})

    r6_keys_present = sorted({k for r in data for k in r["revenue_r6"]})
    r5_keys_present = sorted({k for r in data for k in r["revenue_r5"]})

    verification = {
        "gifu42_count": len(data),
        "missing": {k: v for k, v in missing.items() if v},
        "gifu_shi": {
            "area_code": "212016",
            "excel_rows": gifu_rec["excel_rows"],
            "revenue_r6_地方税": gifu_rec["revenue_r6"].get("地方税"),
            "revenue_r6_地方交付税": gifu_rec["revenue_r6"].get("地方交付税"),
            "revenue_r6_国庫支出金": gifu_rec["revenue_r6"].get("国庫支出金"),
            "revenue_r6_地方債": gifu_rec["revenue_r6"].get("地方債"),
            "expenditure_r6_議会費": gifu_rec["expenditure_purpose_r6"].get("議会費"),
            "expenditure_r6_民生費": gifu_rec["expenditure_purpose_r6"].get("民生費"),
            "expenditure_r6_公債費": gifu_rec["expenditure_purpose_r6"].get("公債費"),
            "実質公債費比率": gifu_rec["実質公債費比率"],
            "将来負担比率": gifu_rec["将来負担比率"],
            "ratio_source_tokens": gifu_rec["ratio_source_tokens"],
            "実質赤字比率_早期健全化基準": gifu_rec["実質赤字比率_早期健全化基準"],
            "連結実質赤字比率_早期健全化基準": gifu_rec["連結実質赤字比率_早期健全化基準"],
            "歳入総額_概況": gifu_shi["歳入総額_概況"],
            "歳出総額_概況": gifu_shi["歳出総額_概況"],
            "revenue_r6_sum": gifu_shi["revenue_r6_sum"],
            "expenditure_r6_sum": gifu_shi["expenditure_r6_sum"],
        },
        "internal_balance_r6": {
            "gifu42_revenue_sum": gifu_rev,
            "gifu42_gaikyo_revenue_sum": gifu_gaikyo_in,
            "gifu42_expenditure_sum": gifu_exp,
            "gifu42_gaikyo_expenditure_sum": gifu_gaikyo_out,
            "revenue_mismatch_count": len(mismatches_in),
            "expenditure_mismatch_count": len(mismatches_out),
            "revenue_mismatches_sample": mismatches_in[:5],
            "expenditure_mismatches_sample": mismatches_out[:5],
        },
        "national_totals_gaikyo": nat,
        "kenzen_vs_gaikyo_mismatch_count": len(kenzen_mismatch),
        "kenzen_vs_gaikyo_mismatches": kenzen_mismatch[:10],
        "revenue_keys_present_r6": r6_keys_present,
        "revenue_keys_present_r5": r5_keys_present,
        "dash_将来負担比率_count": sum(1 for r in data if r["ratio_source_tokens"].get("将来負担比率") == "-"),
        "dash_実質赤字比率_count": sum(1 for r in data if r["ratio_source_tokens"].get("実質赤字比率") == "-"),
    }

    # 原本ごとの取得日・ハッシュ。fetch_m2_2_raw.py が作る fetch_raw.json があればそれを使う。
    # 無い場合（リポジトリには入っていない中間ファイル）は、既存の data/m2_2_finance.json の sources に記録された
    # 取得日を引き継ぎ、ハッシュは原本ファイルから計算し直す（記録と違えば新しい版として今日の日付にする）。
    fetch_json = INSPECT / "fetch_raw.json"
    if fetch_json.exists():
        fetch = json.loads(fetch_json.read_text(encoding="utf-8"))
        fetch_by = {x["path"]: x for x in fetch if "path" in x}
    else:
        import hashlib
        prev = {}
        if OUT_JSON.exists():
            for old in json.loads(OUT_JSON.read_text(encoding="utf-8")).get("sources", []):
                prev[old["source_url"].rsplit("/", 1)[-1]] = old
        fetch_by = {}
        for f in sorted(RAW.iterdir()):
            h = hashlib.sha256(f.read_bytes()).hexdigest()
            old = prev.get(f.name)
            if old and old.get("source_file_hash") == h:
                fetch_by[f.name] = {"sha256": h, "fetched_at": old["fetched_at"]}
            else:
                print("注意: 取得記録が無い、またはハッシュが変わった原本:", f.name)
                fetch_by[f.name] = {"sha256": h, "fetched_at": date.today().isoformat()}
    page = "https://www.soumu.go.jp/iken/zaisei/r06_shichouson.html"
    page_r5 = "https://www.soumu.go.jp/iken/zaisei/r05_shichouson.html"
    page_k = "https://www.soumu.go.jp/iken/zaisei/kenzenka/kenzen_shikin.html"

    def src(role, fname, url, page_url, fetched_at=None):
        meta = fetch_by[fname]
        rec = {
            "role": role,
            "source_url": url,
            "source_page": page_url,
            "fetched_at": fetched_at or meta["fetched_at"],
            "source_file_hash": meta["sha256"],
        }
        return rec

    sources = [
        src("R6資料のみかた", "001061668.pdf", "https://www.soumu.go.jp/main_content/001061668.pdf", page),
        src("R6都市概況", "001061669.xlsx", "https://www.soumu.go.jp/main_content/001061669.xlsx", page),
        src("R6都市歳入内訳", "001061670.xlsx", "https://www.soumu.go.jp/main_content/001061670.xlsx", page),
        src("R6都市目的別歳出", "001061671.xlsx", "https://www.soumu.go.jp/main_content/001061671.xlsx", page),
        src("R6町村概況", "001061674.xlsx", "https://www.soumu.go.jp/main_content/001061674.xlsx", page),
        src("R6町村歳入内訳", "001061675.xlsx", "https://www.soumu.go.jp/main_content/001061675.xlsx", page),
        src("R6町村目的別歳出", "001061676.xlsx", "https://www.soumu.go.jp/main_content/001061676.xlsx", page),
        src("R5都市歳入内訳", "000999901.xlsx", "https://www.soumu.go.jp/main_content/000999901.xlsx", page_r5),
        src("R5町村歳入内訳", "000999906.xlsx", "https://www.soumu.go.jp/main_content/000999906.xlsx", page_r5),
        src("R5都市目的別歳出", "000999902.xlsx", "https://www.soumu.go.jp/main_content/000999902.xlsx", page_r5),
        src("R5町村目的別歳出", "000999908.xlsx", "https://www.soumu.go.jp/main_content/000999908.xlsx", page_r5),
        src("R6健全化判断比率概要", "001060123.xlsx", "https://www.soumu.go.jp/main_content/001060123.xlsx", page_k),
    ]

    payload = {
        "year": "2024年度（令和6年度）",
        "prior_year": "2023年度（令和5年度）",
        "updated": "2026-08-30",
        "unit_revenue_expenditure": "千円（ソースヘッダ『（単位　千円）』）",
        "unit_ratio": "％（概況は比率列、001060123は『（単位：％）』）",
        "code_format": "6桁（全国地方公共団体コード、検査数字付き。例 212016=岐阜市）",
        "revenue_kan_keys_r6": REVENUE_KAN_KEYS_R6,
        "revenue_kan_keys_r5": REVENUE_KAN_KEYS_R5,
        "expenditure_purpose_keys": EXPENDITURE_PURPOSE_KEYS,
        "record_keys": RECORD_KEYS,
        "statutory_criteria_common": {
            "source": "001061668.pdf 資料のみかた（参考）健全化判断比率に係る早期健全化基準等",
            "実質赤字比率_財政再生基準_市区町村": 20,
            "連結実質赤字比率_財政再生基準_市区町村": 30,
            "実質公債費比率_早期健全化基準": 25,
            "実質公債費比率_財政再生基準": 35,
            "将来負担比率_早期健全化基準_市区町村": 350,
            "将来負担比率_財政再生基準": None,
            "将来負担比率_財政再生基準_note": "資料のみかたに『－』（設定なし）",
        },
        "prefecture_layer": {
            "in_shichoson_kessan_files": "無（市町村別ファイル。都道府県名の見出し行はあるが団体コード付きの岐阜県行は無い）",
            "in_kenzen_001060123_都道府県シート": kenzen_pref,
        },
        "sources": sources,
        "verification": {
            "gifu42_count": verification["gifu42_count"],
            "gifu_shi": verification["gifu_shi"],
            "internal_balance_r6": {
                "gifu42_revenue_sum": verification["internal_balance_r6"]["gifu42_revenue_sum"],
                "gifu42_gaikyo_revenue_sum": verification["internal_balance_r6"]["gifu42_gaikyo_revenue_sum"],
                "gifu42_expenditure_sum": verification["internal_balance_r6"]["gifu42_expenditure_sum"],
                "gifu42_gaikyo_expenditure_sum": verification["internal_balance_r6"]["gifu42_gaikyo_expenditure_sum"],
                "revenue_mismatch_count": verification["internal_balance_r6"]["revenue_mismatch_count"],
                "expenditure_mismatch_count": verification["internal_balance_r6"]["expenditure_mismatch_count"],
            },
            "national_totals_gaikyo": verification["national_totals_gaikyo"],
            "kenzen_vs_gaikyo_mismatch_count": verification["kenzen_vs_gaikyo_mismatch_count"],
            "dash_将来負担比率_count": verification["dash_将来負担比率_count"],
            "dash_実質赤字比率_count": verification["dash_実質赤字比率_count"],
        },
        "notes": [
            "2-2a/2-2b/2-2eは市町村別決算状況調の歳入内訳・目的別歳出。単位は各ファイルヘッダ『（単位　千円）』。",
            "労働費・諸支出金は目的別歳出に款合計列が無く、項列（労働費=col35+36、諸支出金=col91+92+93）の合計。他の目的は款合計列。",
            "R5歳入は『特別地方消費税交付金』が款として残る。R6歳入ヘッダに同名款は無い。2-2eは款名で突合する。",
            "2-2c/2-2d/2-2fはR6概況の健全化判断比率列。『－』は null にし ratio_source_tokens にソース表記を残した。",
            "2-2gの団体別早期健全化基準（実質赤字・連結実質赤字）は001060123の市区町村シート col6/col8。ヘッダ行に列名は無いが、値の範囲が資料のみかたの11.25〜15％／16.25〜20％と一致する。",
            "固定の法定基準（公債費25/35、将来負担350、赤字再生20、連結再生30）は資料のみかたの参考表。団体別表ではない。",
            "e-Stat 00200251はFY2018終端のため再探索していない。",
        ],
        "data": data,
    }
    ordered = {k: payload[k] for k in TOP_KEYS}
    OUT_JSON.write_text(json.dumps(ordered, ensure_ascii=False, indent=2), encoding="utf-8")
    INSPECT.mkdir(parents=True, exist_ok=True)
    VERIFY_PATH.write_text(json.dumps(verification, ensure_ascii=False, indent=2), encoding="utf-8")
    print("JSON", OUT_JSON, "records", len(data), "bytes", OUT_JSON.stat().st_size)
    print("VERIFY", VERIFY_PATH)
    print("missing", {k: v for k, v in missing.items() if v})
    print("rev_mismatch", len(mismatches_in), "exp_mismatch", len(mismatches_out), "kenzen_mismatch", len(kenzen_mismatch))
    print("gifu_shi", verification["gifu_shi"])


def header_map_from_rows(rows, row_idx: int) -> dict[str, int]:
    row = rows[row_idx - 1]
    out: dict[str, int] = {}
    for i, v in enumerate(row):
        key = norm(v)
        if key and key not in out:
            out[key] = i
    return out


if __name__ == "__main__":
    main()
