# -*- coding: utf-8 -*-
"""まちのすがた「昼夜間人口比率」市区町村版JSONに全国・類似団体比較(comparison)を追加する。

対象（content-spec-v1.md §7.2 / docs/similar-group-aggregation-design.md）:
  1-2b 昼夜間人口比率（岐阜42市区町村・比較値=全国平均・類似団体平均・順位・百分位つき）
  content-spec-v1.md §7.2 の定義：1-2b の主比較は「比較=類団」（類似団体平均）。

入力（取得済みローカルファイルのみ。ネットワークアクセスをしない）:
  poc/recon-20260828/1-2b/estat_000032214141_kokusei2020_daynight.xlsx
    （統計表ID 000032214141・第１－１表。既存 data/m1_2b_daynight.json の生値と同一の全国集約ファイル）
  poc/recon-20260828/0/ruikei_kubun_r06.xlsx（類似団体区分マスタ・市町村一覧シート・1741団体）

更新（既存ファイルを書き換え・新規ファイルは作らない）:
  data/m1_2b_daynight.json の各レコード（42件）に comparison フィールドを追加する。
  --pref 23 など（D-24 の試作）：data/m1_2b_daynight_23.json（build_kokusei_json.py --pref 23 の出力）に追加する。
  対象団体は自治体マスタの市区町村行（scripts/pref_args.py）。全国の比較値は県に依らず同じ母集団Mで算出する。

集計（docs/similar-group-aggregation-design.md §3・§4・§9・§10。m1_2c_taxable_income.py と同一の実装パターン）:
  母集団M = 区分マスタに載る団体（1741）∩ 全国集約ファイルに値がある団体（'-'等の非該当値を除く）。
  束ね写像：政令指定都市・中核市・施行時特例市・特別区 → "大都市"。他は類団区分の生ラベルをそのまま使う。
  順位：降順 competition ranking（同値は同順位、次の順位を飛ばす）。
  percentile = (rank_out_of − rank + 1) / rank_out_of × 100（小数第1位）。
  average_method = "unweighted_mean"（単純算術平均。人口加重等はしない）。

規律:
  - 列位置はハードコードせず、ヘッダのラベル文字列（地域名・男女・年齢・昼夜間人口比率）から同定し自己検証する。
  - 表題（第１－１表・昼夜間人口比率）をコードでassertしてから使う（推測しない）。
  - 国勢調査は5桁コード。区分マスタは6桁のため、区分マスタ側を code6[:5] に落として結合する
    （§10の規律。地域名セルは "21201_岐阜市" 形式のため正規表現 ^(\\d{5})_ で5桁コードを抜き出す）。
  - 評価語・評価色・複合スコアは付けない（中立性・D-4/D-6/D-18）。
"""
from __future__ import annotations

import json
import re
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from pref_args import load_master_cities, out_path, parse_pref  # noqa: E402
from ward_compare import ward_comparison  # noqa: E402

PREF = parse_pref(sys.argv[1:])

DAYNIGHT_SRC = ROOT / "poc" / "recon-20260828" / "1-2b" / "estat_000032214141_kokusei2020_daynight.xlsx"
RUIKEI_SRC = ROOT / "poc" / "recon-20260828" / "0" / "ruikei_kubun_r06.xlsx"
OUT = out_path("m1_2b_daynight", PREF)

UPDATED = date.today().isoformat()

DAITOSHI_KUBUN = {1, 2, 8, 9}  # 政令指定都市・特別区・中核市・施行時特例市 → "大都市"

# §9実データ検証（similar-group-aggregation-design.md）の岐阜実測アンカー。
# 区分マスタのみに依存する束ね後区分・団体数のため指標に依らず共通（m1_2c_taxable_income.py と同一）。
# 一致しなければ中断する。
ANCHOR = {
    "gifu_shi_group": "大都市",
    "gifu_shi_group_n": 128,
    "ogaki_shi_group": "Ⅳ－２",
    "ogaki_shi_group_n": 75,
    "takayama_shi_group": "Ⅱ－１",
    "takayama_shi_group_n": 127,
}


class Abort(Exception):
    """欠損・単位不一致・アンカー不一致などで処理を止めるための例外。推測で続行しない。"""


def s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def cell(row: tuple, c: int):
    if row is None or c < 0 or c >= len(row):
        return None
    return row[c]


# --------------------------------------------------------------------------------------
# 区分マスタ（類似団体区分）
# --------------------------------------------------------------------------------------
def load_ruikei_master() -> dict[str, dict]:
    from openpyxl import load_workbook

    wb = load_workbook(RUIKEI_SRC, read_only=True, data_only=True)
    sheet_name = next((n for n in wb.sheetnames if "市町村一覧" in n), None)
    if sheet_name is None:
        raise Abort(f"ruikei: シート『市町村一覧』が見つからない（実在={wb.sheetnames}）")
    ws = wb[sheet_name]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    print(f"  [ruikei] engine=openpyxl sheet={sheet_name!r}")
    hidx = None
    for i, row in enumerate(rows):
        if s(row[0] if row else None) == "団体コード":
            hidx = i
            break
    if hidx is None:
        raise Abort("ruikei: ヘッダ行（1列目='団体コード'）が見つからない")

    out: dict[str, dict] = {}
    for row in rows[hidx + 1 :]:
        code = s(row[0]) if len(row) > 0 else ""
        if not re.fullmatch(r"\d{6}", code):
            continue
        name = s(row[1])
        danntai_kubun_raw = row[2]
        ruikei_raw = s(row[4])
        if not ruikei_raw:
            continue
        try:
            danntai_kubun = int(danntai_kubun_raw)
        except (TypeError, ValueError):
            raise Abort(f"ruikei: {code} の団体区分が数値でない: {danntai_kubun_raw!r}")
        group_code = "大都市" if danntai_kubun in DAITOSHI_KUBUN else ruikei_raw
        out[code] = {"name": name, "danntai_kubun": danntai_kubun, "ruikei_raw": ruikei_raw, "group_code": group_code}
    if len(out) != 1741:
        raise Abort(f"ruikei: 市町村一覧の件数が1741ではない（実測={len(out)}）")
    print(f"  [ruikei] 市町村一覧 {len(out)}件を読み込み")
    return out


# --------------------------------------------------------------------------------------
# 全国昼夜間人口比率（国勢調査 第１－１表）
# --------------------------------------------------------------------------------------
def resolve_daynight_columns(ws):
    """(地域名列, 男女列, 年齢列, 昼夜間人口比率列, データ開始行) を返す。列位置は自己同定する。"""
    header_rows = list(ws.iter_rows(min_row=1, max_row=12, values_only=True))

    label_row_idx = None
    for i, row in enumerate(header_rows):
        if any(s(v) == "地域識別コード" for v in row):
            label_row_idx = i
            break
    if label_row_idx is None:
        raise Abort("daynight: ヘッダ行（『地域識別コード』を含む行）が見つからない")
    label_row = header_rows[label_row_idx]

    def col_of(label: str):
        for c, v in enumerate(label_row):
            if s(v) == label:
                return c
        return None

    area_col = col_of("地域名")
    sex_col = col_of("男女")
    age_col = col_of("年齢")
    if area_col is None or sex_col is None or age_col is None:
        raise Abort(
            f"daynight: ヘッダ列（地域名/男女/年齢）が見つからない area={area_col} sex={sex_col} age={age_col}"
        )

    ratio_col = None
    for i in range(0, label_row_idx):
        for c, v in enumerate(header_rows[i]):
            if s(v) == "昼夜間人口比率":
                ratio_col = c
                break
        if ratio_col is not None:
            break
    if ratio_col is None:
        raise Abort("daynight: 『昼夜間人口比率』列が見つからない")

    title_texts = " ".join(s(v) for row in header_rows[:3] for v in row if v)
    if "第１－１表" not in title_texts or "昼夜間人口比率" not in title_texts:
        raise Abort(f"daynight: 表題に『第１－１表』『昼夜間人口比率』を確認できない（実測={title_texts!r}）。推測で続行しない。")
    print(f"  [daynight] 表題確認: {title_texts[:60]}...")

    data_start = label_row_idx + 1
    return area_col, sex_col, age_col, ratio_col, data_start


def load_daynight_by_code5() -> dict[str, float]:
    from openpyxl import load_workbook

    wb = load_workbook(DAYNIGHT_SRC, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    area_col, sex_col, age_col, ratio_col, data_start = resolve_daynight_columns(ws)
    print(
        f"  [daynight] 列: 地域名列=c{area_col + 1} 男女列=c{sex_col + 1} "
        f"年齢列=c{age_col + 1} 昼夜間人口比率列=c{ratio_col + 1} データ開始行=r{data_start + 1}"
    )

    out: dict[str, float] = {}
    n_scanned = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < data_start:
            continue
        n_scanned += 1
        if s(cell(row, sex_col)) != "0_総数" or s(cell(row, age_col)) != "00_総数":
            continue
        area_v = s(cell(row, area_col))
        m = re.match(r"^(\d{5})_", area_v)
        if not m:
            continue
        code5 = m.group(1)
        if code5 == "00000" or code5[-3:] == "000":
            continue  # 全国・都道府県の集計行を除く（市区町村のみ残す）
        ratio_v = cell(row, ratio_col)
        if not isinstance(ratio_v, (int, float)) or isinstance(ratio_v, bool):
            continue  # '-' 等の非該当値
        if code5 in out:
            raise Abort(f"daynight: コード{code5}が重複している")
        out[code5] = float(ratio_v)
    print(f"  [daynight] 全{n_scanned}行走査・市区町村{len(out)}件抽出")
    return out


# --------------------------------------------------------------------------------------
# 集計
# --------------------------------------------------------------------------------------
def ranks_desc(values: list[float]) -> list[int]:
    """降順の competition ranking（同値は同順位、次の順位を飛ばす。1,2,2,4）。"""
    return [1 + sum(1 for w in values if w > v) for v in values]


def main():
    for p in (DAYNIGHT_SRC, RUIKEI_SRC, OUT):
        if not p.exists():
            raise Abort(f"入力ファイルが無い: {p}")

    print("== 既存JSON読み込み ==")
    payload = json.loads(OUT.read_text(encoding="utf-8"))
    existing = payload["data"]
    cities = load_master_cities(PREF)
    if len(existing) != len(cities):
        raise Abort(f"既存JSONのレコード数が{len(cities)}（自治体マスタ）ではない（実測={len(existing)}）")

    print("== 区分マスタ読み込み ==")
    ruikei = load_ruikei_master()

    print("== 全国昼夜間人口比率ファイル読み込み ==")
    daynight_by_code5 = load_daynight_by_code5()

    print("== 母集団Mの構築 ==")
    population: dict[str, dict] = {}  # code6 -> {value, group_code, name}
    skipped_no_source = []
    for code6, meta in ruikei.items():
        code5 = code6[:5]
        v = daynight_by_code5.get(code5)
        if v is None:
            skipped_no_source.append(code6)
            continue
        population[code6] = {"value": v, "group_code": meta["group_code"], "name": meta["name"]}
    print(f"  母集団M = {len(population)}件（区分マスタ1741 − 本ファイルに無い{len(skipped_no_source)}件）")

    codes = list(population.keys())
    values = [population[c]["value"] for c in codes]
    n_total = len(codes)
    national_average = round(sum(values) / n_total, 1)

    group_values: dict[str, list[float]] = {}
    for c in codes:
        group_values.setdefault(population[c]["group_code"], []).append(population[c]["value"])
    group_average = {g: round(sum(vs) / len(vs), 1) for g, vs in group_values.items()}
    group_n = {g: len(vs) for g, vs in group_values.items()}

    ranks = dict(zip(codes, ranks_desc(values)))

    print("== アンカー自己検証（§9実データ検証・2026-08-29の区分構成に基づく） ==")

    def check(label, got, want):
        ok = got == want
        print(f"    {'OK ' if ok else 'NG '} {label}: 実測={got!r} 期待={want!r}")
        return ok

    gifu_shi = population.get("212016")
    ogaki_shi = population.get("212024")
    takayama_shi = population.get("212032")
    if not (gifu_shi and ogaki_shi and takayama_shi):
        raise Abort("岐阜市/大垣市/高山市のいずれかが母集団Mに無い")
    results = [
        check("岐阜市 束ね後区分", gifu_shi["group_code"], ANCHOR["gifu_shi_group"]),
        check("岐阜市 束ね後区分の団体数", group_n[gifu_shi["group_code"]], ANCHOR["gifu_shi_group_n"]),
        check("大垣市 束ね後区分", ogaki_shi["group_code"], ANCHOR["ogaki_shi_group"]),
        check("大垣市 束ね後区分の団体数", group_n[ogaki_shi["group_code"]], ANCHOR["ogaki_shi_group_n"]),
        check("高山市 束ね後区分", takayama_shi["group_code"], ANCHOR["takayama_shi_group"]),
        check("高山市 束ね後区分の団体数", group_n[takayama_shi["group_code"]], ANCHOR["takayama_shi_group_n"]),
    ]
    if not all(results):
        raise Abort("アンカー値が一致しない。JSONを書かずに中断した。")

    print(f"== 既存{len(existing)}件との整合確認 ==")
    gifu_codes = sorted(c for c in codes if c in cities)
    if set(gifu_codes) != set(cities):
        raise Abort(f"母集団Mの県内が{len(gifu_codes)}件（自治体マスタ{len(cities)}件と一致しない）")

    by_area_code = {r["area_code"]: r for r in existing}
    if set(by_area_code.keys()) != set(gifu_codes):
        raise Abort("既存JSONのarea_codeと母集団Mの県内のキー集合が一致しない。JSONを書かずに中断した。")

    print("== 既存レコードへ comparison を付与（生値の再抽出一致も併せて検証） ==")
    for rec in existing:
        code6 = rec["area_code"]
        p = population[code6]
        if abs(rec["昼夜間人口比率"] - p["value"]) > 1e-6:
            raise Abort(
                f"{rec['市区町村名']}（{code6}）: 既存生値{rec['昼夜間人口比率']}と全国ファイルからの再抽出値"
                f"{p['value']}が不一致。JSONを書かずに中断した。"
            )
        rank = ranks[code6]
        pct = round((n_total - rank + 1) / n_total * 100, 1)
        rec["comparison"] = {
            "value": p["value"],
            "national_average": national_average,
            "group_code": p["group_code"],
            "group_average": group_average[p["group_code"]],
            "rank": rank,
            "rank_out_of": n_total,
            "percentile": pct,
            "average_method": "unweighted_mean",
            "source_note": (
                f"全国{n_total}市区町村中の順位（母集団=類似団体区分マスタ登載団体1741のうち、"
                "本統計に数値がある団体）。類似団体平均は総務省の類似団体区分（束ね後）の単純算術平均。"
            ),
        }

    # 政令市の区（試作のみ）：区の値を母集団Mに当てはめた相当順位・親の市の区分の平均
    for w in payload.get("wards", []):
        w["comparison"] = ward_comparison(
            ruikei, {c: p["value"] for c, p in population.items()}, w["昼夜間人口比率"], w["parent_area_code"]
        )
    if payload.get("wards"):
        print(f"  区{len(payload['wards'])}件に相当順位つきcomparisonを付与")

    payload["updated"] = UPDATED
    payload.setdefault("notes", [])
    stale_note = "本ファイルは生値のみを格納する。比較値（全国平均・類団平均・全国順位）は含めず、集計段で付与する。"
    comparison_note_prefix = "comparison（全国平均・類似団体平均・全国順位・百分位）は"
    payload["notes"] = [
        n for n in payload["notes"] if n != stale_note and not n.startswith(comparison_note_prefix)
    ]
    payload["notes"].append(
        "comparison（全国平均・類似団体平均・全国順位・百分位）は2026-09-26追加（F-5・A-14）。"
        "docs/similar-group-aggregation-design.md §3・§4の集計段の定義に従う。"
        f"母集団M = 区分マスタ（総務省 類似団体別市町村財政指数表・{len(ruikei)}団体）"
        f"∩ 全国集約ファイル（本ファイル冒頭sourcesと同一の統計表ID000032214141）に数値がある団体 = {n_total}団体。"
        "national_average は母集団Mの単純算術平均、group_average は束ね後区分ごとの単純算術平均"
        "（average_method=\"unweighted_mean\"）。束ね写像：政令指定都市・中核市・施行時特例市・特別区は『大都市』"
        "に統合。それ以外は類団区分の原ラベルをそのまま使う（一般市・町村で同一ラベルの区分は同一グループとして束ねる）。"
        "content-spec-v1.md §7.2の定義どおり本指標の主比較は類似団体平均（比較=類団）。"
        "rank は母集団M内での値の降順順位（1=最大値。同値は同順位・次の順位を飛ばすcompetition ranking）。"
        "percentile = (rank_out_of − rank + 1) / rank_out_of × 100（小数第1位）。"
        "評価語・評価色・複合スコアは付与していない（20260808_decisions.md D-4／D-6／D-18）。"
    )

    print("== 書き出し ==")
    with OUT.open("w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)
    print(f"  wrote {OUT.relative_to(ROOT)}  ({len(existing)}件にcomparison付与)")

    print(f"== 県コード{PREF}の{len(existing)}件 昼夜間人口比率順（サニティチェック用） ==")
    ranked = sorted(existing, key=lambda r: -r["comparison"]["value"])
    for r in ranked[:3] + ranked[-3:]:
        c = r["comparison"]
        print(
            f"  {r['市区町村名']}（{r['area_code']}）: {c['value']} "
            f"全国{c['rank']}位/{c['rank_out_of']}（上位{c['percentile']}%） "
            f"区分={c['group_code']} 区分平均={c['group_average']} 全国平均={c['national_average']}"
        )


if __name__ == "__main__":
    try:
        main()
    except Abort as exc:
        print(f"\n!! 中断: {exc}", file=sys.stderr)
        sys.exit(1)
