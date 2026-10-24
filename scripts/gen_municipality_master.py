"""
自治体マスタ生成（岐阜V1）。

一次ソース: 総務省「全国地方公共団体コード」（都道府県コード及び市区町村コード）
  公表ページ: https://www.soumu.go.jp/denshijiti/code.html
出力: data/municipality_master.json（手書き禁止。本スクリプトが単一ソース）

city_slug はカナ列から frontend-architecture.md §5.4 規約で機械導出する。
既公開3スラッグ（gifu-pref / gifu-shi / mizuho-shi）は凍結。不一致なら中断。

kaigo_insurer_name は新規取得しない。既存 data/m2_kaigo_premium.json の
「構成市区町村」から引く（広域連合3連合9市町、それ以外は city_name と同一）。

実行タイミング: ビルド前の明示工程。自治体追加時に再実行し生成物をコミット。

  python scripts/gen_municipality_master.py            # マスタ生成＋指標JSONへ area_code 付与
  python scripts/gen_municipality_master.py --master-only
  python scripts/gen_municipality_master.py --attach-only
  python scripts/gen_municipality_master.py --inspect
  python scripts/gen_municipality_master.py --pref 23   # 試作（愛知。data/municipality_master_23.json に出力し、指標JSONには触れない）
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE_PAGE = "https://www.soumu.go.jp/denshijiti/code.html"
FALLBACK_XLSX_URL = "https://www.soumu.go.jp/main_content/000925835.xlsx"
CACHE_PATH = ROOT / "data" / "cache" / "soumu_lg_code.xlsx"
OUTPUT_PATH = ROOT / "data" / "municipality_master.json"
M2_PATH = ROOT / "data" / "m2_kaigo_premium.json"
M3_PATH = ROOT / "data" / "m3_expenditure.json"

PREF_FILTER = "21"
PREF_SLUG_V1 = "gifu-pref"
# 県コード → 地域スラッグ（lib/prefectures.ts の prefSlug と同じ値。D-24 共通化で --pref を足した。
# 岐阜（21）の出力は従来と同一。岐阜以外は試作で、出力先も別ファイル）
PREF_SLUGS = {"21": "gifu-pref", "23": "aichi-pref", "24": "mie-pref", "33": "okayama-pref", "34": "hiroshima-pref", "35": "yamaguchi-pref"}
GIFU_PREF_CODE = "21"
USER_AGENT = "Mozilla/5.0 (jichitai-watch municipality-master generator)"

# 公開済みスラッグ凍結（frontend-architecture.md §5.1 / §5.5）
FROZEN_CITY_SLUGS = {
    "212016": "gifu-shi",
    "212164": "mizuho-shi",
}
FROZEN_PREF_SLUG = PREF_SLUG_V1

# §5.4.2 / §5.4.3 の規約実例。生成結果がこれらと不一致なら変換器の欠陥。
SPEC_SLUG_EXAMPLES = {
    "212024": "ogaki-shi",  # 大垣 おお→o
    "212199": "gujo-shi",  # 郡上 おう→o
    "213411": "yoro-cho",  # 養老 おう→o
    "213811": "godo-cho",  # 神戸町 ごうど（こうべではない）
    "213837": "anpachi-cho",  # 安八 んは常に n（ampachi にしない）
    "215015": "sakahogi-cho",  # 坂祝
    "215040": "hichiso-cho",  # 七宗
    "215066": "shirakawa-cho",  # 白川町
    "216046": "shirakawa-mura",  # 白川村
}

# 形態素境界をまたぐ語は機械的なおう畳みを適用しない（§5.4.2）。
# 輪之内＝わ・の・うち。ノ+ウを長音と誤認すると wanochi になる。
SLUG_OVERRIDES = {
    "213829": "wanouchi-cho",
    # 2026-10-23 運営者裁定（§5.4.6 の例外：総務省カナの癖で機械の結果が読みと違う。試作のみで未公開のため直した）
    "332089": "soja-shi",  # 総社市＝そうじゃ（カナ ｿｳｼﾞﾔ は小さい字を大きく書いたもの。統計年鑑 2026 表2-4：Soja-shi）
    "332127": "setouchi-shi",  # 瀬戸内市＝せとうち（瀬戸・内の境目。統計年鑑 2026 表2-4：Setouchi-shi）
    "335860": "shinjo-son",  # 新庄村＝しんじょうそん（カナ ｼﾝｼﾞﾖｳｿﾝ は小さい字を大きく書いたもの）
}

ATTACH_TARGETS = [
    (ROOT / "data" / "m1_kodomo_iryohi.json", None),
    (ROOT / "data" / "m1_1a_population_timeseries.json", "市区町村コード"),
    (ROOT / "data" / "m1_1b_single_household.json", "市区町村コード"),
    (ROOT / "data" / "m1_1c_aging_rate.json", "市区町村コード"),
    (ROOT / "data" / "m1_2a_industry.json", "市区町村コード"),
    (ROOT / "data" / "m1_2b_daynight.json", "市区町村コード"),
]


# ------------------------------------------------------------------ hepburn (§5.4)

YOUON = {
    "キャ": "kya",
    "キュ": "kyu",
    "キョ": "kyo",
    "シャ": "sha",
    "シュ": "shu",
    "ショ": "sho",
    "チャ": "cha",
    "チュ": "chu",
    "チョ": "cho",
    "ニャ": "nya",
    "ニュ": "nyu",
    "ニョ": "nyo",
    "ヒャ": "hya",
    "ヒュ": "hyu",
    "ヒョ": "hyo",
    "ミャ": "mya",
    "ミュ": "myu",
    "ミョ": "myo",
    "リャ": "rya",
    "リュ": "ryu",
    "リョ": "ryo",
    "ギャ": "gya",
    "ギュ": "gyu",
    "ギョ": "gyo",
    "ジャ": "ja",
    "ジュ": "ju",
    "ジョ": "jo",
    "ビャ": "bya",
    "ビュ": "byu",
    "ビョ": "byo",
    "ピャ": "pya",
    "ピュ": "pyu",
    "ピョ": "pyo",
    "ティ": "ti",
    "ディ": "di",
    "ファ": "fa",
    "フィ": "fi",
    "フェ": "fe",
    "フォ": "fo",
    "ヴァ": "va",
    "ヴィ": "vi",
    "ヴェ": "ve",
    "ヴォ": "vo",
    "ウィ": "wi",
    "ウェ": "we",
    "ウォ": "wo",
    "トゥ": "tu",
    "ドゥ": "du",
    "テュ": "tyu",
    "デュ": "dyu",
    "フュ": "fyu",
}

BASIC = {
    "ア": "a",
    "イ": "i",
    "ウ": "u",
    "エ": "e",
    "オ": "o",
    "カ": "ka",
    "キ": "ki",
    "ク": "ku",
    "ケ": "ke",
    "コ": "ko",
    "サ": "sa",
    "シ": "shi",
    "ス": "su",
    "セ": "se",
    "ソ": "so",
    "タ": "ta",
    "チ": "chi",
    "ツ": "tsu",
    "テ": "te",
    "ト": "to",
    "ナ": "na",
    "ニ": "ni",
    "ヌ": "nu",
    "ネ": "ne",
    "ノ": "no",
    "ハ": "ha",
    "ヒ": "hi",
    "フ": "fu",
    "ヘ": "he",
    "ホ": "ho",
    "マ": "ma",
    "ミ": "mi",
    "ム": "mu",
    "メ": "me",
    "モ": "mo",
    "ヤ": "ya",
    "ユ": "yu",
    "ヨ": "yo",
    "ラ": "ra",
    "リ": "ri",
    "ル": "ru",
    "レ": "re",
    "ロ": "ro",
    "ワ": "wa",
    "ヲ": "o",
    "ン": "n",
    "ガ": "ga",
    "ギ": "gi",
    "グ": "gu",
    "ゲ": "ge",
    "ゴ": "go",
    "ザ": "za",
    "ジ": "ji",
    "ズ": "zu",
    "ゼ": "ze",
    "ゾ": "zo",
    "ダ": "da",
    "ヂ": "ji",
    "ヅ": "zu",
    "デ": "de",
    "ド": "do",
    "バ": "ba",
    "ビ": "bi",
    "ブ": "bu",
    "ベ": "be",
    "ボ": "bo",
    "パ": "pa",
    "ピ": "pi",
    "プ": "pu",
    "ペ": "pe",
    "ポ": "po",
    "ヴ": "vu",
}

# 接尾辞はカナ末尾から切り、ハイフン区切りで付す（§5.4.3）。長い方から。
SUFFIXES = (
    ("チョウ", "cho"),
    ("マチ", "machi"),
    ("ソン", "son"),
    ("ムラ", "mura"),
    ("ケン", "ken"),
    ("シ", "shi"),
    ("ク", "ku"),
)


def to_fullwidth_katakana(text: str) -> str:
    """半角カナ（総務省コード表の表記）を全角へ。NFKC は濁点結合も正規化する。"""
    return unicodedata.normalize("NFKC", text).strip()


def kana_to_morae(kana: str) -> list[str]:
    i = 0
    morae: list[str] = []
    while i < len(kana):
        ch = kana[i]
        if ch in {"ッ", "っ"}:
            morae.append("っ")
            i += 1
            continue
        if i + 1 < len(kana) and kana[i : i + 2] in YOUON:
            morae.append(YOUON[kana[i : i + 2]])
            i += 2
            continue
        if ch in BASIC:
            morae.append(BASIC[ch])
            i += 1
            continue
        if ch == "ー":
            if morae:
                prev = morae[-1]
                morae.append(prev[-1] if prev else "")
            i += 1
            continue
        if ch in {"・", " ", "　"}:
            i += 1
            continue
        raise ValueError(f"未対応のカナ: {ch!r} in {kana!r}")
    return morae


def collapse_long_vowels(morae: list[str]) -> list[str]:
    """おう・おお→o、うう→u（§5.4.2）。撥音 n は触らない。"""
    out: list[str] = []
    i = 0
    n = len(morae)
    while i < n:
        m = morae[i]
        nxt = morae[i + 1] if i + 1 < n else None
        if m == "っ":
            if nxt and nxt != "っ":
                doubled = nxt[0] if not nxt.startswith("ch") else "t"
                out.append(doubled + nxt)
                i += 2
                continue
            i += 1
            continue
        if nxt == "o" and m.endswith("o"):
            out.append(m)
            i += 2
            continue
        if nxt == "u" and m.endswith("o"):
            out.append(m)
            i += 2
            continue
        if nxt == "u" and m.endswith("u"):
            out.append(m)
            i += 2
            continue
        out.append(m)
        i += 1
    return out


def split_suffix(kana: str) -> tuple[str, str | None]:
    for kana_suf, romaji_suf in SUFFIXES:
        if kana.endswith(kana_suf) and len(kana) > len(kana_suf):
            return kana[: -len(kana_suf)], romaji_suf
    return kana, None


def kana_to_city_slug(kana: str, area_code: str) -> str:
    if area_code in SLUG_OVERRIDES:
        return SLUG_OVERRIDES[area_code]
    kana_n = to_fullwidth_katakana(kana)
    if not kana_n:
        return ""
    stem_kana, suffix = split_suffix(kana_n)
    stem = "".join(collapse_long_vowels(kana_to_morae(stem_kana)))
    slug = f"{stem}-{suffix}" if suffix else stem
    slug = slug.lower()
    if not re.fullmatch(r"[a-z]+(?:-[a-z]+)*", slug):
        raise ValueError(f"スラッグが規約外: {slug!r} (kana={kana_n!r}, code={area_code})")
    return slug


# ------------------------------------------------------------------ download / parse

def resolve_xlsx_url() -> str:
    print(f"[fetch] {SOURCE_PAGE}")
    resp = requests.get(
        SOURCE_PAGE,
        timeout=60,
        headers={"User-Agent": USER_AGENT},
    )
    resp.raise_for_status()
    hrefs = re.findall(r'href=["\']([^"\']+\.xlsx)["\']', resp.text, flags=re.I)
    abs_urls: list[str] = []
    for href in hrefs:
        if href.startswith("http"):
            abs_urls.append(href)
        else:
            abs_urls.append("https://www.soumu.go.jp" + (href if href.startswith("/") else "/" + href))
    # コード表が改正一覧表より先に出る（ページ構成）。最初の xlsx を採る。
    if abs_urls:
        print(f"[xlsx] {abs_urls[0]}")
        return abs_urls[0]
    print(f"[xlsx] page parse failed; fallback {FALLBACK_XLSX_URL}")
    return FALLBACK_XLSX_URL


def download_xlsx(url: str, cache: Path) -> Path:
    if cache.exists() and cache.stat().st_size > 0:
        print(f"[cache] {cache}")
        return cache
    print(f"[download] {url}")
    resp = requests.get(url, timeout=120, headers={"User-Agent": USER_AGENT})
    resp.raise_for_status()
    cache.parent.mkdir(parents=True, exist_ok=True)
    cache.write_bytes(resp.content)
    print(f"[saved] {cache} ({len(resp.content) // 1024} KB)")
    return cache


def cell_str(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def normalize_code(raw: object) -> str | None:
    s = cell_str(raw).replace(" ", "").replace("　", "")
    if s.endswith(".0"):
        s = s[:-2]
    if not re.fullmatch(r"\d{5,6}", s):
        return None
    return s.zfill(6) if len(s) == 5 else s


def parse_effective_from(wb) -> str:
    """シート名やタイトルからコード表の有効日を取る。取れなければファイル更新の西暦推定をしない。"""
    for name in wb.sheetnames:
        m = re.search(r"R(\d+)\.(\d+)\.(\d+)", name)
        if m:
            reiwa = int(m.group(1))
            month = int(m.group(2))
            day = int(m.group(3))
            year = 2018 + reiwa
            return date(year, month, day).isoformat()
        m = re.search(r"令和\s*(\d+)\s*年\s*(\d+)\s*月\s*(\d+)\s*日", name)
        if m:
            year = 2018 + int(m.group(1))
            return date(year, int(m.group(2)), int(m.group(3))).isoformat()
    ws = wb[wb.sheetnames[0]]
    for r in range(1, min(8, (ws.max_row or 1) + 1)):
        for c in range(1, min(8, (ws.max_column or 1) + 1)):
            text = cell_str(ws.cell(r, c).value)
            m = re.search(r"令和\s*(\d+)\s*年\s*(\d+)\s*月\s*(\d+)\s*日", text)
            if m:
                year = 2018 + int(m.group(1))
                return date(year, int(m.group(2)), int(m.group(3))).isoformat()
            m = re.search(r"(20\d{2})[./-](\d{1,2})[./-](\d{1,2})", text)
            if m:
                return date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
    raise SystemExit("総務省コード表から有効日（effective_from）を読めませんでした")


def inspect_sheet(ws, max_rows: int = 12) -> None:
    print(f"\nシート: '{ws.title}'  max_row={ws.max_row}  max_col={ws.max_column}")
    for r in range(1, min(max_rows + 1, (ws.max_row or 1) + 1)):
        cells = []
        for c in range(1, min((ws.max_column or 1) + 1, 8)):
            cells.append(f"[{c}]={cell_str(ws.cell(r, c).value)[:24]!r}")
        print(f"  行{r:02d}: {'  '.join(cells)}")


def detect_columns(ws) -> dict[str, int]:
    """ヘッダー行から列番号を取る。表記ゆれ（漢字/カナ・団体コード）に耐える。"""
    for r in range(1, min(16, (ws.max_row or 1) + 1)):
        labels = {c: cell_str(ws.cell(r, c).value) for c in range(1, (ws.max_column or 1) + 1)}
        joined = " ".join(labels.values())
        if "団体コード" not in joined and "コード" not in joined:
            continue
        mapping: dict[str, int] = {}
        for c, lab in labels.items():
            if "団体コード" in lab or lab in {"コード", "code"}:
                mapping.setdefault("code", c)
            elif "都道府県名" in lab and "カナ" in lab:
                mapping["pref_kana"] = c
            elif "市区町村名" in lab and "カナ" in lab:
                mapping["city_kana"] = c
            elif "都道府県名" in lab:
                mapping.setdefault("pref_name", c)
            elif "市区町村名" in lab:
                mapping.setdefault("city_name", c)
        if {"code", "pref_name", "city_name", "pref_kana", "city_kana"} <= mapping.keys():
            mapping["header_row"] = r
            return mapping
    # ヘッダー無しの定位置（団体コード, 都道府県漢字, 市区町村漢字, 都道府県カナ, 市区町村カナ）
    return {
        "code": 1,
        "pref_name": 2,
        "city_name": 3,
        "pref_kana": 4,
        "city_kana": 5,
        "header_row": 0,
    }


def extract_gifu_rows(ws) -> list[dict[str, str]]:
    cols = detect_columns(ws)
    start = cols["header_row"] + 1
    rows: list[dict[str, str]] = []
    seen: set[str] = set()
    for r in range(start, (ws.max_row or 0) + 1):
        code = normalize_code(ws.cell(r, cols["code"]).value)
        if not code or not code.startswith(PREF_FILTER):
            continue
        pref_name = cell_str(ws.cell(r, cols["pref_name"]).value)
        city_name = cell_str(ws.cell(r, cols["city_name"]).value)
        pref_kana = cell_str(ws.cell(r, cols["pref_kana"]).value)
        city_kana = cell_str(ws.cell(r, cols["city_kana"]).value)
        if pref_name == "" and city_name == "":
            continue
        # 政令指定都市の区は岐阜に無い。都道府県行は city_name 空。
        if code in seen:
            continue
        seen.add(code)
        rows.append(
            {
                "area_code": code,
                "pref_name": pref_name,
                "city_name": city_name,
                "pref_kana": pref_kana,
                "city_kana": city_kana,
            }
        )
    return rows


def m2_path_for(pref: str) -> Path:
    """県の介護保険料JSON。岐阜は data/m2_kaigo_premium.json、他県は _{県コード}（試作）。"""
    return M2_PATH if pref == GIFU_PREF_CODE else ROOT / "data" / f"m2_kaigo_premium_{pref}.json"


def load_kaigo_insurer_map(m2_path: Path | None = None) -> dict[str, str]:
    """M2 の構成市区町村 → 保険者名。単独保険者は呼び出し側で city_name を入れる。"""
    payload = json.loads((m2_path or M2_PATH).read_text(encoding="utf-8"))
    mapping: dict[str, str] = {}
    for rec in payload["data"]:
        members = rec.get("構成市区町村")
        if not members:
            continue
        insurer = rec["市区町村名"]
        for city in members:
            if city in mapping and mapping[city] != insurer:
                raise SystemExit(f"介護保険者の二重所属: {city} → {mapping[city]} / {insurer}")
            mapping[city] = insurer
    return mapping


def build_master_records(src_rows: list[dict[str, str]], effective_from: str) -> list[dict]:
    kaigo_map = load_kaigo_insurer_map()
    records: list[dict] = []
    for src in src_rows:
        code = src["area_code"]
        city_name = src["city_name"]
        is_pref = city_name == ""
        if is_pref:
            city_slug = FROZEN_PREF_SLUG
            kaigo = ""
        else:
            city_slug = kana_to_city_slug(src["city_kana"], code)
            kaigo = kaigo_map.get(city_name, city_name)
        records.append(
            {
                "area_code": code,
                "pref_name": src["pref_name"] or "岐阜県",
                "pref_slug": PREF_SLUG_V1,
                "city_name": city_name,
                "city_slug": city_slug,
                "kaigo_insurer_name": kaigo,
                "status": "active",
                "effective_from": effective_from,
                "effective_to": None,
                "successor_area_code": None,
            }
        )
    records.sort(key=lambda r: r["area_code"])
    return records


def validate_master(records: list[dict]) -> None:
    errors: list[str] = []
    cities = [r for r in records if r["city_name"]]
    prefs = [r for r in records if not r["city_name"]]
    if len(prefs) != 1:
        errors.append(f"都道府県レコードは1件であるべき: {len(prefs)}")
    if PREF_FILTER == GIFU_PREF_CODE and len(cities) != 42:
        errors.append(f"岐阜の市区町村は42件であるべき: {len(cities)}")
    for r in records:
        if r["pref_slug"] != FROZEN_PREF_SLUG:
            errors.append(f"pref_slug 凍結不一致: {r['area_code']} {r['pref_slug']}")
    gifu = PREF_FILTER == GIFU_PREF_CODE
    for code, expected in (FROZEN_CITY_SLUGS.items() if gifu else []):
        hit = next((r for r in records if r["area_code"] == code), None)
        if hit is None:
            errors.append(f"凍結対象がマスタに無い: {code} ({expected})")
        elif hit["city_slug"] != expected:
            errors.append(
                f"凍結スラッグ不一致: {code} {hit['city_name']} 生成={hit['city_slug']!r} 凍結={expected!r}"
            )
    for code, expected in (SPEC_SLUG_EXAMPLES.items() if gifu else []):
        hit = next((r for r in records if r["area_code"] == code), None)
        if hit is None:
            errors.append(f"規約実例の団体が無い: {code} ({expected})")
        elif hit["city_slug"] != expected:
            errors.append(
                f"規約実例不一致: {code} {hit['city_name']} 生成={hit['city_slug']!r} 期待={expected!r}"
            )
    slugs = [r["city_slug"] for r in cities]
    dup = {s for s in slugs if slugs.count(s) > 1}
    if dup:
        errors.append(f"city_slug 衝突（§5.4.5 タイブレーク未実装のまま衝突）: {sorted(dup)}")
    union_members = 0
    for r in cities:
        if r["kaigo_insurer_name"] != r["city_name"]:
            union_members += 1
    if gifu and union_members != 9:
        errors.append(f"広域連合構成市町村は9件であるべき: {union_members}")
    if errors:
        print("VALIDATION FAILED:", file=sys.stderr)
        for e in errors:
            print(f"  - {e}", file=sys.stderr)
        raise SystemExit(1)


# ------------------------------------------------------------------ attach area_code

def strip_gun_prefix(name: str) -> str:
    name = name.replace("ヶ", "ケ").replace("ヵ", "カ")
    m = re.match(r"^(.+郡)(.+)$", name)
    return m.group(2) if m else name


def code_digits(raw: object) -> str | None:
    if raw is None:
        return None
    s = str(raw).strip()
    if s.endswith(".0"):
        s = s[:-2]
    s = re.sub(r"\D", "", s)
    return s or None


def resolve_area_code(
    rec: dict,
    by_code: dict[str, dict],
    by_code5: dict[str, list[str]],
    by_name: dict[tuple[str, str], str],
    code_field: str | None,
) -> str:
    digits = code_digits(rec.get(code_field)) if code_field else None
    if digits:
        if len(digits) == 6 and digits in by_code:
            return digits
        five = digits[:5] if len(digits) >= 5 else digits
        matches = by_code5.get(five, [])
        if len(matches) == 1:
            return matches[0]
        if len(matches) > 1:
            raise SystemExit(f"5桁コードが複数のマスタに当たる: {digits} → {matches}")
    pref = rec.get("都道府県") or rec.get("pref_name") or "岐阜県"
    city = strip_gun_prefix(str(rec.get("市区町村名") or rec.get("city_name") or ""))
    key = (pref, city)
    if key in by_name:
        return by_name[key]
    raise SystemExit(f"area_code を解決できない: code={digits!r} pref={pref!r} city={city!r}")


def without_area_code(obj: object) -> object:
    if isinstance(obj, dict):
        return {k: without_area_code(v) for k, v in obj.items() if k != "area_code"}
    if isinstance(obj, list):
        return [without_area_code(x) for x in obj]
    return obj


def attach_area_codes(master_records: list[dict]) -> None:
    cities = [r for r in master_records if r["city_name"]]
    by_code = {r["area_code"]: r for r in cities}
    by_code5: dict[str, list[str]] = {}
    for r in cities:
        by_code5.setdefault(r["area_code"][:5], []).append(r["area_code"])
    by_name = {(r["pref_name"], r["city_name"]): r["area_code"] for r in cities}

    for path, code_field in ATTACH_TARGETS:
        original_text = path.read_text(encoding="utf-8")
        payload = json.loads(original_text)
        original_data = json.loads(json.dumps(payload["data"], ensure_ascii=False))
        attached_codes: list[str] = []
        new_rows: list[dict] = []
        for rec in payload["data"]:
            code = resolve_area_code(rec, by_code, by_code5, by_name, code_field)
            attached_codes.append(code)
            items = [(k, v) for k, v in rec.items() if k != "area_code"]
            new_rec = {"area_code": code}
            new_rec.update(items)
            new_rows.append(new_rec)
        payload["data"] = new_rows

        stripped = without_area_code(new_rows)
        if stripped != original_data and without_area_code(original_data) != stripped:
            # original に area_code が既にある再実行では original_data 側にもある。追加列以外を比較。
            orig_stripped = without_area_code(original_data)
            if orig_stripped != stripped:
                raise SystemExit(f"{path.name}: area_code 以外に差分が出た。書き込み中断。")

        missing = sorted(set(by_code) - set(attached_codes))
        extra = sorted(set(attached_codes) - set(by_code))
        if missing or extra:
            raise SystemExit(
                f"{path.name}: マスタとの area_code 集合が不一致 missing={missing} extra={extra}"
            )

        text = json.dumps(payload, ensure_ascii=False, indent=2)
        if original_text.endswith("\n"):
            text += "\n"
        path.write_text(text, encoding="utf-8")
        print(f"[attach] {path.name}: {len(attached_codes)} 件")

    m3 = json.loads(M3_PATH.read_text(encoding="utf-8"))
    m3_by_name = {r["市区町村名"]: r["area_code"] for r in m3["data"]}
    for frozen_code, slug in FROZEN_CITY_SLUGS.items():
        rec = by_code[frozen_code]
        m3_code = m3_by_name.get(rec["city_name"])
        if m3_code != frozen_code:
            raise SystemExit(
                f"M3 area_code 不一致: {rec['city_name']} master={frozen_code} m3={m3_code}"
            )
    print("[check] M3 既存 area_code（岐阜市 212016 / 瑞穂市 212164）と一致")


def write_master(records: list[dict], source_url: str, effective_from: str) -> None:
    today = date.today().isoformat()
    payload = {
        "updated": today,
        "source_as_of": effective_from,
        "notes": [
            f"総務省「全国地方公共団体コード」（都道府県コード及び市区町村コード）から{next(r['pref_name'] for r in records if not r['city_name'])}({PREF_FILTER})分を抽出。",
            "手書き禁止。scripts/gen_municipality_master.py が単一ソース。",
            "city_slug はカナ列から frontend-architecture.md §5.4 で機械導出。"
            + ("gifu-pref/gifu-shi/mizuho-shi は凍結検証済み。" if PREF_FILTER == GIFU_PREF_CODE else "試作（凍結検証の対象外）。"),
            "kaigo_insurer_name は data/m2_kaigo_premium.json の構成市区町村から解決（新規取得しない）。",
            "status は V1 全 active。合併リダイレクト（C-8）は実装しない。",
        ],
        "sources": [
            {
                "role": "全国地方公共団体コード",
                "source_url": source_url,
                "source_page": SOURCE_PAGE,
                "fetched_at": today,
            }
        ],
        "data": records,
    }
    OUTPUT_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"[write] {OUTPUT_PATH}  ({len(records)} 件)")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--inspect", action="store_true")
    parser.add_argument("--master-only", action="store_true")
    parser.add_argument("--attach-only", action="store_true")
    parser.add_argument(
        "--kaigo-only",
        action="store_true",
        help="既存の試作マスタ（岐阜以外）の kaigo_insurer_name だけを、その県の M2 JSON の構成市区町村から更新する（総務省表の再取得なし）",
    )
    parser.add_argument("--pref", default=GIFU_PREF_CODE, help="県コード2桁（既定 21＝岐阜。岐阜以外は試作）")
    args = parser.parse_args()

    global PREF_FILTER, PREF_SLUG_V1, FROZEN_PREF_SLUG, OUTPUT_PATH
    if args.pref not in PREF_SLUGS:
        raise SystemExit(f"未対応の県コード: {args.pref}（対応: {sorted(PREF_SLUGS)}）")
    PREF_FILTER = args.pref
    PREF_SLUG_V1 = FROZEN_PREF_SLUG = PREF_SLUGS[args.pref]
    if args.pref != GIFU_PREF_CODE:
        # 試作：岐阜の出力・指標JSONを触らない
        OUTPUT_PATH = ROOT / "data" / f"municipality_master_{args.pref}.json"
        args.master_only = True
        if args.attach_only:
            raise SystemExit("--attach-only は岐阜（21）だけ。")

    if args.kaigo_only:
        if args.pref == GIFU_PREF_CODE:
            raise SystemExit("--kaigo-only は岐阜以外（試作）だけ。岐阜は通常の生成を使う。")
        payload = json.loads(OUTPUT_PATH.read_text(encoding="utf-8"))
        kaigo_map = load_kaigo_insurer_map(m2_path_for(args.pref))
        changed = 0
        for r in payload["data"]:
            if r["city_name"] == "":
                continue
            new = kaigo_map.get(r["city_name"])
            if new is None:
                continue  # 広域連合の構成でない団体（区を含む）は触らない
            if r["kaigo_insurer_name"] != new:
                r["kaigo_insurer_name"] = new
                changed += 1
        OUTPUT_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"[kaigo-only] {OUTPUT_PATH}  更新 {changed} 件")
        return

    if args.attach_only:
        if not OUTPUT_PATH.exists():
            raise SystemExit("municipality_master.json が無い。先にマスタを生成すること。")
        records = json.loads(OUTPUT_PATH.read_text(encoding="utf-8"))["data"]
        attach_area_codes(records)
        return

    # 保存済みの表があれば総務省のページを見に行かない（この作業場所は総務省に接続できない。2026-10-19）。
    # data/cache は git 管理外なので、リポジトリに保存した写し（運営者が 2026-10-19 に添付した R6.1.1 現在の表）から補う。
    saved = ROOT / "poc" / "recon-20261019-soumu-code" / "soumu_000925835_lg_code_r6.xlsx"
    if not CACHE_PATH.exists() and saved.exists():
        CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
        CACHE_PATH.write_bytes(saved.read_bytes())
    xlsx_url = FALLBACK_XLSX_URL if CACHE_PATH.exists() and CACHE_PATH.stat().st_size > 0 else resolve_xlsx_url()
    cache = download_xlsx(xlsx_url, CACHE_PATH)
    wb = load_workbook(cache, read_only=True, data_only=True)

    if args.inspect:
        for name in wb.sheetnames:
            inspect_sheet(wb[name])
        return

    # 現在の団体一覧シートを優先（改正一覧・仕様シートを避ける）
    sheet = None
    for name in wb.sheetnames:
        if "改正" in name or "仕様" in name:
            continue
        sheet = wb[name]
        break
    if sheet is None:
        sheet = wb[wb.sheetnames[0]]

    effective_from = parse_effective_from(wb)
    src_rows = extract_gifu_rows(sheet)
    if not src_rows:
        # 2シート目（「現在の団体」）を試す
        for name in wb.sheetnames[1:]:
            src_rows = extract_gifu_rows(wb[name])
            if src_rows:
                sheet = wb[name]
                break
    print(f"[sheet] {sheet.title}  gifu rows={len(src_rows)}  effective_from={effective_from}")
    records = build_master_records(src_rows, effective_from)
    validate_master(records)
    write_master(records, xlsx_url, effective_from)

    if not args.master_only:
        attach_area_codes(records)

    if PREF_FILTER == GIFU_PREF_CODE:
        print("[ok] freeze gifu-pref / gifu-shi / mizuho-shi 一致")
    else:
        print(f"[ok] 試作 {PREF_FILTER}: {sum(1 for r in records if r['city_name'])} 団体（県1件を除く）")


if __name__ == "__main__":
    main()
